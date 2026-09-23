"""Excel workbook with the SAP2000 vs Anejo 10 (CYPECAD) comparison, with charts.

    python tools/excel_report.py [output/comparacion_SAP2000.json] [-o output/Comparacion_SAP2000_vs_CYPE.xlsx]

Reads the report written by tools/compare_sap_tables.py (or sap_oapi_run.py) and writes one
sheet per comparison, each with its bar / scatter charts.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
HEAD = PatternFill("solid", fgColor="1F4E78")
GOOD = PatternFill("solid", fgColor="E2EFDA")
WARN = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="F8CBAD")


def pile_key(p: str) -> int:
    return int(p[1:])


def write_table(ws, top: int, headers: list[str], rows: list[list], pct_col: int | None = None) -> int:
    for j, h in enumerate(headers, 1):
        c = ws.cell(top, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, r in enumerate(rows, top + 1):
        for j, v in enumerate(r, 1):
            ws.cell(i, j, v)
        if pct_col is not None and isinstance(r[pct_col - 1], (int, float)):
            d = abs(r[pct_col - 1])
            ws.cell(i, pct_col).fill = GOOD if d <= 5 else WARN if d <= 10 else BAD
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, len(headers[j - 1]) + 2)
    return top + len(rows)


def bar(ws, title: str, ytitle: str, cats, data_cols: list[int], first: int, last: int, anchor: str):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.y_axis.title = ytitle
    ch.height, ch.width = 8, 18
    for c in data_cols:
        ch.add_data(Reference(ws, min_col=c, min_row=first - 1, max_row=last), titles_from_data=True)
    ch.set_categories(cats)
    ws.add_chart(ch, anchor)


def pct(m, c):
    return round(100 * (m - c) / abs(c), 1) if c not in (None, 0) and abs(c) > 1e-9 else None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("report", nargs="?", type=Path, default=ROOT / "output" / "comparacion_SAP2000.json")
    ap.add_argument("-o", "--out", type=Path, default=ROOT / "output" / "Comparacion_SAP2000_vs_CYPE.xlsx")
    args = ap.parse_args()
    rep = json.loads(args.report.read_text(encoding="utf-8"))
    wb = Workbook()

    # ---- 1. summary ------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Comparación {rep['source']} vs Anejo 10 (CYPECAD) — Muelle de Trasmallo, módulo 40 m"
    ws["A1"].font = Font(bold=True, size=13)
    env = [r for r in rep["pile_envelopes"] if r["where"] == "base"]

    def worst(fam, what):
        rs = [r for r in env if r["family"] == fam and r["what"] == what]
        key = {"N max": lambda r: r["model"], "N min": lambda r: -r["model"]}.get(what, lambda r: abs(r["model"]))
        return max(rs, key=key)

    rows = []
    for t in rep["totals"]:
        if t["comp"] == "sum N" or t["hyp"] == "TB1":
            rows.append([f"Suma {t['comp'][4:]} {t['hyp']}", "kN", t["model"], t["cype"], t["diff_%"]])
    for fam, lab in (("ELU", "ELU"), ("CIM", "Cimentación"), ("ELS", "ELS")):
        for what in ("N max", "N min", "|My| max"):
            r = worst(fam, what)
            rows.append([f"Pilote {what} {lab} ({r['pile']}, comb {r['model_comb']})",
                         "kN·m" if "M" in what else "kN", r["model"], r["cype"], r["diff_%"]])
    beams = rep.get("beam_envelopes", [])
    for what in ("M- cara pilote mar", "M+ vano", "V cara mar", "V cara tierra"):
        rs = [b for b in beams if b["what"] == what and b["diff_%"] is not None]
        if rs:
            d = [b["diff_%"] for b in rs]
            rows.append([f"Vigas transversales {what} (rango de diferencia %)", "%",
                         min(d), max(d), None])
    end = write_table(ws, 3, ["Magnitud", "Unidad", "SAP2000", "CYPE", "Dif. %"], rows, pct_col=5)
    ws.column_dimensions["A"].width = 58
    ws.cell(end + 2, 1, "Verde ≤ 5 %, amarillo 5–10 %, rojo > 10 %.  Filas 'rango': columnas SAP/CYPE = "
                        "diferencia mínima / máxima en %.")

    # ---- 2. pile envelopes -------------------------------------------------------------------
    for fam, title in (("ELU", "Pilotes_ELU"), ("CIM", "Pilotes_Cimentacion"), ("ELS", "Pilotes_ELS")):
        ws = wb.create_sheet(title)
        piles = sorted({r["pile"] for r in env}, key=pile_key)
        rows = []
        for p in piles:
            g = {r["what"]: r for r in env if r["family"] == fam and r["pile"] == p}
            rows.append([p, g["N max"]["model"], g["N max"]["cype"], g["N max"]["diff_%"],
                         g["N min"]["model"], g["N min"]["cype"], g["N min"]["diff_%"],
                         abs(g["|My| max"]["model"]), abs(g["|My| max"]["cype"]), g["|My| max"]["diff_%"],
                         abs(g["|Mx| max"]["model"]), abs(g["|Mx| max"]["cype"]), g["|Mx| max"]["diff_%"]])
        hdr = ["Pilote", "N max SAP (kN)", "N max CYPE", "Dif %", "N min SAP (kN)", "N min CYPE", "Dif %",
               "|My| base SAP (kN·m)", "|My| CYPE", "Dif %", "|Mx| base SAP", "|Mx| CYPE", "Dif %"]
        last = write_table(ws, 1, hdr, rows)
        for col in (4, 7, 10, 13):
            for i in range(2, last + 1):
                v = ws.cell(i, col).value
                if isinstance(v, (int, float)):
                    ws.cell(i, col).fill = GOOD if abs(v) <= 5 else WARN if abs(v) <= 10 else BAD
        cats = Reference(ws, min_col=1, min_row=2, max_row=last)
        bar(ws, f"N máx pilotes — {fam}", "kN", cats, [2, 3], 2, last, "O2")
        bar(ws, f"N mín (tracción) — {fam}", "kN", cats, [5, 6], 2, last, "O19")
        bar(ws, f"|My| base pilotes — {fam}", "kN·m", cats, [8, 9], 2, last, "O36")
        bar(ws, f"|Mx| base pilotes — {fam}", "kN·m", cats, [11, 12], 2, last, "O53")

    # ---- 3. beams ---------------------------------------------------------------------------
    if beams:
        ws = wb.create_sheet("Vigas_ELU")
        porticos = []
        for b in beams:
            if b["portico"] not in porticos:
                porticos.append(b["portico"])
        rows = []
        for p in porticos:
            g = {b["what"]: b for b in beams if b["portico"] == p}
            ref = lambda b: b["cype_drawing"] if b["cype_drawing"] is not None else b["cype_listing"]  # noqa: E731
            rows.append([p.replace("Pórtico", "P."), g["M- cara pilote mar"]["section"],
                         abs(g["M- cara pilote mar"]["model"]), abs(ref(g["M- cara pilote mar"])),
                         g["M- cara pilote mar"]["diff_%"],
                         g["M+ vano"]["model"], ref(g["M+ vano"]), g["M+ vano"]["diff_%"],
                         g["V cara mar"]["model"], g["V cara mar"]["cype_listing"], g["V cara mar"]["diff_%"],
                         abs(g["V cara tierra"]["model"]), abs(g["V cara tierra"]["cype_listing"]),
                         g["V cara tierra"]["diff_%"]])
        hdr = ["Pórtico", "Sección", "|M-| cara mar SAP", "|M-| CYPE", "Dif %", "M+ vano SAP", "M+ CYPE",
               "Dif %", "V cara mar SAP", "V CYPE", "Dif %", "|V| cara tierra SAP", "|V| CYPE", "Dif %"]
        last = write_table(ws, 1, hdr, rows)
        for col in (5, 8, 11, 14):
            for i in range(2, last + 1):
                v = ws.cell(i, col).value
                if isinstance(v, (int, float)):
                    ws.cell(i, col).fill = GOOD if abs(v) <= 5 else WARN if abs(v) <= 10 else BAD
        cats = Reference(ws, min_col=1, min_row=2, max_row=last)
        bar(ws, "Momento negativo en cara de pilote (mar) — ELU", "kN·m", cats, [3, 4], 2, last, "P2")
        bar(ws, "Momento positivo en vano — ELU", "kN·m", cats, [6, 7], 2, last, "P19")
        bar(ws, "Cortante en cara de pilote (mar) — ELU", "kN", cats, [9, 10], 2, last, "P36")
        bar(ws, "Cortante en cara de pilote (tierra) — ELU", "kN", cats, [12, 13], 2, last, "P53")
        ws.cell(last + 2, 1, "CYPE M-: pico del diagrama (imágenes del listado); M+ y V: listado §2.")

    # ---- 4. per-hypothesis tables + scatter ------------------------------------------------
    for key, title in (("pile_hypotheses", "Arranques_hipotesis"),
                       ("pile_head_hypotheses", "Cabeza_hipotesis")):
        rs = rep.get(key)
        if not rs:
            continue
        ws = wb.create_sheet(title)
        rows = [[r["pile"], r["hyp"], r["comp"], r["model"], r["cype"], r["diff"]]
                for r in sorted(rs, key=lambda r: (r["comp"], r["hyp"], pile_key(r["pile"])))]
        last = write_table(ws, 1, ["Pilote", "Hipótesis", "Esfuerzo", "SAP2000", "CYPE", "Dif."], rows)
        # scatter SAP vs CYPE for N and My, with the 1:1 line
        for comp, anchor in (("N", "H2"), ("My", "H22")):
            idx = [i for i, r in enumerate(rows, 2) if r[2] == comp]
            if not idx:
                continue
            ch = ScatterChart()
            ch.title = f"{comp}: SAP2000 vs CYPE (todas las hipótesis y pilotes)"
            ch.x_axis.title = "CYPE"
            ch.y_axis.title = "SAP2000"
            ch.style = 13
            ch.height, ch.width = 9, 14
            s = Series(Reference(ws, min_col=4, min_row=idx[0], max_row=idx[-1]),
                       Reference(ws, min_col=5, min_row=idx[0], max_row=idx[-1]), title="Pilotes")
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            ch.series.append(s)
            lo = min(min(rows[i - 2][3], rows[i - 2][4]) for i in idx)
            hi = max(max(rows[i - 2][3], rows[i - 2][4]) for i in idx)
            c0 = 9 + (0 if comp == "N" else 3)
            ws.cell(1, c0 + 10, f"1:1 {comp}")
            ws.cell(2, c0 + 10, lo), ws.cell(3, c0 + 10, hi)
            ref11 = Reference(ws, min_col=c0 + 10, min_row=2, max_row=3)
            s2 = Series(ref11, ref11, title="1:1")
            ch.series.append(s2)
            ws.add_chart(ch, anchor)

    # ---- 5. totals ---------------------------------------------------------------------------
    ws = wb.create_sheet("Totales")
    rows = [[t["hyp"], t["comp"], t["model"], t["cype"], t["diff_%"]] for t in rep["totals"]]
    last = write_table(ws, 1, ["Hipótesis", "Suma", "SAP2000", "CYPE §3.7", "Dif %"], rows, pct_col=5)
    nrows = [i for i in range(2, last + 1) if ws.cell(i, 2).value == "sum N"]
    ws2 = ws
    for k, i in enumerate(nrows):
        ws2.cell(20 + k, 1, ws2.cell(i, 1).value)
        ws2.cell(20 + k, 2, ws2.cell(i, 3).value)
        ws2.cell(20 + k, 3, ws2.cell(i, 4).value)
    ws2.cell(19, 1, "Hipótesis"), ws2.cell(19, 2, "SAP2000"), ws2.cell(19, 3, "CYPE")
    bar(ws2, "Suma de reacciones verticales por hipótesis", "kN",
        Reference(ws2, min_col=1, min_row=20, max_row=19 + len(nrows)), [2, 3], 20, 19 + len(nrows), "G2")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print("written", args.out)


if __name__ == "__main__":
    main()
