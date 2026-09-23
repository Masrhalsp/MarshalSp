"""Read SAP2000's concrete-frame-design results ("Eurocode 2-2004") and compare them with the
Código Estructural checks of diseno/python/codigo (pilotes.json / vigas.json, when they exist).

    python3 diseno/sap/leer_diseno_sap.py <export.xlsx> [<export2.xlsx> ...]
    python3 diseno/sap/leer_diseno_sap.py diseno/sap/output/diseno_SAP2000.json     (OAPI script output)

Inputs:
* GUI export (Display > Show Tables > Design Data > Concrete Frame Design > "Concrete Design 1 -
  Column Summary Data - Eurocode 2-2004" and "Concrete Design 2 - Beam Summary Data - Eurocode
  2-2004" > File > Export to Excel): one sheet per table, row 1 "TABLE:  <name>", row 2 field keys,
  row 3 units, data from row 4 (spec §5, verified on the user's v27.1 exports).  Field names are
  matched with tolerant patterns (the v27 EC2 torsion/column keys are UNVERIFIED) and every area is
  converted with the units row (m², cm², mm², in²; per m / mm / in / ft), so any display units work.
* or the JSON of sap_oapi_diseno.py (same canonical rows).

Output: diseno/sap/output/comparacion_diseno_SAP.json / .md with
* piles: SAP max P-M-M ratio (and at head / foot) vs our eta N-M (1st/2nd order), CYPE 'Aprov.' and
  our forces run through SAP's own column procedure (:func:`sap_pile_estimates`: without a P-Delta
  analysis SAP's M0e rule leaves the sway second-order moment of the piles unchecked, so SAP's ratio
  follows our first-order eta, not the second-order one);
* beams, bending: SAP As top/bottom at the design sections of vigas.py (pile faces, span, cantilever
  faces) vs our required As (ULS, and with As,min) and the provided As;
* beams, shear: SAP Asw/s at the faces vs our required Asw/s (V at d, cot theta = 1 for basis CYPE);
* SAP-only summary tables (always).
The basis (CYPE: ELU combos / ROM: ELR combos) is detected from SAP's combination names.
"""

from __future__ import annotations

import argparse
import functools
import json
import math
import re
import statistics
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
for _p in (ROOT / "diseno" / "python" / "codigo", ROOT / "diseno" / "python", HERE):       # HERE first on sys.path
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import datos_diseno_sap as D  # noqa: E402

OUT_DIR = HERE / "output"
DISENO_OUT = ROOT / "diseno" / "python" / "output"
REF_JSON = ROOT / "diseno" / "ref" / "cype_design_reference.json"

# canonical field -> regex over SAP field keys (case-insensitive); spec §5 + v20.1 real beam fields [G8]
ALIASES_COMMON = {"frame": r"^Frame$", "loc": r"^(Location|Station|Loc)$", "design_sect": r"^DesignSect(ion)?$",
                  "status": r"^Status$", "error": r"^(ErrMsg|ErrorSummary|Errors?)$",
                  "warning": r"^(WarnMsg|WarningSummary|Warnings?)$"}
ALIASES_COLUMN = {"option": r"^(DesignOpt|MyOption|DesignOption)$", "pmm_combo": r"^PMMCombo$",
                  "pmm_area": r"^PMMArea$", "pmm_ratio": r"^PMMRatio$",
                  "vmaj_combo": r"^(VMajCombo|VmajorCombo|VMajorCombo)$",
                  "av_major": r"^(VMajRebar|VMajArea|AVmajor|VmajorArea|VMajorRebar)$",
                  "vmin_combo": r"^(VMinCombo|VminorCombo|VMinorCombo)$",
                  "av_minor": r"^(VMinRebar|VMinArea|AVminor|VminorArea|VMinorRebar)$"}
ALIASES_BEAM = {"top_combo": r"^F?TopCombo$", "as_top": r"^F?Top(Area|Rebar)$", "bot_combo": r"^F?BotCombo$",
                "as_bot": r"^F?Bot(Area|Rebar)$", "v_combo": r"^(VCombo|VmajorCombo|VMajCombo)$",
                "asw_s": r"^(VRebar|VArea|VmajorArea|VMajRebar|VMajorRebar)$",
                "tl_combo": r"^(TLngCombo|TLCombo)$", "asl_t": r"^(TLng(Area|Rebar)|TLArea)$",
                "tt_combo": r"^(TTrnCombo|TTCombo)$", "ast_s": r"^(TTrn(Area|Rebar)|TTArea)$"}
AREA_FIELDS = {"pmm_area", "as_top", "as_bot", "asl_t"}
PER_LENGTH_FIELDS = {"av_major", "av_minor", "asw_s", "ast_s"}
CANON_UNITS = {"pmm_area": "pmm_area_cm2", "as_top": "as_top_cm2", "as_bot": "as_bot_cm2", "asl_t": "asl_t_cm2",
               "av_major": "av_major_cm2_m", "av_minor": "av_minor_cm2_m", "asw_s": "asw_s_cm2_m",
               "ast_s": "ast_s_cm2_m"}
LENGTH = {"m": 1.0, "cm": 0.01, "mm": 0.001, "in": 0.0254, "ft": 0.3048}


# ============================================================================================
# units
# ============================================================================================
def _clean(u) -> str:
    return str(u or "").strip().lower().replace("²", "2").replace("^", "").replace(" ", "")


def length_to_m(unit) -> float:
    u = _clean(unit)
    return LENGTH.get(u, 1.0) if u else 1.0


def area_to_cm2(unit) -> float:
    u = _clean(unit)
    if not u:
        return 1e4                                     # kN-m model default: m2
    m = re.fullmatch(r"([a-z]+)2", u)
    if not m or m.group(1) not in LENGTH:
        raise ValueError(f"unknown area unit {unit!r}")
    return (LENGTH[m.group(1)] * 100.0) ** 2


def per_length_to_cm2_m(unit) -> float:
    u = _clean(unit)
    if not u:
        return 1e4                                     # m2/m
    a, _, b = u.partition("/")
    if not b or b not in LENGTH:
        raise ValueError(f"unknown area-per-length unit {unit!r}")
    return area_to_cm2(a) / LENGTH[b]


# ============================================================================================
# readers
# ============================================================================================
def read_sap_xlsx(path: Path) -> dict[str, dict]:
    """{table name: {"fields", "units", "rows"}} of a SAP2000 Excel export (title in A1)."""
    from openpyxl import load_workbook
    out = {}
    for ws in load_workbook(path, read_only=True, data_only=True).worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            title = str(next(it)[0] or "")
            fields = [str(f) if f is not None else None for f in next(it)]
            units = list(next(it))
        except StopIteration:
            continue
        name = title.split("TABLE:", 1)[1].strip() if "TABLE:" in title else ws.title
        rows = [dict(zip(fields, r)) for r in it if r and any(v is not None for v in r)]
        out[name] = {"fields": fields, "units": dict(zip(fields, units)), "rows": rows}
    return out


def normalise(tab: dict, aliases: dict) -> list[dict]:
    """Canonical rows (areas cm², per-length cm²/m, loc m) of one SAP table."""
    cols, factor = {}, {}
    for canon, rx in aliases.items():
        hit = [f for f in tab["fields"] if f and re.match(rx, f, re.I)]
        if not hit:
            continue
        cols[canon] = hit[0]
        unit = tab["units"].get(hit[0])
        if canon in AREA_FIELDS:
            factor[canon] = area_to_cm2(unit)
        elif canon in PER_LENGTH_FIELDS:
            factor[canon] = per_length_to_cm2_m(unit)
        elif canon == "loc":
            factor[canon] = length_to_m(unit)
    out = []
    for r in tab["rows"]:
        row = {}
        for canon, f in cols.items():
            v = r.get(f)
            if canon in factor and v is not None and v != "":
                v = float(v) * factor[canon]
            row[CANON_UNITS.get(canon, canon)] = v
        if row.get("frame"):
            out.append(row)
    return out


def _frame_index(model: dict) -> dict:
    fr = {f["name"]: f for f in model["frames"]}
    return {fd.frame: (fd, fr[fd.frame]) for fd in D.frame_designs(model)}


def from_gui_export(paths: list[Path], model: dict | None = None) -> tuple[list[dict], list[dict], dict]:
    """Pile and beam rows (canonical, as sap_oapi_diseno.read_results) from GUI Excel exports."""
    model = model or D.base_model()
    idx = _frame_index(model)
    joints = model["joints"]
    col_t, beam_t, names = [], [], []
    for p in paths:
        for name, tab in read_sap_xlsx(Path(p)).items():
            names.append(name)
            if re.search(r"column summary", name, re.I):
                col_t.append(tab)
            elif re.search(r"beam summary", name, re.I):
                beam_t.append(tab)
    if not col_t and not beam_t:
        raise ValueError(f"no 'Column/Beam Summary Data' table in {paths}: found {names}")
    piles, beams = [], []
    for tab in col_t:
        for r in normalise(tab, {**ALIASES_COMMON, **ALIASES_COLUMN}):
            fd, f = idx.get(r["frame"], (None, None))
            if fd is None:
                continue
            opt = str(r.get("option") or "").lower()
            piles.append({"frame": r["frame"], "member": fd.member,
                          "option": "check" if opt.startswith("check") or opt == "1" else (opt or "check"),
                          "loc": _rnd(r.get("loc")), "z": _rnd(joints[f["i"]][2] + (r.get("loc") or 0.0)),
                          "pmm_combo": r.get("pmm_combo"), "pmm_area_cm2": _rnd(r.get("pmm_area_cm2"), 3),
                          "pmm_ratio": _rnd(r.get("pmm_ratio"), 5), "vmaj_combo": r.get("vmaj_combo"),
                          "av_major_cm2_m": _rnd(r.get("av_major_cm2_m"), 3), "vmin_combo": r.get("vmin_combo"),
                          "av_minor_cm2_m": _rnd(r.get("av_minor_cm2_m"), 3),
                          "error": r.get("error") or "", "warning": r.get("warning") or "",
                          "status": r.get("status") or ""})
    for tab in beam_t:
        for r in normalise(tab, {**ALIASES_COMMON, **ALIASES_BEAM}):
            fd, f = idx.get(r["frame"], (None, None))
            if fd is None:
                continue
            xi, yi, _ = joints[f["i"]]
            origin = yi if fd.kind == "beam_t" else xi
            beams.append({"frame": r["frame"], "member": fd.member, "kind": fd.kind, "loc": _rnd(r.get("loc")),
                          "pos": _rnd(origin + (r.get("loc") or 0.0)), "top_combo": r.get("top_combo"),
                          "as_top_cm2": _rnd(r.get("as_top_cm2"), 3), "bot_combo": r.get("bot_combo"),
                          "as_bot_cm2": _rnd(r.get("as_bot_cm2"), 3), "v_combo": r.get("v_combo"),
                          "asw_s_cm2_m": _rnd(r.get("asw_s_cm2_m"), 3), "tl_combo": r.get("tl_combo"),
                          "asl_t_cm2": _rnd(r.get("asl_t_cm2"), 3), "tt_combo": r.get("tt_combo"),
                          "ast_s_cm2_m": _rnd(r.get("ast_s_cm2_m"), 3),
                          "error": r.get("error") or "", "warning": r.get("warning") or "",
                          "status": r.get("status") or ""})
    return piles, beams, {"source": [str(p) for p in paths], "tables": names}


def from_oapi_json(path: Path) -> tuple[list[dict], list[dict], dict]:
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    return d["piles"], d["beams"], {"source": [str(path)], "oapi_meta": {k: d["meta"].get(k) for k in
                                                                        ("base", "kphi_piles", "tan_theta",
                                                                         "table_overwrites", "to_do")},
                                    "analysis_identity": d.get("analysis_identity")}


def _rnd(v, nd: int = 4):
    return None if v is None or v == "" else round(float(v), nd)


def detect_base(piles: list[dict], beams: list[dict]) -> str:
    combos = {r.get(k) for r in piles for k in ("pmm_combo",)} | {r.get(k) for r in beams
                                                                  for k in ("top_combo", "bot_combo", "v_combo")}
    combos = {str(c) for c in combos if c}
    if any(c.startswith("ELR") for c in combos):
        return "ROM"
    if any(c.startswith("ELU") for c in combos):
        return "CYPE"
    return "?"


# ============================================================================================
# our forces through SAP's EC2 column procedure (what SAP should print for the piles)
# ============================================================================================
def _kphi_value(kphi) -> float:
    kp = D.KPHI_OPTIONS[kphi] if isinstance(kphi, str) else float(kphi)
    return kp or 1.0                                   # 0 = program default = 1.0


def _add(m: float, extra: float) -> float:
    """|m| increased by extra >= 0, keeping the sign of m."""
    return math.copysign(abs(m) + extra, m if m else 1.0)


@functools.lru_cache(maxsize=4)
def sap_pile_estimates(base: str, kphi=D.KPHI_DEFAULT) -> tuple:
    """Pile P-M-M ratio that SAP2000 should print with the settings of this folder, from OUR forces
    (SAP27 export) and SAP's EC2 column procedure (CFD-EC-2-2004 §3.4.2):

    * first order at each end station + Mi = N·max(ei, emin) (ei = θ0·αh·l0/2 = 12.9 mm, emin = 20 mm),
      in one direction at a time;
    * nominal curvature: e2 = Kr·Kφ·εyd/(0.45 d)·l0²/c with c = 8, d = 327.5 mm (outer bar row) and
      the Kφ overwrite, M2 = N·e2 in both directions, λ > λlim assumed (SAP's A, B, C are not
      documented; λ = 58);
    * variant "M0e" = the manual literally: MEd = M0e + Mi + M2 with M0e = 0.6 M02 + 0.4 M01 >= 0.4 M02
      (end moments).  The piles bend in double curvature (sway), so M0e = 0.4 M02 and the first-order
      end check governs: the sway second-order moment at the pile ends is NOT covered, because SAP
      expects it from a P-Delta analysis (manual 3.4.2.2) and the model is linear first order;
    * variant "estación" = M2 added to the end-station moment (CYPE / Código Estructural practice for
      a sway member); SAP should show this only if its details print MEd = M02 + Mi + M2.

    Capacity: fibre section 12Ø25 Mode.CODIGO (``eta_fast``), a proxy for SAP's P-M-M surface
    (rectangular block, 48 x 21 points): expect about ±3 %.  Returns one dict per pile."""
    import armado_pilote as AP                          # noqa: PLC0415 - heavy imports only here
    import esfuerzos as E                               # noqa: PLC0415
    from seccion import Mode                            # noqa: PLC0415

    lay = AP.PROVIDED["fuste"]
    sec = lay.section(Mode.CODIGO)
    h, l0 = lay.b, D.L0_PILE * 1000.0
    d = h / 2 + lay.c                                   # 327.5 mm
    alpha_h = min(1.0, max(2.0 / 3.0, 2.0 / math.sqrt(D.L0_PILE)))
    ei = max(D.EC2_PREFS[7] * alpha_h * l0 / 2.0, h / 30.0, 20.0)
    e2 = _kphi_value(kphi) * (lay.steel.fyd / lay.steel.Es) / (0.45 * d) * l0 ** 2 / 8.0

    def eta(N: float, M: dict) -> float:
        return float(sec.eta_fast((N * 1e3, M["Mx"] * 1e6, M["My"] * 1e6)))

    def with_mi(N: float, M: dict) -> list[dict]:
        out = []
        for ax in ("Mx", "My"):
            m = dict(M)
            m[ax] = _add(m[ax], abs(N) * ei / 1000.0)
            out.append(m)
        return out

    stations = E.load_stations()["piles"]
    rows = []
    for pile, sts in stations.items():
        flex = [s for s in sts if not s.rigid]
        ends = {"cabeza": max(flex, key=lambda s: s.pos), "pie": min(flex, key=lambda s: s.pos)}
        best = {"M0e": (-1.0, "", ""), "estacion": (-1.0, "", "")}

        def keep(kind: str, value: float, combo: str, what: str) -> None:
            if value > best[kind][0]:
                best[kind] = (value, combo, what)

        for combo, fac in E.uls(base).items():
            f = {k: E.combine(s.values, fac, ("N", "Mx", "My")) for k, s in ends.items()}
            for end, r in f.items():
                M = {"Mx": r["Mx"], "My": r["My"]}
                for m in with_mi(r["N"], M):
                    v = eta(r["N"], m)
                    keep("M0e", v, combo, f"1er orden + Mi, {end}")
                    keep("estacion", v, combo, f"1er orden + Mi, {end}")
                if r["N"] <= 0:
                    continue
                m2 = r["N"] * e2 / 1000.0
                for m in with_mi(r["N"], {k: _add(M[k], m2) for k in M}):
                    keep("estacion", eta(r["N"], m), combo, f"M + Mi + M2, {end}")
                m0e = {}
                for k in ("Mx", "My"):
                    a, b = f["cabeza"][k], f["pie"][k]
                    m02, m01 = (a, b) if abs(a) >= abs(b) else (b, a)
                    val = abs(0.6 * m02 + 0.4 * m01) if (0.6 * m02 + 0.4 * m01) * m02 > 0 else 0.0
                    m0e[k] = math.copysign(max(val, 0.4 * abs(m02)), m02 if m02 else 1.0)
                for m in with_mi(r["N"], {k: _add(m0e[k], m2) for k in m0e}):
                    keep("M0e", eta(r["N"], m), combo, f"M0e + Mi + M2 (N {end})")
        rows.append({"pile": pile,
                     "sap_est_M0e": round(best["M0e"][0], 3), "combo_M0e": best["M0e"][1],
                     "gov_M0e": best["M0e"][2],
                     "sap_est_estacion": round(best["estacion"][0], 3), "combo_estacion": best["estacion"][1],
                     "gov_estacion": best["estacion"][2]})
    meta = {"ei_mm": round(ei, 2), "e2_mm": round(e2, 2), "d_mm": d, "Kphi": _kphi_value(kphi), "c": 8}
    return tuple(rows), meta


def _estimates(base: str, kphi=D.KPHI_DEFAULT) -> tuple[dict, dict]:
    """``sap_pile_estimates`` keyed by pile; empty (with the reason) if the forces are unavailable."""
    if base not in D.BASES:
        return {}, {"error": f"base {base!r}"}
    try:
        rows, meta = sap_pile_estimates(base, kphi)
    except Exception as exc:                            # noqa: BLE001 - optional information
        return {}, {"error": f"{type(exc).__name__}: {exc}"}
    return {r["pile"]: r for r in rows}, meta


# ============================================================================================
# comparison
# ============================================================================================
def _pct(a, b):
    return None if a is None or b in (None, 0) else round((a / b - 1) * 100, 2)


def _load(path: Path):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def sap_pile_summary(piles: list[dict]) -> list[dict]:
    out = []
    for member in dict.fromkeys(r["member"] for r in piles):
        rows = [r for r in piles if r["member"] == member and r.get("pmm_ratio") is not None]
        if not rows:
            continue
        top = max(rows, key=lambda r: r["pmm_ratio"])
        head, foot = max(rows, key=lambda r: r["z"]), min(rows, key=lambda r: r["z"])
        out.append({"pile": member, "pmm_ratio_max": top["pmm_ratio"], "z_max": top["z"], "combo": top["pmm_combo"],
                    "pmm_ratio_head": head["pmm_ratio"], "pmm_ratio_foot": foot["pmm_ratio"],
                    "av_major_max_cm2_m": max((r["av_major_cm2_m"] or 0) for r in rows),
                    "av_minor_max_cm2_m": max((r["av_minor_cm2_m"] or 0) for r in rows),
                    "ok": top["pmm_ratio"] <= 1.0,
                    "messages": "; ".join(sorted({r.get("error") or r.get("warning") or "" for r in rows} - {""}))})
    return out


def _sap_face(rows: list[dict], member: str, pos: float, key: str, tol: float = 0.006) -> dict | None:
    cand = [r for r in rows if r["member"] == member and abs(r["pos"] - pos) <= tol and r.get(key) is not None]
    return max(cand, key=lambda r: r[key]) if cand else None


FACE_POS = {"cara pilote mar - voladizo": -0.20, "cara pilote mar - vano": 0.20,
            "cara pilote tierra - vano": 3.25, "cara pilote tierra - vuelo": 3.65}


def face_position(section: str, pos: float) -> float:
    """Support face of a vigas.py shear section (the VRd,s row itself sits at d from it)."""
    if section in FACE_POS:
        return FACE_POS[section]
    m = re.search(r"X=(-?[\d.]+)", section)
    return float(m.group(1)) if m else pos


def _sap_nearest(rows: list[dict], member: str, pos: float, key: str, tol: float = 0.13) -> dict | None:
    cand = [r for r in rows if r["member"] == member and abs(r["pos"] - pos) <= tol and r.get(key) is not None]
    if not cand:
        return None
    dmin = min(abs(r["pos"] - pos) for r in cand)
    return max((r for r in cand if abs(r["pos"] - pos) <= dmin + 1e-6), key=lambda r: r[key])


def _span_of(member: str, section: str, pos: float) -> tuple[float, float]:
    """Region where SAP's maximum is taken for a 'máx. positivo' row."""
    if member.startswith("VT"):
        return 0.20, 3.25
    m = re.search(r"tramo (\d+)", section)
    k = int(m.group(1)) if m else 1 + int(pos // D.tm.AXIS_SPACING)
    return D.tm.AXES_X[k - 1], D.tm.AXES_X[k]


def sap_beam_value(beams: list[dict], member: str, section: str, pos: float, key: str) -> dict | None:
    if "máx." in section:
        a, b = _span_of(member, section, pos)
        cand = [r for r in beams if r["member"] == member and a - 1e-6 <= r["pos"] <= b + 1e-6 and r.get(key) is not None]
        return max(cand, key=lambda r: r[key]) if cand else None
    return _sap_face(beams, member, pos, key)


def _asw_cot1(r: dict) -> float | None:
    """Our required Asw/s converted to cot θ = 1 (SAP runs with the TanTheta = 1 overwrite): the
    vigas.py value is V/(z·fywd·cot θ) without the minimum, so it scales with cot θ."""
    a = r.get("Asw_s_req_cm2_m")
    return None if a is None else round(a * (r.get("cot_theta") or 1.0), 3)


def _aprov(ref: dict | None) -> dict:
    out: dict = {}
    if ref:
        for r in ref["piles"]["appendix_1_section_3"]["armado_pilares_3_2"]["rows"]:
            out.setdefault(r["pilar"], {})[r["planta"]] = r["aprov_pct"]
    return out


def compare(piles: list[dict], beams: list[dict], base: str, pil: dict | None, vig: dict | None,
            ref: dict | None, kphi=D.KPHI_DEFAULT) -> dict:
    tables: dict[str, list[dict]] = {"sap_pilotes": sap_pile_summary(piles)}
    sap_p = {r["pile"]: r for r in tables["sap_pilotes"]}
    # ---- piles -----------------------------------------------------------------------------------
    aprov = _aprov(ref)
    est, _ = _estimates(base, kphi) if sap_p else ({}, {})
    cype_cmp = []
    for pile, s in sap_p.items():
        ap = aprov.get(pile, {}).get("Forjado 1")
        if ap:
            cype_cmp.append({"quantity": f"{pile} P-M-M ratio (SAP) vs Aprov. fuste (CYPE)",
                             "ours": s["pmm_ratio_max"], "cype": ap / 100, "diff_pct": _pct(s["pmm_ratio_max"], ap / 100),
                             "src": "Anejo 10 Apéndice 1 §3.2 (cype_design_reference.json)"})
    if pil or est:
        var = {r["pile"]: r for r in pil["tables"].get("variants", [])} if pil else {}
        cases = [("CYPE", "codigo"), ("CYPE", "cype")] if base != "ROM" else [("ROM", "codigo")]
        rows = []
        for pile, s in sap_p.items():
            v = var.get(pile, {})
            row = {"pile": pile, "sap_ratio_max": s["pmm_ratio_max"], "sap_z": s["z_max"], "sap_combo": s["combo"],
                   "sap_ratio_head": s["pmm_ratio_head"], "sap_ratio_foot": s["pmm_ratio_foot"]}
            e = est.get(pile, {})
            row["sap_est_M0e"], row["sap_est_estacion"] = e.get("sap_est_M0e"), e.get("sap_est_estacion")
            row["diff_pct_vs_est_M0e"] = _pct(s["pmm_ratio_max"], e.get("sap_est_M0e"))
            for b, mode in cases if pil else ():
                eta = max(x for x in (v.get(f"nm1[{b},{mode}]"), v.get(f"nm2[{b},{mode}]")) if x is not None) \
                    if v.get(f"nm1[{b},{mode}]") is not None else None
                row[f"ours_nm1[{b},{mode}]"] = v.get(f"nm1[{b},{mode}]")
                row[f"ours_eta[{b},{mode}]"] = eta
                row[f"diff_pct[{b},{mode}]"] = _pct(s["pmm_ratio_max"], eta)
            ap = aprov.get(pile, {})
            row["cype_aprov_fuste"] = ap.get("Forjado 1") / 100 if ap.get("Forjado 1") else None
            row["cype_aprov_arranque"] = ap.get("Cimentación") / 100 if ap.get("Cimentación") else None
            row["sap_av_major_cm2_m"], row["sap_av_minor_cm2_m"] = s["av_major_max_cm2_m"], s["av_minor_max_cm2_m"]
            row["prov_asw_s_cm2_m"] = D.provided_summary()[D.PILE_SECTION]["Asw_s_cm2_m"]
            rows.append(row)
        tables["pilotes"] = rows
    # ---- beams -----------------------------------------------------------------------------------
    sap_b = []
    for member in dict.fromkeys(r["member"] for r in beams):
        rows = [r for r in beams if r["member"] == member]
        t = max(rows, key=lambda r: r["as_top_cm2"] or 0)
        b = max(rows, key=lambda r: r["as_bot_cm2"] or 0)
        v = max(rows, key=lambda r: r["asw_s_cm2_m"] or 0)
        sap_b.append({"member": member, "as_top_max_cm2": t["as_top_cm2"], "pos_top": t["pos"],
                      "as_bot_max_cm2": b["as_bot_cm2"], "pos_bot": b["pos"], "asw_s_max_cm2_m": v["asw_s_cm2_m"],
                      "pos_v": v["pos"], "asl_t_max_cm2": max((r["asl_t_cm2"] or 0) for r in rows),
                      "ast_s_max_cm2_m": max((r["ast_s_cm2_m"] or 0) for r in rows),
                      "messages": "; ".join(sorted({r.get("error") or r.get("warning") or "" for r in rows} - {""}))})
    tables["sap_vigas"] = sap_b
    if vig:
        flex_case = "CE-ROM" if base == "ROM" else "CE-CYPE"
        extra_case = None if base == "ROM" else "CYPE"
        req = vig["tables"].get("as_required", [])
        by_key = {(r["case"], r["member"], r["section"], r["face"]): r for r in req}
        rows = []
        for r in req:
            if r["case"] != flex_case:
                continue
            key = "as_top_cm2" if r["face"] == "top" else "as_bot_cm2"
            s = sap_beam_value(beams, r["member"], r["section"], r["pos_m"], key)
            ex = by_key.get((extra_case, r["member"], r["section"], r["face"])) if extra_case else None
            rows.append({"member": r["member"], "section": r["section"], "pos": r["pos_m"], "face": r["face"],
                         "in_verdict": r.get("in_verdict", True), "ours_combo": r["combo"],
                         "ours_As_uls_cm2": r["As_uls_cm2"], "ours_As_uls_cype_mode_cm2": ex["As_uls_cm2"] if ex else None,
                         "ours_As_req_cm2": r["As_req_cm2"], "As_prov_cm2": r["As_prov_cm2"],
                         "sap_As_cm2": s[key] if s else None, "sap_pos": s["pos"] if s else None,
                         "sap_combo": s["top_combo" if r["face"] == "top" else "bot_combo"] if s else None,
                         "diff_pct_vs_uls": _pct(s[key], r["As_uls_cm2"]) if s else None,
                         "sap_over_prov": round(s[key] / r["As_prov_cm2"], 3) if s and r["As_prov_cm2"] else None,
                         "note": "" if s else "SAP: sin estación de diseño (tramo rígido 'No Design' en el eje del pilote)"})
        tables["vigas_flexion"] = rows
        shear_case = "CE-ROM" if base == "ROM" else "CYPE"
        rows = []
        for r in vig["tables"].get("shear_torsion", []):
            if r["case"] != shear_case or not r["check"].startswith("VEd(d) <= VRd,s"):
                continue
            face = face_position(r["section"], r["pos_m"])
            sf = _sap_face(beams, r["member"], face, "asw_s_cm2_m")
            sd = _sap_nearest(beams, r["member"], r["pos_m"], "asw_s_cm2_m")
            cot1 = _asw_cot1(r)
            rows.append({"member": r["member"], "section": r["section"], "face_pos": face, "pos_d": r["pos_m"],
                         "ours_combo": r["combo"], "ours_cot_theta": r.get("cot_theta"),
                         "ours_Asw_s_req_cm2_m": r.get("Asw_s_req_cm2_m"), "ours_Asw_s_cot1_cm2_m": cot1,
                         "Asw_s_prov_cm2_m": r.get("Asw_s_prov_cm2_m"),
                         "sap_Asw_s_face_cm2_m": sf["asw_s_cm2_m"] if sf else None,
                         "sap_combo_face": sf["v_combo"] if sf else None,
                         "sap_Asw_s_d_cm2_m": sd["asw_s_cm2_m"] if sd else None, "sap_pos_d": sd["pos"] if sd else None,
                         "diff_pct_d": _pct(sd["asw_s_cm2_m"], cot1) if sd else None,
                         "sap_face_over_prov": (round(sf["asw_s_cm2_m"] / r["Asw_s_prov_cm2_m"], 3)
                                                if sf and r.get("Asw_s_prov_cm2_m") else None)})
        tables["vigas_cortante"] = rows
    # ---- summary ---------------------------------------------------------------------------------
    summary = {"base": base, "sap_piles_max_ratio": max((r["pmm_ratio_max"] for r in tables["sap_pilotes"]), default=None),
               "sap_piles_over_1": [r["pile"] for r in tables["sap_pilotes"] if not r["ok"]]}
    if "pilotes" in tables:
        k = "diff_pct[ROM,codigo]" if base == "ROM" else "diff_pct[CYPE,codigo]"
        d = [r[k] for r in tables["pilotes"] if r.get(k) is not None]
        summary["piles_diff_pct_vs_ours_ce"] = {"mean": round(statistics.mean(d), 2), "min": min(d), "max": max(d)} if d else None
        d = [r["diff_pct_vs_est_M0e"] for r in tables["pilotes"] if r.get("diff_pct_vs_est_M0e") is not None]
        summary["piles_diff_pct_vs_sap_procedure_estimate"] = ({"mean": round(statistics.mean(d), 2), "min": min(d),
                                                                "max": max(d)} if d else None)
    if "vigas_flexion" in tables:
        d = [r["diff_pct_vs_uls"] for r in tables["vigas_flexion"]
             if r["diff_pct_vs_uls"] is not None and r["in_verdict"] and (r["ours_As_uls_cm2"] or 0) > 2.0]
        summary["beams_flexure_diff_pct_vs_ours_uls"] = ({"n": len(d), "median": round(statistics.median(d), 2),
                                                          "min": min(d), "max": max(d)} if d else None)
        summary["beams_sap_As_over_provided"] = [f"{r['member']} {r['section']} {r['face']}: {r['sap_over_prov']}"
                                                 for r in tables["vigas_flexion"]
                                                 if r["sap_over_prov"] and r["sap_over_prov"] > 1.0]
    if "vigas_cortante" in tables:
        d = [r["diff_pct_d"] for r in tables["vigas_cortante"] if r["diff_pct_d"] is not None]
        summary["beams_shear_diff_pct"] = ({"n": len(d), "median": round(statistics.median(d), 2), "min": min(d),
                                            "max": max(d)} if d else None)
    return {"tables": tables, "summary": summary, "cype_comparison": cype_cmp,
            "proposal": {"note": "no aplica: SAP2000 es una comprobación independiente; propuestas en "
                                 "diseno/python/output/pilotes.json y vigas.json"}}


NOTES_PILES = [
    "Pilotes: SAP (sin análisis P-Delta, como este modelo) aplica el método de la curvatura nominal con "
    "M0e = 0.6·M02 + 0.4·M01 >= 0.4·M02 (manual CFD-EC-2-2004 §3.4.2.2).  Los pilotes flectan en doble "
    "curvatura (traslacionales), así que M0e = 0.4·M02 y gobierna la comprobación de 1er orden en los "
    "extremos + Mi = N·max(ei, emin): el momento de 2º orden del pilote traslacional en cabeza/pie NO se "
    "comprueba (SAP lo espera de un análisis P-Delta).  Valor esperado en SAP = 'sap_est_M0e' (≈ nm1 propio "
    "+ 1-6 %), no nm2.  Si los detalles de SAP muestran MEd = M02 + Mi + M2, el valor esperado es "
    "'sap_est_estacion' (≈ nm2 con e2 de c = 8 y el Kφ del overwrite).  El pilote se verifica con nm2 "
    "(diseno/python/output/pilotes.md), no con SAP.",
    "SAP e2 = Kr·Kφ·εyd/(0.45 d)·l0²/8 con d = 327.5 mm: Kφ 1.1129 da e2 = 92.1 mm = CYPE; el e2 estricto del "
    "Código (d = h/2 + is = 307 mm) es 98.3 mm (Kφ 'ce_is' = 1.1874).",
]


def expected(base: str, pil: dict | None, vig: dict | None, ref: dict | None, kphi=D.KPHI_DEFAULT) -> dict:
    """Our results at the stations SAP reports (what SAP should roughly show), before running SAP;
    for the piles also our forces through SAP's own column procedure (:func:`sap_pile_estimates`)."""
    tables: dict[str, list[dict]] = {}
    aprov = _aprov(ref)
    est, est_meta = _estimates(base, kphi)
    if pil or est:
        modes = [("CYPE", "codigo"), ("CYPE", "cype")] if base != "ROM" else [("ROM", "codigo")]
        var = {v["pile"]: v for v in pil["tables"].get("variants", [])} if pil else {}
        rows = []
        for pile in dict.fromkeys([*var, *est]):
            v, e = var.get(pile, {}), est.get(pile, {})
            row = {"pile": pile}
            for b, mode in modes:
                row[f"nm1[{b},{mode}]"] = v.get(f"nm1[{b},{mode}]")
                row[f"nm2[{b},{mode}]"] = v.get(f"nm2[{b},{mode}]")
            row.update({k: e.get(k) for k in ("sap_est_M0e", "combo_M0e", "gov_M0e", "sap_est_estacion",
                                             "combo_estacion")})
            ap = aprov.get(pile, {})
            row["cype_aprov_fuste"] = ap.get("Forjado 1") / 100 if ap.get("Forjado 1") else None
            row["cype_aprov_arranque"] = ap.get("Cimentación") / 100 if ap.get("Cimentación") else None
            rows.append(row)
        tables["pilotes"] = rows
    if vig:
        case = "CE-ROM" if base == "ROM" else "CE-CYPE"
        tables["vigas_flexion"] = [
            {"member": r["member"], "section": r["section"], "pos": r["pos_m"], "face": r["face"], "combo": r["combo"],
             "MEd_kNm": r["MEd_kNm"], "As_uls_cm2": r["As_uls_cm2"], "As_min_cm2": r["As_min_cm2"],
             "As_prov_cm2": r["As_prov_cm2"]}
            for r in vig["tables"].get("as_required", []) if r["case"] == case and r.get("in_verdict", True)]
        shear_case = "CE-ROM" if base == "ROM" else "CYPE"
        st = [r for r in vig["tables"].get("shear_torsion", []) if r["case"] == shear_case]
        vmax = {(r["member"], r["section"]): r for r in st if r["check"].startswith("VEd(cara) <= VRd,max")}
        fywd_sap = D.REBAR.fy / 1e3 / 1.15                # SAP: fyk/gamma_s = 434.8 MPa, z = 0.9 d, TanTheta 1

        def sap_face(r: dict) -> dict:
            f = vmax.get((r["member"], r["section"]))
            if not f or not f.get("z_mm"):
                return {"V_face_kN": None, "sap_est_Asw_s_face_cm2_m": None}
            return {"V_face_kN": f["demand"], "combo_face": f["combo"],
                    "sap_est_Asw_s_face_cm2_m": round(f["demand"] * 1e3 / (f["z_mm"] * fywd_sap) * 10.0, 3)}

        tables["vigas_cortante"] = [
            {"member": r["member"], "section": r["section"], "face_pos": face_position(r["section"], r["pos_m"]),
             "pos_d": r["pos_m"], "combo": r["combo"], "cot_theta": r.get("cot_theta"),
             "Asw_s_req_cm2_m": r.get("Asw_s_req_cm2_m"), "Asw_s_req_cot1_cm2_m": _asw_cot1(r),
             **sap_face(r), "Asw_s_prov_cm2_m": r.get("Asw_s_prov_cm2_m")}
            for r in st if r["check"].startswith("VEd(d) <= VRd,s")]
    n_beams = len(tables.get("vigas_flexion", [])) + len(tables.get("vigas_cortante", []))
    return {"meta": {"module": "diseno/sap/leer_diseno_sap.py --esperados", "base": base,
                     "title": f"Valores propios esperados en las estaciones de diseño de SAP (base {base}; "
                              "no son resultados de SAP)",
                     "source": [(DISENO_OUT / "pilotes.json").relative_to(ROOT).as_posix(),
                                (DISENO_OUT / "vigas.json").relative_to(ROOT).as_posix(),
                                "sap2000/resultados_sap/SAP27_Element_Forces_Frames.xlsx (sap_est_*)"],
                     "n_rows": {"piles": len(tables.get("pilotes", [])), "beams": n_beams},
                     "sap_procedure": est_meta,
                     "notes": ["Valores propios (Código Estructural) en las estaciones que SAP reporta; no son "
                               "resultados de SAP.  Vigas: As_uls = armadura de tracción ELU (SAP no añade "
                               "As,min del Código en la misma forma); Asw/s con V a d de la cara; "
                               "'Asw_s_req_cot1' = el mismo valor con cot θ = 1 (SAP con TanTheta = 1; SAP usa "
                               "además V en la cara y fywd = 434.8 MPa).", *NOTES_PILES]},
            "tables": tables, "summary": {"base": base}}


def run(inputs: list[Path], pilotes: Path = DISENO_OUT / "pilotes.json", vigas: Path = DISENO_OUT / "vigas.json",
        model: dict | None = None) -> dict:
    if len(inputs) == 1 and Path(inputs[0]).suffix.lower() == ".json":
        piles, beams, src = from_oapi_json(inputs[0])
    else:
        piles, beams, src = from_gui_export([Path(p) for p in inputs], model)
    base = detect_base(piles, beams)
    pil, vig, ref = _load(Path(pilotes)), _load(Path(vigas)), _load(REF_JSON)
    kphi = (src.get("oapi_meta") or {}).get("kphi_piles")
    kphi = D.KPHI_DEFAULT if kphi is None else kphi       # 0.0 = SAP default (Kphi 1.0), kept
    res = compare(piles, beams, base, pil, vig, ref, kphi)
    res["meta"] = {"module": "diseno/sap/leer_diseno_sap.py", **src, "base": base,
                   "ours": {"pilotes": str(pilotes) if pil else None, "vigas": str(vigas) if vig else None},
                   "n_rows": {"piles": len(piles), "beams": len(beams)},
                   "notes": [
                       "SAP ratio = OL/OC along the load ray (P, M2, M3), like our eta (CYPE 'mismas excentricidades').",
                       "SAP e2 uses c = 8 (CE/CYPE pi^2); Kphi overwrite 1.1129 makes e2 equal; ei/emin rules differ.",
                       *NOTES_PILES,
                       "Shear: 'diff_pct_d' compares SAP with our Asw/s converted to cot theta = 1 (SAP TanTheta = 1).",
                       "Beams: SAP designs the web rectangle, flexure from M3 only, shear with V at the face, "
                       "fywd = fyk/gamma_s = 434.8 MPa, nu = 0.6(1 - fck/250) (ours: V at d, fywd 400, nu1 0.6).",
                       "Our flexure column 'ours_As_uls' is the ULS tension steel (no As,min); 'ours_As_req' includes As,min.",
                       "SAP does not check SLS (crack width 0.1 mm XS3, stresses): see diseno/python/output/*.md."]}
    res = {"meta": res.pop("meta"), **res}
    return res


# ============================================================================================
# report
# ============================================================================================
def _fmt(v) -> str:
    if v is None:
        return "-"
    if isinstance(v, bool):
        return "sí" if v else "NO"
    if isinstance(v, float):
        return f"{v:.3f}" if abs(v) < 10 else f"{v:.2f}"
    return str(v)


def _md_table(rows: list[dict], cols: list[str] | None = None) -> list[str]:
    if not rows:
        return ["(sin filas)", ""]
    cols = cols or list(rows[0])
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    out += ["| " + " | ".join(_fmt(r.get(c)) for c in cols) + " |" for r in rows]
    return out + [""]


def write_report(res: dict, out_dir: Path = OUT_DIR, stem: str = "comparacion_diseno_SAP") -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    js = out_dir / f"{stem}.json"
    js.write_text(json.dumps(res, indent=1, ensure_ascii=False), encoding="utf-8")
    t, s, m = res["tables"], res["summary"], res["meta"]
    L = [f"# {m.get('title', 'Diseño SAP2000 v27.1 (Eurocode 2-2004) vs comprobaciones Código Estructural')}", "",
         f"Fuente: {', '.join(m['source'])}  ·  base de combinaciones: **{m['base']}**  ·  "
         f"filas: {m['n_rows']['piles']} pilotes, {m['n_rows']['beams']} vigas", "",
         "## Resumen", "", "```", json.dumps(s, indent=1, ensure_ascii=False), "```", "", "## Notas", ""]
    L += [f"- {n}" for n in m["notes"]] + [""]
    if m.get("sap_procedure"):
        L += ["Procedimiento SAP usado en `sap_est_*` (nuestras fuerzas): "
              + ", ".join(f"{k} = {v}" for k, v in m["sap_procedure"].items()), ""]
    if res.get("cype_comparison"):
        L += ["## SAP vs CYPE (Anejo 10)", ""] + _md_table(res["cype_comparison"])
    if m.get("analysis_identity"):
        L += ["## Identidad del análisis (OAPI)", "", "```", json.dumps(m["analysis_identity"], indent=1), "```", ""]
    titles = {"pilotes": "Pilotes: ratio P-M-M (SAP) / η N-M propia / Aprov. CYPE",
              "sap_pilotes": "Pilotes: resumen SAP", "vigas_flexion": "Vigas: As a flexión (cm²)",
              "vigas_cortante": "Vigas: Asw/s (cm²/m)", "sap_vigas": "Vigas: máximos SAP"}
    for k in ("pilotes", "sap_pilotes", "vigas_flexion", "vigas_cortante", "sap_vigas"):
        if k in t:
            L += [f"## {titles[k]}", ""] + _md_table(t[k])
    md = out_dir / f"{stem}.md"
    md.write_text("\n".join(L), encoding="utf-8")
    return {"json": js, "md": md}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("inputs", nargs="*", type=Path, help="GUI Excel export(s) or diseno_SAP2000.json")
    ap.add_argument("--pilotes", type=Path, default=DISENO_OUT / "pilotes.json")
    ap.add_argument("--vigas", type=Path, default=DISENO_OUT / "vigas.json")
    ap.add_argument("--esperados", choices=tuple(D.BASES), default=None,
                    help="write only our values at SAP's stations (valores_esperados_SAP_<base>.md), no SAP input")
    ap.add_argument("--kphi", default=D.KPHI_DEFAULT, help="pile Kphi of the SAP run for the sap_est_* columns")
    ap.add_argument("-o", "--out", type=Path, default=OUT_DIR)
    args = ap.parse_args()
    if args.esperados:
        kphi = args.kphi if args.kphi in D.KPHI_OPTIONS else float(args.kphi)
        res = expected(args.esperados, _load(args.pilotes), _load(args.vigas), _load(REF_JSON), kphi)
        paths = write_report(res, args.out, f"valores_esperados_SAP_{args.esperados}")
        for k, p in paths.items():
            print(f"{k}: {p}")
        return
    if not args.inputs:
        ap.error("give the SAP Excel export(s) or diseno_SAP2000.json (or --esperados)")
    res = run(args.inputs, args.pilotes, args.vigas)
    paths = write_report(res, args.out)
    print(json.dumps(res["summary"], indent=1, ensure_ascii=False))
    for k, p in paths.items():
        print(f"{k}: {p}")


if __name__ == "__main__":
    main()
