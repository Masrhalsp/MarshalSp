"""SAP2000 v27.1 concrete frame design ("Eurocode 2-2004") of the Muelle de Trasmallo model
through the CSI OAPI - an independent check of the Código Estructural design.

Windows, SAP2000 v27.1 installed, Python 3.9+ (pip install comtypes openpyxl numpy scipy):

    python diseno\\sap\\sap_oapi_diseno.py                  # start SAP, build the model, design
    python diseno\\sap\\sap_oapi_diseno.py --attach         # use the SAP2000 window already open
    python diseno\\sap\\sap_oapi_diseno.py --from-s2k diseno\\sap\\output\\Muelle_Trasmallo_40m_rect.$2k
    python diseno\\sap\\sap_oapi_diseno.py --base ROM       # design with ELR01-22 (psi0,Qa = 1.0)

Steps (diseno/sap/sap_design_spec.md, section in brackets):

1. model: built with ``sap2000/tools/sap_oapi_run.define_model`` from the design model of
   ``datos_diseno_sap`` (beam sections = web rectangles + modifiers; analysis identical), or opened
   from a .$2k;
2. code "Eurocode 2-2004" and the 17 EC2 preferences, read back (A.1, A.2);
3. rebar B500SD, bar sizes (A.4, A.5); beam rectangles re-defined with modifiers computed from
   SAP's own ``GetSectProps``; column data of the pile (12Ø25, CHECK) and beam data (A.6);
4. design procedures: piles, transverse beams outside the piles and edge beams designed, rigid
   segments "No Design" (A.7);
5. overwrites per frame: DC Low; piles unbraced-length ratio 6.70/7.25 and beta = 1 (A.3); the
   table-only items TanTheta = 1 and KPhi through ``DatabaseTables`` (A.10, UNVERIFIED keys: on
   failure the script continues and reports what to set in the GUI);
6. combinations ELR01-22 defined; design combos ELU01-22 (or ELR01-22), auto combos off (A.8);
7. analysis, ``DesignConcrete.StartDesign``, summaries per frame, ``VerifyPassed`` (A.9);
8. analysis identity: frame forces of the 6 load cases compared with
   sap2000/resultados_sap/SAP27_Element_Forces_Frames.xlsx (current model, General sections).

Seismic variant (``--sismo``): the model also gets the response-spectrum analysis
(sap2000/tools/sismo_sap.py: mass source, MODAL run, EQX/EQY/EQZ, SIS*), SISXQ ... SISZ are design
combinations too (SAP uses the persistent gamma_c/gamma_s 1.5/1.15 for them; the Código Estructural
checks use the accidental 1.3/1.0), and the outputs are Muelle_Trasmallo_40m_diseno_sismo.sdb and
diseno_SAP2000_sismo.json/.xlsx.  Without ``--sismo`` nothing changes.

Output: diseno/sap/output/diseno_SAP2000.json and .xlsx (areas in cm², Asw/s in cm²/m,
positions in m: piles z from the fixity, transverse beams global y, edge beams global x).
Then ``python3 diseno/sap/leer_diseno_sap.py diseno/sap/output/diseno_SAP2000.json``.
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
for _p in (HERE, ROOT / "sap2000" / "tools", ROOT / "sap2000" / "model", ROOT / "diseno" / "python"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import datos_diseno_sap as D  # noqa: E402
import sap_oapi_run as sor  # noqa: E402
from sap_oapi_run import KN_M_C, check, rc  # noqa: E402

tm = D.tm
OUT_DIR = HERE / "output"
SAP_FRAMES_XLSX = ROOT / "sap2000" / "resultados_sap" / "SAP27_Element_Forces_Frames.xlsx"
M2_TO_CM2 = 1e4


# ============================================================================================
# settings
# ============================================================================================
def set_code_and_preferences(m) -> dict[int, float]:
    """A.1 + A.2: code string, 17 preferences, read back (a silent reset shows up here)."""
    check(m.DesignConcrete.SetCode(D.DESIGN_CODE), "DesignConcrete.SetCode")
    code = check(m.DesignConcrete.GetCode(""), "DesignConcrete.GetCode")[0]
    if code != D.DESIGN_CODE:
        raise RuntimeError(f"design code is {code!r}, expected {D.DESIGN_CODE!r}")
    ec2 = m.DesignConcrete.Eurocode_2_2004
    for item, val in D.EC2_PREFS.items():
        check(ec2.SetPreference(item, float(val)), f"EC2 SetPreference {item}")
    got = {}
    for item, val in D.EC2_PREFS.items():
        got[item] = float(check(ec2.GetPreference(item, 0.0), f"EC2 GetPreference {item}")[0])
        if abs(got[item] - val) > 1e-9:
            raise RuntimeError(f"EC2 preference {item}: set {val}, SAP has {got[item]}")
    return got


def define_rebar(m) -> list[str]:
    """A.4 + A.5: B500SD (uniaxial rebar) and the bar sizes the pile uses; returns the sizes added."""
    r = D.REBAR
    check(m.PropMaterial.SetMaterial(r.name, 6), f"SetMaterial {r.name}")
    check(m.PropMaterial.SetMPUniaxial(r.name, r.E, r.alpha), f"SetMPUniaxial {r.name}")
    check(m.PropMaterial.SetWeightAndMass(r.name, 1, r.unit_weight), f"SetWeightAndMass {r.name}")
    check(m.PropMaterial.SetORebar_1(r.name, r.fy, r.fu, r.expected * r.fy, r.expected * r.fu, 1, 1,
                                     r.strain_hardening, r.strain_ultimate, r.final_slope, False),
          f"SetORebar_1 {r.name}")
    added = []
    for name, (area, dia) in D.BAR_SIZES.items():
        if not _bar_exists(m, name):
            check(m.PropRebar.SetProp(name, area, dia), f"PropRebar.SetProp {name}")
            added.append(name)
    return added


def _bar_exists(m, name: str) -> bool:
    """A.5: only the return code of ``PropRebar.GetProp`` is used (its output layout is not settled);
    called as the 3P libraries do (name only), then with placeholders; if both forms are rejected by
    comtypes the size is assumed to exist (it is in SAP's default list)."""
    for args in ((name,), (name, 0.0, 0.0)):
        try:
            return rc(m.PropRebar.GetProp(*args)) == 0
        except TypeError:
            continue
    return True


def define_design_sections(m, base_sections: dict, info_areas: bool = True) -> dict:
    """A.6: beam rectangles (SetRectangle resets the section -> modifiers -> rebar), modifiers from
    SAP's own rectangle properties; pile column data (check) and beam data."""
    report = {}
    for name, (t3, t2) in D.DESIGN_RECT.items():
        fs = base_sections[name]
        check(m.PropFrame.SetRectangle(name, fs.material, t3, t2), f"SetRectangle {name}")
        sp = check(m.PropFrame.GetSectProps(name, *([0.0] * 12)), f"GetSectProps {name}")
        got = dict(zip(D.PROP_KEYS, (float(v) for v in sp[:6])))
        mods = D.rect_modifiers(fs.section.props, got, fs.modifiers)
        check(m.PropFrame.SetModifiers(name, [mods[k] for k in D.MOD_KEYS]), f"SetModifiers {name}")
        formula = D.rect_modifiers(fs.section.props, D.sap_rect_props(t3, t2), fs.modifiers)
        report[name] = {"t3": t3, "t2": t2, "sap_props": got, "modifiers": mods,
                        "max_rel_dev_vs_formula": max(abs(mods[k] / formula[k] - 1) for k in D.MOD_KEYS)}
    for br in D.BEAM_REBAR.values():
        check(m.PropFrame.SetRebarBeam(*br.oapi_args(info_areas)), f"SetRebarBeam {br.section}")
    check(m.PropFrame.SetRebarColumn(*D.PILE_REBAR.oapi_args()), f"SetRebarColumn {D.PILE_SECTION}")
    r = check(m.PropFrame.GetRebarColumn(D.PILE_SECTION, "", "", 0, 0, 0.0, 0, 0, 0, "", "", 0.0, 0, 0, False),
              "GetRebarColumn")
    # r = [MatPropLong, MatPropConfine, Pattern, ConfineType, Cover, NumberCBars, NumberR3Bars,
    #      NumberR2Bars, RebarSize, TieSize, TieSpacingLongit, Number2DirTieBars, Number3DirTieBars,
    #      ToBeDesigned, ret]
    if not (r[6] == D.PILE_REBAR.n_r3_bars and r[7] == D.PILE_REBAR.n_r2_bars and r[8] == D.PILE_REBAR.bar
            and not r[13]):
        raise RuntimeError(f"pile rebar read back {r}")
    return report


def define_design_groups(m, model: dict) -> None:
    for g, objs in D.design_groups(model).items():
        check(m.GroupDef.SetGroup(g), f"SetGroup {g}")
        for _, name in objs:
            check(m.FrameObj.SetGroupAssign(name, g), f"SetGroupAssign {name} -> {g}")


def set_design_procedures(m, model: dict) -> dict:
    """A.7; read back: 2 = concrete, 9 = no design."""
    fds = D.frame_designs(model)
    for fd in fds:
        check(m.FrameObj.SetDesignProcedure(fd.frame, fd.procedure, 0), f"SetDesignProcedure {fd.frame}")
    bad = []
    for fd in fds:
        got = check(m.FrameObj.GetDesignProcedure(fd.frame, 0), f"GetDesignProcedure {fd.frame}")[0]
        if got != (2 if fd.design else 9):
            bad.append((fd.frame, got))
    if bad:
        raise RuntimeError(f"design procedures not as intended: {bad[:5]}")
    return {"designed": sum(fd.design for fd in fds), "no_design": sum(not fd.design for fd in fds)}


def set_overwrites(m, model: dict, kphi) -> dict:
    """A.3: documented items per object (DC Low; piles ratio 0.924138 and beta 1), read back."""
    ec2 = m.DesignConcrete.Eurocode_2_2004
    rows = D.overwrites(model, kphi)
    for row in rows:
        for item, val in row["oapi"].items():
            check(ec2.SetOverwrite(row["frame"], item, float(val), 0), f"SetOverwrite {row['frame']} item {item}")
    readback = {}
    for row in (next(r for r in rows if r["kind"] == "pile"), next(r for r in rows if r["kind"] != "pile")):
        for item, val in row["oapi"].items():
            v, progdet = check(ec2.GetOverwrite(row["frame"], item, 0.0, False), "GetOverwrite")[:2]
            if abs(v - val) > 1e-6 or progdet:
                raise RuntimeError(f"overwrite {row['frame']} item {item}: set {val}, SAP has {v} (progdet {progdet})")
            readback[f"{row['frame']}:{item}"] = v
    return {"frames": len(rows), "readback": readback}


def table_key(m, pattern: str) -> str | None:
    raw = m.DatabaseTables.GetAllTables(0, [], [], [], [])
    if rc(raw) != 0:
        return None
    keys = [k for k in raw[1] if re.search(pattern, k, re.I)]
    return keys[0] if len(keys) == 1 else None


def apply_table_overwrites(m, model: dict, kphi, tan_theta: float = D.TAN_THETA) -> dict:
    """A.10: TanTheta (all designed frames) and KPhi (piles) are not documented OAPI items; they are
    edited in the EC2 overwrite database table.  Table key and field names are UNVERIFIED for v27:
    any failure is reported (and the edit cancelled) instead of raised."""
    rows_ow = D.overwrites(model, kphi, tan_theta)
    want = {r["frame"]: r["table_only"] for r in rows_ow}
    status = {"applied": False, "table": None, "rows_changed": 0, "rows_expected": len(want), "reason": "",
              "values": {"TanTheta": tan_theta,
                         "KPhi_piles": next(r for r in rows_ow if r["kind"] == "pile")["table"]["KPhi"]}}
    try:
        key = table_key(m, r"^overwrites\W+concrete design\W+eurocode 2-2004$")
        if key is None:
            status["reason"] = "EC2 overwrite table not found by DatabaseTables.GetAllTables"
            return status
        status["table"] = key
        raw = m.DatabaseTables.GetTableForEditingArray(key, "All", 0, [], 0, [])
        if rc(raw) != 0:
            status["reason"] = f"GetTableForEditingArray returned {rc(raw)}"
            return status
        version, fields, nrec, data = raw[0], list(raw[1]), int(raw[2]), list(raw[3])
        low = {f.lower(): i for i, f in enumerate(fields)}
        fi = low.get("frame")
        ti = low.get("tantheta")
        ki = low.get("kphi")
        if fi is None or ti is None or ki is None:
            m.DatabaseTables.CancelTableEditing()
            status["reason"] = f"fields Frame/TanTheta/KPhi not all found in {fields}"
            return status
        nf = len(fields)
        rows = [list(data[i * nf:(i + 1) * nf]) for i in range(nrec)]
        for r in rows:
            w = want.get(r[fi])
            if w:
                r[ti], r[ki] = f"{w['TanTheta']:g}", f"{w['KPhi']:g}"
                status["rows_changed"] += 1
        flat = [c for r in rows for c in r]
        ret = m.DatabaseTables.SetTableForEditingArray(key, version, fields, nrec, flat)
        if rc(ret) != 0:
            m.DatabaseTables.CancelTableEditing()
            status["reason"] = f"SetTableForEditingArray returned {rc(ret)}"
            return status
        res = m.DatabaseTables.ApplyEditedTables(True, 0, 0, 0, 0, "")
        if rc(res) != 0 or res[0] or res[1]:
            status["reason"] = f"ApplyEditedTables: fatal {res[0]}, errors {res[1]}; log: {str(res[4])[:400]}"
            return status
        status["applied"] = True
        if status["rows_changed"] < len(want):
            status["reason"] = (f"the table lists {status['rows_changed']} of the {len(want)} designed frames: "
                                "set the others in the GUI")
    except Exception as exc:                      # noqa: BLE001 - optional step, report and continue
        status["reason"] = f"{type(exc).__name__}: {exc}"
        try:
            m.DatabaseTables.CancelTableEditing()
        except Exception:                         # noqa: BLE001
            pass
    return status


def define_rom_combos(m) -> dict:
    """ELR01-22 (ROM 2.0-11, psi0,Qa = 1.0) as linear-add combos (the .$2k design files have them)."""
    added, existing = [], []
    for name, fac in D.combos("ROM").items():
        if m.RespCombo.Add(name, 0) != 0:
            existing.append(name)
            continue
        for pat, sf in zip(tm.PATTERN_ORDER, fac):
            if sf:
                check(m.RespCombo.SetCaseList(name, 0, pat, float(sf)), f"SetCaseList {name} {pat}")
        added.append(name)
    return {"added": len(added), "existing": len(existing)}


def select_design_combos(m, base: str, seismic: bool = False) -> list[str]:
    """A.8: auto combos off, only the ULS family of ``base`` (+ SIS* if ``seismic``) as Strength
    (robust deselection)."""
    check(m.DesignConcrete.SetComboAutoGenerate(False), "SetComboAutoGenerate")
    wanted = D.design_combo_names(base, seismic)
    for b in D.BASES:
        for name in D.design_combo_names(b):
            ret = m.DesignConcrete.SetComboStrength(name, name in wanted)
            if name in wanted:
                check(ret, f"SetComboStrength {name}")
    for fam, fams in tm.COMBOS.items():
        for k in range(1, len(fams) + 1):
            m.DesignConcrete.SetComboStrength(f"{fam}{k:02d}", f"{fam}{k:02d}" in wanted)
    for name in (D.SEISMIC_COMBOS if seismic else ()):
        check(m.DesignConcrete.SetComboStrength(name, True), f"SetComboStrength {name}")
    names = check(m.DesignConcrete.GetComboStrength(0, []), "GetComboStrength")[1]
    for extra in set(names or []) - set(wanted):   # e.g. DCon* created before auto-gen was off
        check(m.DesignConcrete.SetComboStrength(extra, False), f"deselect {extra}")
    names = check(m.DesignConcrete.GetComboStrength(0, []), "GetComboStrength")[1]
    if sorted(names or []) != sorted(wanted):
        raise RuntimeError(f"design combos {names}, expected {wanted}")
    return sorted(names)


# ============================================================================================
# analysis, design, results
# ============================================================================================
def run_analysis_and_design(m, save_path: Path, wait: float = 0.0, run_modal: bool = False) -> None:
    check(m.File.Save(str(save_path)), "File.Save")
    if not run_modal:                             # the seismic variant needs MODAL for EQX/EQY/EQZ
        try:
            m.Analyze.SetRunCaseFlag("MODAL", False)
        except Exception:                         # noqa: BLE001 - MODAL may not exist
            pass
    if wait:
        time.sleep(wait)                          # v27.0 #11904: licence not yet acquired
    check(m.Analyze.RunAnalysis(), "Analyze.RunAnalysis")
    check(m.DesignConcrete.StartDesign(), "DesignConcrete.StartDesign")
    if not m.DesignConcrete.GetResultsAvailable():
        raise RuntimeError("SAP2000 reports no concrete design results")


def _r(v, nd: int = 4):
    return round(float(v), nd)


def read_results(m, model: dict) -> tuple[list, list, dict]:
    """A.9: GetSummaryResultsColumn/Beam for every designed frame, VerifyPassed."""
    check(m.SetPresentUnits(KN_M_C), "SetPresentUnits")
    frames = {f["name"]: f for f in model["frames"]}
    joints = model["joints"]
    piles, beams = [], []
    for fd in D.frame_designs(model):
        if not fd.design:
            continue
        f = frames[fd.frame]
        xi, yi, zi = joints[f["i"]]
        if fd.design_type == "column":
            col = check(m.DesignConcrete.GetSummaryResultsColumn(fd.frame, 0, *[[] for _ in range(12)], 0),
                        f"GetSummaryResultsColumn {fd.frame}")
            n, frame, opt, loc, pcombo, parea, pratio, vmc, avmaj, vnc, avmin, err, warn = col[:13]
            piles += [{"frame": frame[k], "member": fd.member, "option": "check" if int(opt[k]) == 1 else "design",
                       "loc": _r(loc[k]), "z": _r(zi + loc[k]), "pmm_combo": pcombo[k],
                       "pmm_area_cm2": _r(parea[k] * M2_TO_CM2, 3), "pmm_ratio": _r(pratio[k], 5),
                       "vmaj_combo": vmc[k], "av_major_cm2_m": _r(avmaj[k] * M2_TO_CM2, 3),
                       "vmin_combo": vnc[k], "av_minor_cm2_m": _r(avmin[k] * M2_TO_CM2, 3),
                       "error": err[k] or "", "warning": warn[k] or ""} for k in range(n)]
        else:
            bm = check(m.DesignConcrete.GetSummaryResultsBeam(fd.frame, 0, *[[] for _ in range(14)], 0),
                       f"GetSummaryResultsBeam {fd.frame}")
            n, frame, loc, tc, ta, bc, ba, vc, va, tlc, tla, ttc, tta, err, warn = bm[:15]
            origin = yi if fd.kind == "beam_t" else xi
            beams += [{"frame": frame[k], "member": fd.member, "kind": fd.kind, "loc": _r(loc[k]),
                       "pos": _r(origin + loc[k]), "top_combo": tc[k], "as_top_cm2": _r(ta[k] * M2_TO_CM2, 3),
                       "bot_combo": bc[k], "as_bot_cm2": _r(ba[k] * M2_TO_CM2, 3), "v_combo": vc[k],
                       "asw_s_cm2_m": _r(va[k] * M2_TO_CM2, 3), "tl_combo": tlc[k],
                       "asl_t_cm2": _r(tla[k] * M2_TO_CM2, 3), "tt_combo": ttc[k],
                       "ast_s_cm2_m": _r(tta[k] * M2_TO_CM2, 3), "error": err[k] or "", "warning": warn[k] or ""}
                      for k in range(n)]
    ver = check(m.DesignConcrete.VerifyPassed(0, 0, 0, []), "VerifyPassed")
    verify = {"not_passed_or_unchecked": int(ver[0]), "failed": int(ver[1]), "not_checked": int(ver[2]),
              "frames": list(ver[3] or [])}
    return piles, beams, verify


BEAM_SECTIONS = (("cara pilote mar - voladizo", -0.20), ("cara pilote mar - vano", 0.20),
                 ("cara pilote tierra - vano", 3.25), ("cara pilote tierra - vuelo", 3.65))
SPAN = (0.20, 3.25)


def _at(rows: list[dict], pos: float, tol: float = 0.006) -> list[dict]:
    return [r for r in rows if abs(r["pos"] - pos) <= tol]


def _max(rows: list[dict], key: str) -> dict | None:
    return max(rows, key=lambda r: r[key]) if rows else None


def summarise(piles: list[dict], beams: list[dict]) -> dict:
    """Per member: governing pile ratio / beam steel at the design sections of vigas.py (D4)."""
    out_p = {}
    for member in dict.fromkeys(r["member"] for r in piles):
        rows = [r for r in piles if r["member"] == member]
        top, head, foot = _max(rows, "pmm_ratio"), max(rows, key=lambda r: r["z"]), min(rows, key=lambda r: r["z"])
        out_p[member] = {"pmm_ratio_max": top["pmm_ratio"], "z_max": top["z"], "combo": top["pmm_combo"],
                         "pmm_ratio_head": head["pmm_ratio"], "pmm_ratio_foot": foot["pmm_ratio"],
                         "av_major_max_cm2_m": max(r["av_major_cm2_m"] for r in rows),
                         "av_minor_max_cm2_m": max(r["av_minor_cm2_m"] for r in rows),
                         "messages": sorted({r["error"] or r["warning"] for r in rows} - {""})}
    out_b: dict = {}
    for member in dict.fromkeys(r["member"] for r in beams):
        rows = [r for r in beams if r["member"] == member]
        secs = []
        if member.startswith("VT"):
            for label, y in BEAM_SECTIONS:
                at = _at(rows, y)
                if at:
                    t, b, v = _max(at, "as_top_cm2"), _max(at, "as_bot_cm2"), _max(at, "asw_s_cm2_m")
                    secs.append({"section": label, "pos": y, "as_top_cm2": t["as_top_cm2"], "top_combo": t["top_combo"],
                                 "as_bot_cm2": b["as_bot_cm2"], "bot_combo": b["bot_combo"],
                                 "asw_s_cm2_m": v["asw_s_cm2_m"], "v_combo": v["v_combo"]})
            span = [r for r in rows if SPAN[0] - 1e-6 <= r["pos"] <= SPAN[1] + 1e-6]
            b = _max(span, "as_bot_cm2")
            if b:
                secs.append({"section": "vano: máx. positivo", "pos": b["pos"], "as_bot_cm2": b["as_bot_cm2"],
                             "bot_combo": b["bot_combo"]})
        t, b, v = _max(rows, "as_top_cm2"), _max(rows, "as_bot_cm2"), _max(rows, "asw_s_cm2_m")
        tl = _max(rows, "asl_t_cm2")
        out_b[member] = {"sections": secs,
                         "max": {"as_top_cm2": t["as_top_cm2"], "pos_top": t["pos"], "as_bot_cm2": b["as_bot_cm2"],
                                 "pos_bot": b["pos"], "asw_s_cm2_m": v["asw_s_cm2_m"], "pos_v": v["pos"],
                                 "asl_t_cm2": tl["asl_t_cm2"], "ast_s_cm2_m": max(r["ast_s_cm2_m"] for r in rows)},
                         "messages": sorted({r["error"] or r["warning"] for r in rows} - {""})}
    return {"piles": out_p, "beams": out_b}


def check_analysis_identity(m, model: dict, xlsx: Path = SAP_FRAMES_XLSX) -> dict:
    """Frame forces of the 6 load cases (piles, transverse beams) vs the SAP2000 v27.1 export of the
    current model (General sections): the rectangle + modifiers must not change the analysis.
    Tolerance = rounding of the Excel export (4-5 significant digits)."""
    import esfuerzos as E                         # noqa: PLC0415 - openpyxl only needed here
    ref: dict = {}
    for r in E.read_table(xlsx):
        if r["OutputCase"] in tm.PATTERN_ORDER:
            ref.setdefault((r["Frame"], r["OutputCase"]), []).append(r)
    sor.select_cases(m, list(tm.PATTERN_ORDER))
    comps = ("P", "V2", "V3", "T", "M2", "M3")
    worst = {c: (0.0, None) for c in comps}
    n = bad = 0
    for f in model["frames"]:
        if f["kind"] not in ("pile", "beam_t"):
            continue
        for row in sor.frame_force(m, f["name"]):
            cand = [r for r in ref.get((f["name"], row["case"]), []) if abs(float(r["Station"]) - row["station"]) < 1e-4]
            if not cand:
                continue
            n += 1
            for c in comps:
                d = abs(row[c] - float(cand[0][c]))
                if d > 0.02 + 2e-3 * abs(float(cand[0][c])):
                    bad += 1
                if d > worst[c][0]:
                    worst[c] = (d, f"{f['name']} @ {row['station']:.3f} {row['case']}")
    return {"reference": str(xlsx.relative_to(ROOT)) if xlsx.is_relative_to(ROOT) else str(xlsx),
            "rows_compared": n, "values_out_of_tolerance": bad, "ok": n > 0 and bad == 0,
            "max_abs_diff": {c: {"diff": round(v, 5), "at": at} for c, (v, at) in worst.items()}}


# ============================================================================================
# outputs
# ============================================================================================
def write_outputs(result: dict, out_dir: Path, stem: str = "diseno_SAP2000") -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    js = out_dir / f"{stem}.json"
    js.write_text(json.dumps(result, indent=1, ensure_ascii=False), encoding="utf-8")
    paths = {"json": js}
    try:
        from openpyxl import Workbook
    except ImportError:
        return paths
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name: str, rows: list[dict]) -> None:
        if not rows:
            return
        ws = wb.create_sheet(name[:31])
        cols = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) if not isinstance(r.get(c), (list, dict)) else json.dumps(r.get(c), ensure_ascii=False)
                       for c in cols])

    s = result["summary"]
    sheet("resumen_pilotes", [{"pilote": k, **v} for k, v in s["piles"].items()])
    sheet("resumen_vigas", [{"miembro": k, **sec} for k, v in s["beams"].items() for sec in v["sections"]]
          + [{"miembro": k, "section": "máximos", **v["max"]} for k, v in s["beams"].items()])
    sheet("pilotes", result["piles"])
    sheet("vigas", result["beams"])
    meta = result["meta"]
    sheet("meta", [{"clave": k, "valor": json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v}
                   for k, v in meta.items()])
    xl = out_dir / f"{stem}.xlsx"
    wb.save(xl)
    paths["xlsx"] = xl
    return paths


def run(m, *, base: str = "CYPE", kphi=D.KPHI_DEFAULT, info_areas: bool = True, from_s2k: Path | None = None,
        save_path: Path | None = None, out_dir: Path = OUT_DIR,
        verify_analysis: bool = True, wait: float = 0.0, xlsx: Path = SAP_FRAMES_XLSX,
        sismo: bool = False) -> dict:
    """Complete OAPI design run on the SapModel ``m`` (real SAP2000 or the test mock);
    ``sismo`` = seismic variant (module docstring)."""
    t0 = time.time()
    suffix = "_sismo" if sismo else ""
    save_path = save_path or OUT_DIR / f"Muelle_Trasmallo_40m_diseno{suffix}.sdb"
    model = D.base_model()
    dm = D.design_model(model)
    if from_s2k:
        check(m.File.OpenFile(str(Path(from_s2k).resolve())), "File.OpenFile (.$2k import)")
    else:
        sor.define_model(m, dm, seismic=sismo)
    try:
        m.SetModelIsLocked(False)
    except Exception:                             # noqa: BLE001
        pass
    check(m.SetPresentUnits(KN_M_C), "SetPresentUnits")
    prefs = set_code_and_preferences(m)
    bars_added = define_rebar(m)
    sections = define_design_sections(m, model["sections"], info_areas)
    if from_s2k:
        define_design_groups(m, model)
    combos_rom = define_rom_combos(m)
    procs = set_design_procedures(m, model)
    ows = set_overwrites(m, model, kphi)
    table_ow = apply_table_overwrites(m, model, kphi)
    combos = select_design_combos(m, base, sismo)
    save_path.parent.mkdir(parents=True, exist_ok=True)
    run_analysis_and_design(m, save_path.resolve(), wait, run_modal=sismo)
    piles, beams, verify = read_results(m, model)
    ident = check_analysis_identity(m, model, xlsx) if verify_analysis else None
    kp = D.KPHI_OPTIONS[kphi] if isinstance(kphi, str) else float(kphi)
    todo = []
    if not table_ow["applied"] or table_ow["rows_changed"] < table_ow["rows_expected"]:
        todo.append(f"GUI: Design > Concrete Frame Design > View/Revise Overwrites: piles (group DISENO_PILOTES) "
                    f"K phi = {kp:g}, Tan(theta) = {D.TAN_THETA:g}; beams (DISENO_VIGAS_T, DISENO_VIGAS_BORDE) "
                    f"Tan(theta) = {D.TAN_THETA:g}; then re-run the design ({table_ow['reason']})")
    meta = {
        "module": "diseno/sap/sap_oapi_diseno.py", "date": datetime.datetime.now().isoformat(timespec="seconds"),
        "code": D.DESIGN_CODE, "base": base, "design_combos": combos,
        "preferences": [{"item": p.item, "gui": p.gui, "value": prefs[p.item], "default": p.default, "why": p.why}
                        for p in D.PREFERENCES],
        "kphi_piles": kp, "tan_theta": D.TAN_THETA, "pile_length_ratio": D.PILE_LENGTH_RATIO,
        "l0_piles_m": D.PILE_LENGTH_RATIO * D.L_PILE, "framing_type": "DC Low",
        "overwrites": ows, "table_overwrites": table_ow, "procedures": procs, "rom_combos": combos_rom,
        "bar_sizes_added": bars_added, "beam_sections": sections, "beam_rebar_info_areas": info_areas,
        "provided": D.provided_summary(), "verify_passed": verify, "to_do": todo,
        "units": "areas cm2, Asw/s cm2/m, loc/pos/z m (piles z from the fixity, beams global y, edge beams global x)",
        "notes": ["SAP designs the web rectangle (sagging exact; hogging and torsion slightly conservative, spec B.3)",
                  "SAP does not check SLS (crack width 0.1 mm XS3, stress limits): see diseno/python/codigo/*.py",
                  "e2 uses c = 8 in SAP; Kphi = 1.1129 reproduces CYPE e2 (c = pi^2, Kphi 1.373, d 327.5)",
                  "piles: no P-Delta analysis and SAP's M0e = 0.6 M02 + 0.4 M01 >= 0.4 M02 (double curvature) -> "
                  "SAP's ratio is governed by the first-order end check (~ our nm1), not by the sway second-order "
                  "moment (our nm2): compare with sap_est_M0e of leer_diseno_sap.py (manual CFD-EC-2-2004 "
                  "3.4.2.2 and App. A)"],
        "runtime_s": round(time.time() - t0, 1),
    }
    if sismo:
        meta.update({"sismo": True, "gamma_seismic_combos": D.SEISMIC_GAMMA})
    result = {"meta": meta, "summary": summarise(piles, beams), "analysis_identity": ident,
              "piles": piles, "beams": beams}
    result["meta"]["outputs"] = {k: str(v) for k, v in write_outputs(result, out_dir, f"diseno_SAP2000{suffix}").items()}
    return result


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--attach", action="store_true", help="attach to the SAP2000 instance already open")
    ap.add_argument("--from-s2k", type=Path, default=None, help="open this .$2k instead of building the model")
    ap.add_argument("--base", choices=tuple(D.BASES), default="CYPE",
                    help="design combinations: CYPE = ELU01-22 (Anejo 10), ROM = ELR01-22 (psi0,Qa = 1.0)")
    ap.add_argument("--kphi", default=D.KPHI_DEFAULT,
                    help="ce (1.1129, default), ce_is (1.1874), cype (1.373), sap (1.0) or a number")
    ap.add_argument("--areas-cero", action="store_true", help="beam rebar areas 0 instead of the provided ones")
    ap.add_argument("--sin-verificacion", action="store_true", help="skip the analysis-identity check")
    ap.add_argument("--save", type=Path, default=None,
                    help="default output/Muelle_Trasmallo_40m_diseno.sdb (with --sismo: ..._diseno_sismo.sdb)")
    ap.add_argument("-o", "--out", type=Path, default=OUT_DIR)
    ap.add_argument("--sismo", action="store_true",
                    help="add the RS analysis (EQX/EQY/EQZ) and design SIS* too; outputs *_sismo")
    args = ap.parse_args()
    kphi = args.kphi if args.kphi in D.KPHI_OPTIONS else float(args.kphi)
    _sap, m = sor.connect(args.attach)
    res = run(m, base=args.base, kphi=kphi, info_areas=not args.areas_cero, from_s2k=args.from_s2k,
              save_path=args.save, out_dir=args.out, verify_analysis=not args.sin_verificacion,
              wait=5.0 if (args.from_s2k or not args.attach) else 0.0, sismo=args.sismo)
    for k, v in res["meta"]["outputs"].items():
        print(f"{k}: {v}")
    for p, s in res["summary"]["piles"].items():
        print(f"{p:4s} PMM ratio max {s['pmm_ratio_max']:.3f} at z = {s['z_max']:.2f} ({s['combo']})")
    if res["analysis_identity"]:
        print("analysis identical to SAP27 export:", res["analysis_identity"]["ok"])
    for line in res["meta"]["to_do"]:
        print("TO DO:", line)


if __name__ == "__main__":
    main()
