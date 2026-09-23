#!/usr/bin/env python3
"""validate_codigo.py - Codigo Estructural (2021) / Anejo 19 design formulas, validated against
the numbers CYPECAD 2023 printed in Anejo 10 (A10_CALC_ESTRUCTURAS.docx), Muelle de Trasmallo.

Run:  python3 validate_codigo.py            (about 30-60 s; numpy + scipy; openpyxl optional)
      python3 validate_codigo.py --json out.json   (also dump every check as JSON)

Every function below is written to be lifted as-is into the project code.  Units: N, mm, MPa
(forces printed in kN / kN*m).  Section sign convention = CYPE listing: compression positive,
N = sum(sigma dA), Mx = sum(sigma*y dA), My = sum(sigma*x dA).

Each check prints: id | description | clause | CYPE value | computed | delta | status.
Status PASS = within max(0.5 %, half a unit of the last digit CYPE printed); INFO = no CYPE
target (demonstration / sensitivity) or a CYPE convention that is documented, not matched.
Paragraph references "P###" point at the Anejo 10 text dump (a10_full.txt).
"""
from __future__ import annotations

import json
import math
import os
import sys

import numpy as np
from scipy.optimize import brentq, least_squares

# ============================================================================== reporting
RESULTS: list[dict] = []


def _decimals(s: str) -> int:
    s = s.strip().lstrip('+-')
    return len(s.split('.')[1]) if '.' in s else 0


def check(cid, desc, clause, cype, calc, unit='', tol=0.005, status=None, note=''):
    """cype: string exactly as printed by CYPE (None if there is no CYPE target)."""
    row = dict(id=cid, desc=desc, clause=clause, unit=unit, calc=float(calc), note=note)
    if cype is None:
        row.update(cype=None, delta=None, status=status or 'INFO')
    else:
        c = float(cype)
        half = 0.5 * 10 ** (-_decimals(cype))
        err = abs(calc - c)
        rel = err / abs(c) if c != 0 else (0.0 if err == 0 else float('inf'))
        ok = err <= max(tol * abs(c), half + 1e-12)
        row.update(cype=cype, delta=(calc - c) / abs(c) * 100 if c != 0 else 0.0,
                   status=status or ('PASS' if ok else 'FAIL'), rel=rel)
    RESULTS.append(row)
    return calc


def print_results():
    cur = None
    for r in RESULTS:
        grp = r['id'].split('.')[0]
        if grp != cur:
            cur = grp
            print('\n' + '=' * 132 + f'\n{GROUPS.get(grp, grp)}\n' + '=' * 132)
            print(f"{'id':7s} {'check':52s} {'clause':22s} {'CYPE':>11s} {'computed':>12s} {'delta':>8s}  status")
        cy = r['cype'] if r['cype'] is not None else '-'
        d = f"{r['delta']:+.2f}%" if r['delta'] is not None else ''
        calc = r['calc']
        cs = f'{calc:.6f}' if abs(calc) < 0.01 and calc != 0 else (f'{calc:.4f}' if abs(calc) < 10 else f'{calc:.2f}')
        line = f"{r['id']:7s} {r['desc'][:52]:52s} {r['clause'][:22]:22s} {cy:>11s} {cs:>12s} {d:>8s}  {r['status']}"
        if r['note']:
            line += '  ' + r['note']
        print(line)
    n = {s: sum(1 for r in RESULTS if r['status'] == s) for s in ('PASS', 'FAIL', 'INFO')}
    print('\n' + '=' * 132)
    print(f"SUMMARY: {n['PASS']} PASS, {n['FAIL']} FAIL, {n['INFO']} INFO  (tolerance: max(0.5 %, half unit of CYPE's last printed digit))")


GROUPS = {'M': 'M - MATERIALS (A19.2.4.2.4, A19.3.1, A19.3.2)',
          'C': 'C - LOAD COMBINATIONS (Anejo 18 eq. 6.10, CTE DB SE-AE psi as used by CYPE)',
          'U': 'U - ULS BENDING + AXIAL, FIBRE SECTION (A19.6.1, A19.3.1.7, A19.3.2.7)',
          'S': 'S - COLUMN SLENDERNESS / IMPERFECTIONS / NOMINAL CURVATURE (A19.5.2, A19.5.8)',
          'V': 'V - SHEAR (A19.6.2.2, A19.6.2.3, A19.9.2.2)',
          'D': 'D - DETAILING, MIN/MAX REINFORCEMENT (A19.8.2, A19.9.2, A19.9.5 + AN/UNE-EN 1992-1-1)',
          'K': 'K - SLS: CRACKING CRITERION, CRACK WIDTH, DEFLECTION LIMITS (A19.7.1, A19.7.3, CE Tabla 27.2)',
          'P': 'P - PSI2 SENSITIVITY WITH THE SAP2000 v27.1 RESULTS (quasi-permanent combination)'}

# ============================================================================== 1. materials
AGGREGATE_FACTOR = {'cuarcita': 1.0, 'caliza': 0.9, 'arenisca': 0.7, 'basalto': 1.2}   # A19.3.1.3(2)


def concrete(fck, gamma_c=1.5, alpha_cc=1.0, alpha_ct=1.0, aggregate='cuarcita'):
    """A19 Tabla 3.1 + 3.1.6 (alpha_cc = alpha_ct = 1.00 in Spain) + 3.1.3(2) aggregate factor."""
    fcm = fck + 8.0
    fctm = 0.30 * fck ** (2 / 3) if fck <= 50 else 2.12 * math.log(1 + fcm / 10)
    fctk005, fctk095 = 0.7 * fctm, 1.3 * fctm
    Ecm = 22000.0 * (fcm / 10) ** 0.3 * AGGREGATE_FACTOR[aggregate]
    if fck <= 50:
        eps_c2, eps_cu2, npar, lam, eta = 0.0020, 0.0035, 2.0, 0.8, 1.0
    else:   # NB: BOE Tabla A19.3.1 prints 0,85 in the eps_c2 formula; the tabulated values need 0,085
        eps_c2 = (2.0 + 0.085 * (fck - 50) ** 0.53) / 1000
        eps_cu2 = (2.6 + 35 * ((90 - fck) / 100) ** 4) / 1000
        npar = 1.4 + 23.4 * ((90 - fck) / 100) ** 4
        lam, eta = 0.8 - (fck - 50) / 400, 1.0 - (fck - 50) / 200
    return dict(fck=fck, fcm=fcm, fctm=fctm, fctk005=fctk005, fctk095=fctk095, Ecm=Ecm,
                eps_c2=eps_c2, eps_cu2=eps_cu2, n=npar, lam=lam, eta=eta,
                fcd=alpha_cc * fck / gamma_c, fctd=alpha_ct * fctk005 / gamma_c, gamma_c=gamma_c)


def fctm_fl(fctm, h_mm):
    """A19.3.1.8 (3.23): mean flexural tensile strength."""
    return max((1.6 - h_mm / 1000.0) * fctm, fctm)


def steel(fyk=500.0, gamma_s=1.15, Es=200000.0):
    fyd = fyk / gamma_s
    return dict(fyk=fyk, fyd=fyd, Es=Es, eps_yd=fyd / Es)


# ============================================================================== 2. fibre section
class RCSection:
    """Reinforced-concrete section = union of rectangles (x0, x1, y0, y1) + bars (x, y, area[, active]).

    ULS assumptions A19.6.1(2): plane sections, perfect bond, no concrete tension, parabola-rectangle
    concrete (3.1.7, n, eps_c2, eps_cu2), steel with horizontal top branch (3.2.7(2)b).
    Strain domains A19 Figura 6.1: pivot A (steel eps_su), B (eps_cu2), C (eps_c2 at (1-eps_c2/eps_cu2)h).
    CYPE convention reproduced with strain_factor=0.995: its failure plane reaches only 99.5 % of
    eps_cu2 and eps_su (=10 permil, a CYPE limit for the horizontal branch).  Gross concrete area
    (bars do not displace concrete) - also CYPE's convention.
    """

    def __init__(self, rects, bars, fcd, fyd, Es=200000.0, eps_c2=0.002, eps_cu2=0.0035, n=2.0,
                 eps_su=0.010, strain_factor=0.995, cell=2.0):
        self.rects = [tuple(map(float, r)) for r in rects]
        self.bars = np.array([b[:3] for b in bars], float)
        self.active = np.array([bool(b[3]) if len(b) > 3 else True for b in bars])
        self.fcd, self.fyd, self.Es, self.ec2, self.npar = fcd, fyd, Es, eps_c2, n
        self.ecu2, self.esu = eps_cu2 * strain_factor, eps_su * strain_factor
        self.vx = np.array([v for r in self.rects for v in (r[0], r[1], r[1], r[0])])
        self.vy = np.array([v for r in self.rects for v in (r[2], r[2], r[3], r[3])])
        self.set_mesh(cell)

    def set_mesh(self, cell):
        xs, ys, As = [], [], []
        for (x0, x1, y0, y1) in self.rects:
            nx = max(1, int(round((x1 - x0) / cell)))
            ny = max(1, int(round((y1 - y0) / cell)))
            dx, dy = (x1 - x0) / nx, (y1 - y0) / ny
            X, Y = np.meshgrid(x0 + (np.arange(nx) + 0.5) * dx, y0 + (np.arange(ny) + 0.5) * dy)
            xs.append(X.ravel()); ys.append(Y.ravel()); As.append(np.full(X.size, dx * dy))
        self.fx, self.fy, self.fa = np.concatenate(xs), np.concatenate(ys), np.concatenate(As)
        self.cell = cell

    # --- constitutive laws (compression +)
    def sig_c(self, e):
        e = np.asarray(e, float)
        par = self.fcd * (1.0 - (1.0 - np.clip(e, 0.0, self.ec2) / self.ec2) ** self.npar)
        return np.where(e > 0, np.where(e >= self.ec2, self.fcd, par), 0.0)

    def sig_s(self, e):
        return np.clip(self.Es * np.asarray(e, float), -self.fyd, self.fyd)

    # --- resultants of the plane eps(x, y) = e0 + kx*x + ky*y
    def resultants(self, e0, kx, ky, detail=False):
        ec = e0 + kx * self.fx + ky * self.fy
        Fc = self.sig_c(ec) * self.fa
        bx, by, ba = self.bars.T
        es = e0 + kx * bx + ky * by
        ss = self.sig_s(es) * self.active
        Fs = ss * ba
        N = Fc.sum() + Fs.sum()
        Mx = (Fc * self.fy).sum() + (Fs * by).sum()
        My = (Fc * self.fx).sum() + (Fs * bx).sum()
        if not detail:
            return np.array([N, Mx, My])
        Cc = Fc.sum(); comp, ten = Fs > 0, Fs < 0
        Cs, T = Fs[comp].sum(), -Fs[ten].sum()
        ev = e0 + kx * self.vx + ky * self.vy            # strains at the section vertices
        return dict(N=N, Mx=Mx, My=My, Cc=Cc, Cs=Cs, T=T,
                    ecc=((Fc * self.fx).sum() / Cc, (Fc * self.fy).sum() / Cc) if Cc else (0, 0),
                    ecs=((Fs[comp] * bx[comp]).sum() / Cs, (Fs[comp] * by[comp]).sum() / Cs) if Cs else (0, 0),
                    eT=((Fs[ten] * bx[ten]).sum() / -T, (Fs[ten] * by[ten]).sum() / -T) if T else (0, 0),
                    eps_c_max=ev.max(), sig_c_max=float(self.sig_c(ev.max())),
                    eps_s_min=es[self.active].min(), bar_eps=es, bar_sig=ss, plane=(e0, kx, ky))

    # --- ultimate strain plane, parameter t in [0,3]: 0-1 pivot A, 1-2 pivot B, 2-3 pivot C
    def ult_plane(self, alpha, t):
        ca, sa = math.cos(alpha), math.sin(alpha)          # (ca, sa) points to the compressed side
        s_v = self.vx * ca + self.vy * sa
        s_top, s_bot = s_v.max(), s_v.min(); h = s_top - s_bot
        s_bar = (self.bars[self.active, 0] * ca + self.bars[self.active, 1] * sa).min()
        ecu, ec2, esu = self.ecu2, self.ec2, self.esu
        if t <= 1.0:
            e_top = -esu + t * (ecu + esu)
            k = (e_top + esu) / (s_top - s_bar)
            e_at = lambda s: -esu + k * (s - s_bar)
        elif t <= 2.0:
            xAB = ecu / (ecu + esu) * (s_top - s_bar)
            x = xAB + (t - 1.0) * (h - xAB)
            e_at = lambda s: ecu - ecu / x * (s_top - s)
        else:
            sC = s_top - (1 - ec2 / ecu) * h
            e_top = ecu + (t - 2.0) * (ec2 - ecu)
            k = (e_top - ec2) / (s_top - sC)
            e_at = lambda s: ec2 + k * (s - sC)
        e0 = e_at(0.0); kk = e_at(1.0) - e0
        return e0, kk * ca, kk * sa

    def ult(self, alpha, t, detail=False):
        return self.resultants(*self.ult_plane(alpha, t), detail=detail)

    def capacity_ray(self, S, coarse=10.0):
        """NRd, MRd with the same eccentricities as S=(N, Mx, My) [N, N*mm]; utilisation eta = |S|/|R|
        (CYPE: 'esfuerzos que producen el agotamiento con las mismas excentricidades')."""
        S = np.asarray(S, float)
        L = max(np.ptp(self.vx), np.ptp(self.vy)); sc = np.array([1.0, 1 / L, 1 / L])
        shat = S * sc / np.linalg.norm(S * sc)
        fine = self.cell; self.set_mesh(coarse)
        best = (9.0, 0.0, 0.0)
        for a in np.radians(np.arange(0.0, 360.0, 3.0)):
            for t in np.linspace(0.0, 3.0, 121):
                r = self.ult(a, t) * sc
                nr = np.linalg.norm(r)
                if nr > 0:
                    err = np.linalg.norm(r / nr - shat)
                    if err < best[0]:
                        best = (err, a, t)
        self.set_mesh(fine)

        def res(p):
            r = self.ult(p[0], p[1]) * sc
            return r / np.linalg.norm(r) - shat
        sol = least_squares(res, [best[1], best[2]], bounds=([-10, 0], [20, 3]), xtol=1e-13, ftol=1e-13)
        d = self.ult(sol.x[0], sol.x[1], detail=True)
        R = np.array([d['N'], d['Mx'], d['My']])
        d['eta'] = np.linalg.norm(S * sc) / np.linalg.norm(R * sc)
        d['alpha'], d['t'] = sol.x
        return d

    def capacity_constant_N(self, S):
        """Utilisation at constant N and constant Mx:My ratio (reproduces CYPE's 'Aprov.' in the
        pile summary listing, section 3.5.1 of the appendix)."""
        S = np.asarray(S, float)
        r0 = self.capacity_ray(S)
        ang = math.atan2(S[2], S[1])

        def res(p):
            R = self.ult(p[0], p[1])
            return [(R[0] - S[0]) / 1e5, math.atan2(R[2], R[1]) - ang]
        sol = least_squares(res, [r0['alpha'], r0['t']], bounds=([-10, 0], [20, 3]), xtol=1e-13)
        R = self.ult(*sol.x)
        return math.hypot(S[1], S[2]) / math.hypot(R[1], R[2]), R

    def capacity_uniaxial(self, N, compression_side=+1):
        """MRd,x for a given N, bending about x; compression_side=+1 top (+y), -1 bottom."""
        a = math.pi / 2 if compression_side > 0 else -math.pi / 2
        t = brentq(lambda t: self.ult(a, t)[0] - N, 0.0, 3.0, xtol=1e-12)
        return self.ult(a, t, detail=True)

    def equilibrium(self, S):
        """Strain plane in equilibrium with S=(N, Mx, My) (CYPE 'Equilibrio ... esfuerzos solicitantes')."""
        S = np.asarray(S, float)
        L = max(np.ptp(self.vx), np.ptp(self.vy))
        scale = np.array([1.0, 1 / L, 1 / L]) / max(abs(S[0]), np.linalg.norm(S[1:]) / L, 1.0)
        res = lambda p: (self.resultants(p[0] * 1e-3, p[1] * 1e-5, p[2] * 1e-5) - S) * scale
        best = None
        for e0 in (0.5, 0.0):
            for kx in (-1.0, 0.0, 1.0):
                for ky in (-1.0, 0.0, 1.0):
                    sol = least_squares(res, [e0, kx, ky], xtol=1e-15, ftol=1e-15, gtol=1e-15, max_nfev=3000)
                    if best is None or sol.cost < best.cost:
                        best = sol
            if best.cost < 1e-20:
                break
        p = best.x
        return self.resultants(p[0] * 1e-3, p[1] * 1e-5, p[2] * 1e-5, detail=True)


def required_tension_steel(rects, bar_xy, M, fcd, fyd, compression_side, cell=2.0):
    """As needed (singly reinforced: no compression/side steel, N = 0) so that MRd = |M|.
    Reproduces CYPE 'Area Sup./Inf. Nec.' in the beam listing (appendix section 2)."""
    def mrd(As):
        bars = [(x, y, As / len(bar_xy)) for x, y in bar_xy]
        sec = RCSection(rects, bars, fcd, fyd, cell=cell)
        return abs(sec.capacity_uniaxial(0.0, compression_side)['Mx'])
    return brentq(lambda a: mrd(a) - abs(M), 50.0, 20000.0, xtol=0.01)


# ============================================================================== 3. columns
def slenderness(l0, I, A):
    """A19.5.8.3.2: lambda = l0 / i, i = sqrt(I/A) (gross, uncracked concrete section)."""
    i = math.sqrt(I / A)
    return l0 / i, i


def lambda_lim(phi_ef, omega, n, rm=None):
    """A19.5.8.3.1 (5.13).  B = sqrt(1+2w) as in UNE-EN 1992-1-1 (the BOE text prints
    1 + sqrt(1+2w), an erratum; CYPE prints 'Formulacion UNE-EN 1992-1-1:2013').
    C = 1.7 - rm, or 0.7 when rm is unknown / member unbraced (CYPE: 0.7)."""
    A = 1.0 / (1.0 + 0.2 * phi_ef)
    B = math.sqrt(1.0 + 2.0 * omega)
    C = 0.7 if rm is None else 1.7 - rm
    return 20.0 * A * B * C / math.sqrt(n), A, B, C


def imperfection_ei(l0, l, m=1, theta0=1 / 200):
    """A19.5.2 (5.1), (5.2): ei = theta_i * l0 / 2, theta_i = theta0*alpha_h*alpha_m."""
    alpha_h = min(max(2.0 / math.sqrt(l), 2.0 / 3.0), 1.0)
    alpha_m = math.sqrt(0.5 * (1.0 + 1.0 / m))
    theta_i = theta0 * alpha_h * alpha_m
    return theta_i * l0 / 2.0, theta_i, alpha_h, alpha_m


def nominal_curvature(fyd, Es, d, n, omega, phi_ef, fck, lam, l0, c=math.pi ** 2, n_bal=0.4):
    """A19.5.8.8.2-3: e2 = (1/r) l0^2 / c ; 1/r = Kr Kphi eps_yd/(0.45 d)."""
    eps_yd = fyd / Es
    r0inv = eps_yd / (0.45 * d)
    n_u = 1.0 + omega
    Kr = min((n_u - n) / (n_u - n_bal), 1.0)
    beta = 0.35 + fck / 200.0 - lam / 150.0
    Kphi = max(1.0 + beta * phi_ef, 1.0)
    rinv = Kr * Kphi * r0inv
    return dict(eps_yd=eps_yd, r0inv=r0inv, n_u=n_u, Kr=Kr, beta=beta, Kphi=Kphi, rinv=rinv,
                e2=rinv * l0 ** 2 / c)


def min_eccentricity(h):
    """A19.6.1(4): e0 = max(h/30, 20 mm)."""
    return max(h / 30.0, 20.0)


# ============================================================================== 4. shear
def VRd_c(fck, gamma_c, bw, d, Asl, NEd, Ac, fcd, k1=0.15):
    """A19.6.2.2(1) (6.2.a/b), (6.3): members without shear reinforcement (N, mm)."""
    CRdc = 0.18 / gamma_c
    k = min(1.0 + math.sqrt(200.0 / d), 2.0)
    rho_l = min(Asl / (bw * d), 0.02)
    sigma_cp = min(NEd / Ac, 0.2 * fcd)
    vmin = 0.035 * k ** 1.5 * math.sqrt(fck)
    V1 = (CRdc * k * (100.0 * rho_l * fck) ** (1 / 3) + k1 * sigma_cp) * bw * d
    V2 = (vmin + k1 * sigma_cp) * bw * d
    return dict(V=max(V1, V2), V1=V1, Vmin=V2, CRdc=CRdc, k=k, rho_l=rho_l, sigma_cp=sigma_cp, vmin=vmin)


def nu1_shear(fck, stirrup_stress_below_08fyk=True):
    """A19.6.2.3(3): nu1 = 0.6(1-fck/250); = 0.6 (fck<=60) if fywd <= 0.8 fywk (6.10.a/b)."""
    if stirrup_stress_below_08fyk:
        return 0.6 if fck <= 60 else max(0.9 - fck / 200.0, 0.5)
    return 0.6 * (1.0 - fck / 250.0)


def alpha_cw(sigma_cp, fcd, prestressed=False):
    """A19.6.2.3(3) (6.11): 1 for non-prestressed structures (CYPE: sigma_cp <= 0 -> 1)."""
    if not prestressed or sigma_cp <= 0:
        return 1.0
    if sigma_cp <= 0.25 * fcd:
        return 1.0 + sigma_cp / fcd
    if sigma_cp <= 0.5 * fcd:
        return 1.25
    return 2.5 * (1.0 - sigma_cp / fcd)


def sigma_cp_cype(NEd, As_comp, fyd, Ac):
    """CYPE's sigma_cp for alpha_cw: (NEd - A's fyd)/Ac  (EHE-08 44.2.3.1 heritage)."""
    return (NEd - As_comp * fyd) / Ac


def VRd_max(bw, z, nu1, fcd, cot_theta=1.0, alpha_deg=90.0, acw=1.0):
    """A19.6.2.3(3)/(4) (6.9)/(6.14)."""
    cot_a = 1.0 / math.tan(math.radians(alpha_deg)) if alpha_deg != 90 else 0.0
    return acw * bw * z * nu1 * fcd * (cot_theta + cot_a) / (1.0 + cot_theta ** 2)


def VRd_s(Asw, s, z, fywd, cot_theta=1.0, alpha_deg=90.0):
    """A19.6.2.3(3)/(4) (6.8)/(6.13).  0.5 <= cot(theta) <= 2.0 in Spain (6.7)."""
    a = math.radians(alpha_deg)
    cot_a = 0.0 if alpha_deg == 90 else 1.0 / math.tan(a)
    return Asw / s * z * fywd * (cot_theta + cot_a) * math.sin(a)


def asw_required(VEd, z, fywd, cot_theta=1.0):
    """Asw/s [mm2/mm] from VEd <= VRd,s (vertical stirrups)."""
    return VEd / (z * fywd * cot_theta)


def rho_w_min(fck, fyk):
    """A19.9.2.2(5) (9.5)."""
    return 0.08 * math.sqrt(fck) / fyk


def s_l_max(d, alpha_deg=90.0):
    """A19.9.2.2(6) (9.6)."""
    return 0.75 * d * (1.0 + (0.0 if alpha_deg == 90 else 1 / math.tan(math.radians(alpha_deg))))


def s_t_max(d):
    """A19.9.2.2(8) (9.8)."""
    return min(0.75 * d, 600.0)


# ============================================================================== 5. detailing
def clear_spacing_min(phi, dg, k1=1.0, k2=5.0):
    """A19.8.2(2): max(k1*phi, dg + k2, 20 mm), k1 = 1, k2 = 5 mm.  (CYPE prints s2 = 1.25 dg)."""
    return max(k1 * phi, dg + k2, 20.0)


def column_As_min(NEd, fyd, Ac, fyc_d_cap=400.0):
    """A19.9.5.2(2): per face 0.05 NEd/fyc,d (fyc,d = fyd <= 400); symmetric 0.10 NEd/fyd.
    AN/UNE-EN 1992-1-1 9.5.2(2) (used by CYPE): additionally 0.004 Ac."""
    return dict(per_face=0.05 * NEd / min(fyd, fyc_d_cap), symmetric=0.10 * NEd / fyd, geometric_AN=0.004 * Ac)


def column_As_max(Ac, fcd, fyd):
    """A19.9.5.2(3): 0.04 Ac (0.08 Ac at laps).  AN/UNE-EN (used by CYPE): 0.5 fcd Ac/fyc,d per face."""
    return dict(CE=0.04 * Ac, CE_laps=0.08 * Ac, AN_total=fcd * Ac / fyd)


def s_cl_max(phi_min, b, h):
    """A19.9.5.3(3): min(15 phi_min, 300 mm, min(b, h)) (BOE prints 'min(G,h)', G = b)."""
    return min(15.0 * phi_min, 300.0, min(b, h)), (15.0 * phi_min, 300.0, min(b, h))


def phi_t_min(phi_long_max):
    """A19.9.5.3(1): >= 6 mm and >= phi_max/4."""
    return max(6.0, phi_long_max / 4.0)


def beam_As_min(W, z, fctm_fl_, fyd):
    """A19.9.2.1.1(1) (9.1) Spanish: W/z * fctm,fl/fyd (z ~ 0.8h per code; CYPE uses 0.9d)."""
    return W / z * fctm_fl_ / fyd


# ============================================================================== 6. SLS
def homogenised_stresses(rects, bars, n_mod, M, N=0.0):
    """Uncracked homogenised section (concrete in tension included), bending about x.
    Returns stresses at top/bottom fibre and in the bars (tension negative)."""
    parts = [((x1 - x0) * (y1 - y0), (y0 + y1) / 2, (x1 - x0) * (y1 - y0) ** 3 / 12) for x0, x1, y0, y1 in rects]
    A = sum(a for a, _, _ in parts) + sum((n_mod - 1) * a for _, _, a in bars)
    Sy = sum(a * y for a, y, _ in parts) + sum((n_mod - 1) * a * y for _, y, a in bars)
    yc = Sy / A
    I = sum(i + a * (y - yc) ** 2 for a, y, i in parts) + sum((n_mod - 1) * a * (y - yc) ** 2 for _, y, a in bars)
    ytop = max(r[3] for r in rects); ybot = min(r[2] for r in rects)
    sig = lambda y: N / A + M * (y - yc) / I
    return dict(yc=yc, I=I, A=A, top=sig(ytop), bot=sig(ybot),
                bars=[n_mod * sig(y) for _, y, _ in bars])


def cracked_elastic(rects, bars, n_mod, M):
    """Fully cracked linear-elastic section (no concrete tension), bending about x, N = 0.
    Returns neutral-axis position and steel stresses (sigma_s for 7.3.4)."""
    ytop = max(r[3] for r in rects); ybot = min(r[2] for r in rects)
    sag = M > 0      # sagging: compression on top

    def force(yn):   # net axial force for a neutral axis at y = yn, unit curvature
        F = 0.0
        for x0, x1, y0, y1 in rects:
            b = x1 - x0
            if sag:
                lo, hi = max(y0, yn), y1
                if hi > lo:
                    F += b * ((hi - yn) ** 2 - (lo - yn) ** 2) / 2
            else:
                lo, hi = y0, min(y1, yn)
                if hi > lo:
                    F += b * ((yn - lo) ** 2 - (yn - hi) ** 2) / 2
        for x, y, a in bars:
            F += n_mod * a * ((y - yn) if sag else (yn - y))
        return F
    yn = brentq(force, ybot + 1e-6, ytop - 1e-6)
    # second moment about NA
    I = 0.0
    for x0, x1, y0, y1 in rects:
        b = x1 - x0
        lo, hi = (max(y0, yn), y1) if sag else (y0, min(y1, yn))
        if hi > lo:
            I += b * ((hi - yn) ** 3 - (lo - yn) ** 3) / 3
    I += sum(n_mod * a * (y - yn) ** 2 for x, y, a in bars)
    x = (ytop - yn) if sag else (yn - ybot)
    return dict(yn=yn, x=x, Icr=I, sig_s=[n_mod * abs(M) * abs(y - yn) / I for _, y, _ in bars],
                sig_c=abs(M) * x / I)


def crack_width(sig_s, Es, Ecm, fct_eff, As, Ac_eff, c, phi, kt=0.4, k1=0.8, k2=0.5, k3=3.4, k4=0.425,
                spacing=None, h=None, x=None):
    """A19.7.3.4 (7.8)-(7.11), (7.14)."""
    alpha_e = Es / Ecm
    rho = As / Ac_eff
    d_eps = max((sig_s - kt * fct_eff / rho * (1 + alpha_e * rho)) / Es, 0.6 * sig_s / Es)
    if spacing is not None and spacing > 5 * (c + phi / 2):
        sr = 1.3 * (h - x)
    else:
        sr = k3 * c + k1 * k2 * k4 * phi / rho
    return dict(wk=sr * d_eps, sr_max=sr, d_eps=d_eps, rho_p_eff=rho)


# ============================================================================== CYPE data
A = lambda phi: math.pi * phi ** 2 / 4
C50, C35, C25 = concrete(50, aggregate='caliza'), concrete(35, aggregate='caliza'), concrete(25, aggregate='caliza')
B500 = steel()
fyd, Es = B500['fyd'], B500['Es']


def pile_bars(c, m, phi=25.0):
    """12 bars in a 400x400 pile: corners at +-c, intermediate at +-m (CYPE numbering 1..12)."""
    pts = [(-c, c), (-m, c), (m, c), (c, c), (c, m), (c, -m), (c, -c), (m, -c), (-m, -c), (-c, -c), (-c, -m), (-c, m)]
    return [(x, y, A(phi)) for x, y in pts]


PILE_RECT = [(-200.0, 200.0, -200.0, 200.0)]
PILE_FORJ = pile_bars(127.5, 42.5)            # cover 50 + tie 10 + 12.5 (P257)
PILE_EMP = pile_bars(129.5, 259.0 / 6)        # cover 50 + tie 8 + 12.5 (P312), 43.17
# inner transverse beam "50x55+15x30+15x30" (inverted T), coordinates w.r.t. gross centroid (P381)
YC_T = (500 * 550 * 275 + 2 * 150 * 300 * 150) / (500 * 550 + 2 * 150 * 300)
YB, YT = -YC_T, 550 - YC_T
BEAM_RECT = [(-250, 250, YB, YT), (-400, -250, YB, YB + 300), (250, 400, YB, YB + 300)]
BEAM_TOP = [(x, YT - 70, A(20)) for x in (-180, -97.5, 0, 97.5, 180)]
BEAM_BOT = [(x, YB + 70, A(20)) for x in (330, 180, 97.5, 0, -97.5, -180, -330)]
BEAM_SIDE = [(185, YB + 235, A(10), False), (335, YB + 235, A(10), True),       # CYPE gives 0 stress to
             (-335, YB + 235, A(10), True), (-185, YB + 235, A(10), False)]     # the two web O10 (bars 6, 16)
BEAM_BARS = BEAM_TOP[:5] + [BEAM_SIDE[0], BEAM_SIDE[1]] + BEAM_BOT + [BEAM_SIDE[2], BEAM_SIDE[3]]  # CYPE order 1..16


def run():
    # -------------------------------------------------------------------------- M materials
    check('M.01', 'fcd HA-50 = acc*fck/gc (acc=1.00)', 'A19.3.1.6(1)', '33.33', C50['fcd'], 'MPa')
    check('M.02', 'fcd HA-35', 'A19.3.1.6(1)', '23.33', C35['fcd'], 'MPa')
    check('M.03', 'fyd B500SD = fyk/gs', 'A19.2.4.2.4', '434.78', fyd, 'MPa')
    check('M.04', 'eps_yd = fyd/Es (Es = 200 GPa)', 'A19.3.2.7(4)', '0.00217', B500['eps_yd'])
    check('M.05', 'fctm HA-35 = 0.30 fck^(2/3)', 'A19 Tabla 3.1', '3.21', C35['fctm'], 'MPa')
    check('M.06', 'fctm,fl HA-35, h = 550 mm', 'A19.3.1.8 (3.23)', '3.37', fctm_fl(C35['fctm'], 550), 'MPa')
    check('M.07', 'Ecm HA-35 limestone (x0.9)', 'A19.3.1.3(2)', '30669', C35['Ecm'], 'MPa')
    check('M.08', 'Ecm HA-50 limestone (x0.9)', 'A19.3.1.3(2)', '33550', C50['Ecm'], 'MPa')
    check('M.09', 'Ecm HA-25 limestone (x0.9)', 'A19.3.1.3(2)', '28328', C25['Ecm'], 'MPa')
    check('M.10', 'eps_cu2 (fck<=50)', 'A19 Tabla 3.1', '0.0035', C50['eps_cu2'])
    check('M.11', 'eps_c2 (fck<=50)', 'A19 Tabla 3.1', '0.0020', C50['eps_c2'])
    check('M.12', 'Steel stress at eps=-0.00995 (horizontal branch)', 'A19.3.2.7(2)b', '-434.78',
          float(np.clip(Es * -0.00995, -fyd, fyd)), 'MPa', note='P381 bars 1-5: -434.78 @ -0.009950')
    check('M.13', 'Steel stress at eps=+0.001826 (elastic, Es)', 'A19.3.2.7(4)', '365.17', Es * 0.001826, 'MPa')

    # -------------------------------------------------------------------------- C combinations
    # P3 hypotheses (appendix 3.3, head / base values; rounded to 0.1 in the listing -> tol 0.3 %)
    head = dict(PP=(86.0, 7.0, -0.9, -1.6), CM=(27.5, 2.7, -0.3, -0.6), Qa=(235.7, 23.4, -2.3, -5.4), TB=(146.2, 167.1, 0.1, -51.1))
    base = dict(PP=112.3, CM=27.5, Qa=235.7, TB=146.2)
    comb = lambda d, gG, gQ, gT, i: gG * (d['PP'][i] + d['CM'][i]) + gQ * d['Qa'][i] + gT * d['TB'][i]
    check('C.01', 'NEd head 1.35G+1.05Qa+1.5TB (1.05=1.5*0.7)', 'A18 (6.10); CTE psi0', '620.06', comb(head, 1.35, 1.05, 1.5, 0), 'kN', tol=0.003)
    check('C.02', 'MEd,x head same comb. (listing My)', 'A18 (6.10)', '288.36', comb(head, 1.35, 1.05, 1.5, 1), 'kN*m', tol=0.003)
    check('C.03', 'VEd,y head same comb.', 'A18 (6.10)', '85.36', -comb(head, 1.35, 1.05, 1.5, 3), 'kN', tol=0.003)
    check('C.04', 'NEd base (Empotramiento) same comb.', 'A18 (6.10)', '655.55',
          1.35 * (base['PP'] + base['CM']) + 1.05 * base['Qa'] + 1.5 * base['TB'], 'kN', tol=0.003)
    check('C.05', 'NEd max: 1.35G+1.5Qa+0.9TB (0.9=1.5*0.6)', 'A18 (6.10); CTE psi0', '673.90',
          1.35 * (base['PP'] + base['CM']) + 1.5 * base['Qa'] + 0.9 * base['TB'], 'kN', tol=0.003)
    check('C.06', 'NEd min for VRd,c: 1.0G+1.5TB (head)', 'A18 (6.10)', '332.81', comb(head, 1.0, 0.0, 1.5, 0), 'kN', tol=0.003)
    check('C.07', 'VEd,y for VRd,c: 1.0G+1.5TB (head)', 'A18 (6.10)', '78.86', -comb(head, 1.0, 0.0, 1.5, 3), 'kN', tol=0.003)

    # -------------------------------------------------------------------------- U ULS sections
    t0 = __import__('time').time()
    forj = RCSection(PILE_RECT, PILE_FORJ, C50['fcd'], fyd, cell=1.0)
    emp = RCSection(PILE_RECT, PILE_EMP, C50['fcd'], fyd, cell=1.0)
    r = forj.capacity_ray([620.06e3, 353.51e6, -68.95e6])     # P245-P261 (2nd order, 'Cabeza')
    check('U.01', 'Pile head NRd (2nd-order ray)', 'A19.6.1', '679.16', r['N'] / 1e3, 'kN')
    check('U.02', 'Pile head MRd,x', 'A19.6.1', '387.20', r['Mx'] / 1e6, 'kN*m')
    check('U.03', 'Pile head MRd,y', 'A19.6.1', '-75.52', r['My'] / 1e6, 'kN*m')
    check('U.04', 'Pile head Cc (concrete resultant)', 'A19.6.1', '1422.97', r['Cc'] / 1e3, 'kN')
    check('U.05', 'Pile head Cs', 'A19.6.1', '536.86', r['Cs'] / 1e3, 'kN')
    check('U.06', 'Pile head T', 'A19.6.1', '1280.67', r['T'] / 1e3, 'kN')
    check('U.07', 'Pile head ecc,y', 'A19.6.1', '140.81', r['ecc'][1], 'mm')
    check('U.08', 'Pile head eT,y', 'A19.6.1', '-93.92', r['eT'][1], 'mm')
    check('U.09', 'Pile head eps bar 7 (most tensioned)', 'A19.6.1(3)', '-0.004001', r['bar_eps'][6])
    check('U.10', 'Pile head eps bar 1', 'A19.6.1(3)', '0.001826', r['bar_eps'][0])
    check('U.11', 'Pile head eps_c,max (=0.995*eps_cu2)', 'A19.6.1(3)', '0.0035', r['eps_c_max'],
          note='CYPE stops at 99.5 % of eps_cu2')
    check('U.12', 'Pile head eta = NSd/NRd (instability)', 'A19.6.1', '0.913', 620.06e3 / r['N'])
    r = forj.capacity_ray([620.06e3, 288.36e6, -3.81e6])
    check('U.13', 'Pile head 1st-order NRd', 'A19.6.1', '925.83', r['N'] / 1e3, 'kN')
    check('U.14', 'Pile head 1st-order MRd,x', 'A19.6.1', '430.56', r['Mx'] / 1e6, 'kN*m')
    check('U.15', 'Pile head 1st-order MRd,y', 'A19.6.1', '-5.68', r['My'] / 1e6, 'kN*m')
    check('U.16', 'Pile head eta section = NEd/NRd', 'A19.6.1', '0.670', 620.06e3 / r['N'])
    r = emp.capacity_ray([655.55e3, -283.56e6, 2.79e6])       # P300
    check('U.17', 'Pile fixity 1st-order NRd', 'A19.6.1', '1021.59', r['N'] / 1e3, 'kN')
    check('U.18', 'Pile fixity 1st-order MRd,x', 'A19.6.1', '-441.89', r['Mx'] / 1e6, 'kN*m')
    check('U.19', 'Pile fixity 1st-order MRd,y', 'A19.6.1', '4.34', r['My'] / 1e6, 'kN*m')
    check('U.20', 'Pile fixity eta section', 'A19.6.1', '0.642', 655.55e3 / r['N'])
    r = emp.capacity_ray([655.55e3, -352.07e6, 71.29e6])
    check('U.21', 'Pile fixity 2nd-order NRd', 'A19.6.1', '733.05', r['N'] / 1e3, 'kN')
    check('U.22', 'Pile fixity 2nd-order MRd,x', 'A19.6.1', '-393.69', r['Mx'] / 1e6, 'kN*m')
    check('U.23', 'Pile fixity 2nd-order MRd,y', 'A19.6.1', '79.72', r['My'] / 1e6, 'kN*m')
    check('U.24', 'Pile fixity Cc / Cs / T : Cc', 'A19.6.1', '1439.73', r['Cc'] / 1e3, 'kN')
    check('U.25', 'Pile fixity T', 'A19.6.1', '1269.29', r['T'] / 1e3, 'kN')
    check('U.26', 'Pile fixity eta (instability)', 'A19.6.1', '0.894', 655.55e3 / r['N'])
    e = forj.equilibrium([620.06e3, 353.51e6, -68.95e6])       # P264-P269
    check('U.27', 'Pile head design-load equilibrium Cc', 'A19.6.1', '1321.67', e['Cc'] / 1e3, 'kN')
    check('U.28', 'Pile head design-load equilibrium T', 'A19.6.1', '1126.46', e['T'] / 1e3, 'kN')
    check('U.29', 'Pile head design-load eps bar 7', 'A19.6.1', '-0.002733', e['bar_eps'][6])
    check('U.30', 'Pile head design-load eps_c,max', 'A19.6.1', '0.0026', e['eps_c_max'])
    e = emp.equilibrium([655.55e3, -352.07e6, 71.29e6])        # P319-P324
    check('U.31', 'Pile fixity design-load equilibrium Cc', 'A19.6.1', '1315.83', e['Cc'] / 1e3, 'kN')
    check('U.32', 'Pile fixity design-load equilibrium Cs', 'A19.6.1', '428.49', e['Cs'] / 1e3, 'kN')
    check('U.33', 'Pile fixity design-load eps bar 1', 'A19.6.1', '-0.002603', e['bar_eps'][0])
    # listing 'Aprov.' (appendix 3.5.1) = utilisation at constant N (not along the ray)
    etaN, _ = emp.capacity_constant_N([655.55e3, -352.07e6, 71.29e6])
    check('U.34', "Listing Aprov. P3 Cimentacion (constant-N def.)", 'CYPE listing 3.5.1', '90.7', etaN * 100, '%')
    etaN, _ = forj.capacity_constant_N([620.06e3, 353.51e6, -68.95e6])
    check('U.35', "Listing Aprov. P3 Cabeza (constant-N def.)", 'CYPE listing 3.5.1', '92.5', etaN * 100, '%',
          note='ray def. gives 91.3; NOT a force-rounding effect (listing forces 620.1/353.5/-69.0 give 92.31); '
               'Forjado-1 points run 0.2 pt low (see U.53), matched if bars sit at +-127.0 mm (92.53)')
    beam = RCSection(BEAM_RECT, BEAM_BARS, C35['fcd'], fyd, cell=1.0)
    r = beam.capacity_ray([0.0, -273.71e6, 0.0])               # P369-P385
    check('U.36', 'Beam P3-P4 hogging MRd,x (N=0)', 'A19.6.1', '-325.78', r['Mx'] / 1e6, 'kN*m')
    check('U.37', 'Beam Cc = T', 'A19.6.1', '767.00', r['Cc'] / 1e3, 'kN')
    check('U.38', 'Beam ecc,y', 'A19.6.1', '-219.16', r['ecc'][1], 'mm')
    check('U.39', 'Beam eT,y', 'A19.6.1', '205.59', r['eT'][1], 'mm')
    check('U.40', 'Beam eps top bars (=0.995*eps_su)', 'A19.6.1(3)', '-0.009950', r['bar_eps'][0])
    check('U.41', 'Beam eps_c,max', 'A19.6.1', '0.0017', r['eps_c_max'], note='figure: 1.66 permil')
    check('U.42', 'Beam sig_c,max', 'A19.3.1.7', '22.65', r['sig_c_max'], 'MPa')
    e0, kx, ky = r['plane']
    check('U.43', 'Beam neutral-axis depth x (figure P381)', 'A19.6.1', '68.52', -e0 / ky - YB, 'mm')
    check('U.44', 'Beam eta = MEd/MRd', 'A19.6.1', '0.840', 273.71 / abs(r['Mx'] / 1e6))
    e = beam.equilibrium([0.0, -273.71e6, 0.0])               # P388-P393
    check('U.45', 'Beam design-load Cc', 'A19.6.1', '543.32', e['Cc'] / 1e3, 'kN')
    check('U.46', 'Beam design-load Cs', 'A19.6.1', '93.94', e['Cs'] / 1e3, 'kN')
    check('U.47', 'Beam design-load T', 'A19.6.1', '637.26', e['T'] / 1e3, 'kN')
    check('U.48', 'Beam design-load sig_s top bars', 'A19.6.1', '-392.45', e['bar_sig'][0], 'MPa')
    check('U.49', 'Beam design-load sig_c,max', 'A19.3.1.7', '11.65', e['sig_c_max'], 'MPa')
    check('U.50', 'Beam design-load eps ledge O10 (bar 7)', 'A19.6.1', '-0.000662', e['bar_eps'][6])
    # required areas (appendix 2.1.3, Portico 4, tramo P3-P4)
    As = required_tension_steel(BEAM_RECT, [(x, y) for x, y, _ in BEAM_BOT], 294.61e6, C35['fcd'], fyd, +1)
    check('U.51', "Beam 'Area Inf. Nec.' Mmax=294.61 (singly r.)", 'A19.6.1', '15.10', As / 100, 'cm2')
    As = required_tension_steel(BEAM_RECT, [(x, y) for x, y, _ in BEAM_TOP], 273.71e6, C35['fcd'], fyd, -1)
    check('U.52', "Beam 'Area Sup. Nec.' M(P3)=-273.71", 'A19.6.1', '13.72', As / 100, 'cm2',
          note='listing zone shows -266.03; CYPE sizes with the node value')
    # third listing point (appendix 3.5.1, P3 'Forjado 1 ... Pie', G,Q,V: 655.6 / -352.4 / 71.7 -> 91.9 %)
    etaN, _ = forj.capacity_constant_N([655.6e3, -352.4e6, 71.7e6])
    check('U.53', "Listing Aprov. P3 Pie (constant-N def.)", 'CYPE listing 3.5.1', '91.9', etaN * 100, '%',
          note='same +0.2 pt offset as U.35 (Cimentacion U.34 exact); bars at +-127.0 give 91.92')
    # code-strict alternative to C1: eps_cu2 = 3.5 permil and NO steel strain limit (A19.3.2.7(2)b)
    strict = RCSection(BEAM_RECT, BEAM_BARS, C35['fcd'], fyd, cell=1.0, strain_factor=1.0, eps_su=1.0)
    r = strict.capacity_ray([0.0, -273.71e6, 0.0])
    check('U.54', 'Beam MRd code-strict (no 10 permil steel limit)', 'A19.3.2.7(2)b', None, r['Mx'] / 1e6, 'kN*m',
          note=f"{(r['Mx'] / -325.78e6 - 1) * 100:+.2f} % vs CYPE -325.78 (concrete reaches 3.5 permil); piles unchanged +0.09 %")
    print(f'   [ULS fibre analyses: {__import__("time").time() - t0:.1f} s]', file=sys.stderr)

    # -------------------------------------------------------------------------- S slenderness
    Ac, I = 400.0 * 400.0, 400.0 ** 4 / 12
    lam, ic = slenderness(6700.0, I, Ac)
    check('S.01', 'Radius of gyration ic', 'A19.5.8.3.2', '11.55', ic / 10, 'cm')
    check('S.02', 'Slenderness lambda = l0/i (l0 = 6.70 m)', 'A19.5.8.3.2 (5.14)', '58.02', lam)
    omega = 12 * A(25) * fyd / (Ac * C50['fcd'])
    check('S.03', 'omega = As fyd/(Ac fcd)', 'A19.5.8.3.1', '0.48', omega)
    phi_ef = 1.75
    for cid, N, ref in (('S.04', 620.06e3, '42.58'), ('S.05', 655.55e3, '41.42')):
        n = N / (Ac * C50['fcd'])
        ll, Af, Bf, Cf = lambda_lim(phi_ef, omega, n)
        check(cid, f'lambda_lim, NEd={N / 1e3:.2f} (n={n:.3f})', 'A19.5.8.3.1 (5.13)', ref, ll)
    check('S.06', 'A = 1/(1+0.2 phi_ef), phi_ef=1.75', 'A19.5.8.3.1', '0.74', Af)
    check('S.07', 'B = sqrt(1+2 omega) (UNE-EN form)', 'A19.5.8.3.1', '1.40', Bf, note='BOE erratum 1+sqrt()')
    check('S.08', 'C (rm unknown/unbraced)', 'A19.5.8.3.1', '0.70', Cf)
    check('S.09', 'n = NEd/(Ac fcd), fixity', 'A19.5.8.3.1', '0.12', 655.55e3 / (Ac * C50['fcd']))
    ei, th, ah, am = imperfection_ei(6700.0, 6.70)
    check('S.10', 'alpha_h = 2/sqrt(l), l = 6.70 m', 'A19.5.2(5)', '0.7727', ah)
    check('S.11', 'theta_i = theta0 ah am (theta0 = 1/200)', 'A19.5.2 (5.1)', '0.0039', th)
    check('S.12', 'ei = theta_i l0/2', 'A19.5.2 (5.2)', '12.94', ei, 'mm')
    check('S.13', 'e_min = max(h/30, 20)', 'A19.6.1(4)', '20.00', min_eccentricity(400.0), 'mm')
    check('S.14', 'e0,x head = MEd/NEd', 'A19.6.1(4)', '465.06', 288.36 / 620.06 * 1000, 'mm')
    for cid, d, N, ref in (('S.15', 327.5, 620.06e3, ('0.0148', '0.0203', '92.12')),
                           ('S.18', 329.5, 655.55e3, ('0.0147', '0.0201', '91.56'))):
        nc = nominal_curvature(fyd, Es, d, N / (Ac * C50['fcd']), omega, phi_ef, 50.0, lam, 6700.0)
        k = int(cid[2:])
        check(f'S.{k:02d}', f'1/r0 = eps_yd/(0.45 d), d={d}', 'A19.5.8.8.3 (5.34)', ref[0], nc['r0inv'] * 1000, '1/m')
        check(f'S.{k + 1:02d}', '1/r = Kr Kphi /r0', 'A19.5.8.8.3', ref[1], nc['rinv'] * 1000, '1/m')
        check(f'S.{k + 2:02d}', 'e2 = (1/r) l0^2 / c, c = pi^2', 'A19.5.8.8.2', ref[2], nc['e2'], 'mm')
    check('S.21', 'Kr = (nu-n)/(nu-nbal) <= 1', 'A19.5.8.8.3 (5.36)', '1.000', nc['Kr'], note=f"raw={(nc['n_u'] - 0.1229) / (nc['n_u'] - 0.4):.2f}")
    check('S.22', 'nu = 1 + omega', 'A19.5.8.8.3', '1.48', nc['n_u'])
    check('S.23', 'beta = 0.35 + fck/200 - lambda/150', 'A19.5.8.8.3', '0.213', nc['beta'])
    check('S.24', 'Kphi = 1 + beta phi_ef', 'A19.5.8.8.3 (5.37)', '1.373', nc['Kphi'])
    check('S.25', 'c = pi^2 (sinusoidal curvature)', 'A19.5.8.8.2(4)', '9.870', math.pi ** 2)
    etot = 432.56 + ei + nc['e2']
    check('S.26', 'etot = ee + ei + e2 (fixity, x)', 'A19.5.8.8.2 (5.31)', '537.06', etot, 'mm')
    check('S.27', 'MSd,x = N etot (fixity)', 'A19.5.8.8.2', '352.07', 655.55 * etot / 1000, 'kN*m')
    check('S.28', 'MSd,y = N (ee,y + ei + e2) (fixity)', 'A19.5.8.8.2', '71.29', 655.55 * (4.25 + ei + nc['e2']) / 1000, 'kN*m',
          note='CYPE adds ei and e2 in BOTH axes')
    nc5 = nominal_curvature(fyd, Es, 200 + math.sqrt(sum(y * y for _, y, _ in PILE_EMP) / 12), 0.1229, omega, phi_ef, 50, lam, 6700.0)
    check('S.29', 'e2 with d = h/2 + i_s (5.35, code-strict)', 'A19.5.8.8.3(2)', None, nc5['e2'], 'mm',
          note='CYPE uses d to extreme bar row (91.56); +6.9 %')

    # -------------------------------------------------------------------------- V shear
    nu1 = nu1_shear(50.0)
    check('V.01', 'nu1 = 0.6 (fck<=60, fywd=0.8 fywk)', 'A19.6.2.3(3) (6.10.a)', '0.600', nu1)
    check('V.02', 'sigma_cp (CYPE) = (NEd - As\' fyd)/Ac, dir X', 'A19.6.2.3(3)', '-12.13',
          sigma_cp_cype(620.06e3, 12 * A(25), fyd, Ac), 'MPa')
    check('V.03', 'sigma_cp (CYPE), dir Y, As\' = 4O25', 'A19.6.2.3(3)', '-1.46', sigma_cp_cype(620.06e3, 4 * A(25), fyd, Ac), 'MPa')
    check('V.04', 'sigma_cp (CYPE) fixity dir X', 'A19.6.2.3(3)', '-11.91', sigma_cp_cype(655.55e3, 12 * A(25), fyd, Ac), 'MPa')
    check('V.05', 'alpha_cw (sigma_cp <= 0 -> 1)', 'A19.6.2.3(3) (6.11)', '1.000', alpha_cw(-12.13, C50['fcd']))
    check('V.06', 'Pile VRd,max head (z = 249.16 CYPE)', 'A19.6.2.3(3) (6.9)', '996.63', VRd_max(400, 249.16, nu1, C50['fcd']) / 1e3, 'kN')
    check('V.07', 'Pile VRd,max fixity (z = 251.34 CYPE)', 'A19.6.2.3(3) (6.14)', '1005.35', VRd_max(400, 251.34, nu1, C50['fcd']) / 1e3, 'kN')
    check('V.08', 'Pile VRd,max with z = 0.9 d (d = 263.75)', 'A19.6.2.3(1)', None, VRd_max(400, 0.9 * 263.75, nu1, C50['fcd']) / 1e3, 'kN',
          note="CYPE's column z not reproduced (internal lever arm); -4.7 %, conservative")
    check('V.09', 'eta1 = sqrt((Vx/VRdmax)^2+(Vy/VRdmax)^2) head', 'CYPE biaxial', '0.086', math.hypot(0.98, 85.36) / 996.63)
    Asl = 8 * A(25)
    vc = VRd_c(50.0, 1.5, 400.0, 263.75, Asl, 332.81e3, Ac, C50['fcd'])
    check('V.10', 'Asl = 8 O25 (all bars except compressed row)', 'A19.6.2.2(1)', '39.27', Asl / 100, 'cm2')
    check('V.11', 'CRd,c = 0.18/gc', 'A19.6.2.2(1)', '0.120', vc['CRdc'])
    check('V.12', 'k = 1 + sqrt(200/d) <= 2, d = 263.75', 'A19.6.2.2(1)', '1.871', vc['k'])
    check('V.13', 'rho_l = Asl/(bw d) <= 0.02', 'A19.6.2.2(1)', '0.020', vc['rho_l'])
    check('V.14', 'sigma_cp = NEd/Ac <= 0.2 fcd', 'A19.6.2.2(1)', '2.08', vc['sigma_cp'], 'MPa')
    check('V.15', 'vmin = 0.035 k^1.5 fck^0.5', 'A19.6.2.2(1) (6.3)', '0.63', vc['vmin'], 'MPa')
    check('V.16', 'Pile VRd,c (CYPE label VRd,s)', 'A19.6.2.2(1) (6.2.a)', '142.85', vc['V1'] / 1e3, 'kN')
    check('V.17', 'Pile VRd,c minimum', 'A19.6.2.2(1) (6.2.b)', '99.73', vc['Vmin'] / 1e3, 'kN')
    check('V.18', 'eta2 = VEd/VRd,c (Vx=0.34, Vy=78.86)', 'A19.6.2.1(3)', '0.552', math.hypot(0.34, 78.86) / (vc['V'] / 1e3))
    # beam P3-P4
    d, bw, z = 480.0, 500.0, 0.9 * 480.0
    check('V.19', 'Beam z = 0.9 d', 'A19.6.2.3(1)', '432.00', z, 'mm')
    check('V.20', 'Beam sigma_cp (CYPE), As\' = 7O20', 'A19.6.2.3(3)', '-2.62', sigma_cp_cype(0.0, 7 * A(20), fyd, 365000.0), 'MPa')
    check('V.21', 'Beam VRd,max (theta = 45 deg)', 'A19.6.2.3(3) (6.9)', '1512.00', VRd_max(bw, z, 0.6, C35['fcd']) / 1e3, 'kN')
    fywd = 0.8 * 500.0
    check('V.22', 'fywd = 0.8 fywk (NOTA to 6.2.3(3))', 'A19.6.2.3(3) NOTA', '400.00', fywd, 'MPa')
    Asw = 3 * A(10)
    check('V.23', 'Asw = 3 legs O10', '-', '2.36', Asw / 100, 'cm2')
    check('V.24', 'Beam VRd,s = Asw/s z fywd cot(theta)', 'A19.6.2.3(3) (6.8)', '407.15', VRd_s(Asw, 100.0, z, fywd) / 1e3, 'kN')
    check('V.25', 'Beam eta VEd/VRd,max (VEd = 390.18)', 'A19.6.2', '0.258', 390.18 / 1512.0)
    check('V.26', 'Beam eta VEd/VRd,s', 'A19.6.2', '0.958', 390.18 / (VRd_s(Asw, 100.0, z, fywd) / 1e3))
    check('V.27', "'Area Transv. Nec.' = VEd/(z fywd)", 'A19.6.2.3(3)', '22.58', asw_required(390.18e3, z, fywd) * 1000 / 100, 'cm2/m')
    check('V.28', 'rho_w = Asw/(s bw sin a)', 'A19.9.2.2(5) (9.4)', '0.0047', Asw / (100.0 * bw))
    check('V.29', 'rho_w,min = 0.08 sqrt(fck)/fyk', 'A19.9.2.2(5) (9.5)', '0.0009', rho_w_min(35.0, 500.0))
    check('V.30', "Min. stirrups 50 cm web (listing 'Nec.')", 'A19.9.2.2(5)', '4.73', rho_w_min(35.0, 500.0) * 500 * 1000 / 100, 'cm2/m')
    check('V.31', "Min. stirrups 25x30 edge beam", 'A19.9.2.2(5)', '2.37', rho_w_min(35.0, 500.0) * 250 * 1000 / 100, 'cm2/m')
    check('V.32', 's_l,max = 0.75 d (1 + cot a)', 'A19.9.2.2(6) (9.6)', '360', s_l_max(d), 'mm')
    check('V.33', 's_t,max = 0.75 d <= 600', 'A19.9.2.2(8) (9.8)', '360', s_t_max(d), 'mm')

    # -------------------------------------------------------------------------- D detailing
    check('D.01', 'Clear spacing min, long. O25, dg 20', 'A19.8.2(2)', '25', clear_spacing_min(25, 20), 'mm',
          note='CYPE s2 = 1.25 dg = 25 = dg + 5')
    check('D.02', 'Clear spacing pile bars (85 - 25)', 'geometry', '60', 85 - 25, 'mm')
    check('D.03', 'Clear spacing beam top bars (82.5 - 20)', 'geometry', '63', 82.5 - 20, 'mm')
    check('D.04', 'phi_long >= 12 mm (pile O25)', 'A19.9.5.2(1)', '12', 12.0, 'mm')
    s, parts = s_cl_max(25.0, 400.0, 400.0)
    check('D.05', 's_cl,max = min(15 phi_min, 300, min(b,h))', 'A19.9.5.3(3)', '300', s, 'mm',
          note=f'terms {parts} (CYPE prints s1=300, s2=375, s3=400: s1/s2 labels swapped)')
    check('D.06', 'phi_t >= max(6, phi_max/4)', 'A19.9.5.3(1)', '6.3', phi_t_min(25.0), 'mm')
    cm = column_As_min(673.90e3, fyd, Ac)
    check('D.07', "A's,min = 0.10 NEd/fyd (NEd = 673.90)", 'A19.9.5.2(2) (9.12)', '1.55', cm['symmetric'] / 100, 'cm2')
    check('D.08', 'As,min = 0.004 Ac (AN/UNE-EN 1992-1-1)', 'AN 9.5.2(2)', '6.40', cm['geometric_AN'] / 100, 'cm2',
          note='not in the CE A19 text')
    check('D.09', "A's,min per face 0.05 NEd/fyc,d (fyc,d<=400)", 'A19.9.5.2(2)', None, cm['per_face'] / 100, 'cm2')
    cx = column_As_max(Ac, C50['fcd'], fyd)
    check('D.10', 'As,max = fcd Ac / fyd (AN / CYPE)', 'AN 9.5.2(3)', '122.67', cx['AN_total'] / 100, 'cm2')
    check('D.11', 'As,max = 0.04 Ac (CE A19 text) vs 58.91', 'A19.9.5.2(3)', None, cx['CE'] / 100, 'cm2',
          note='58.91 <= 64.00 also OK')
    W = 866680.7e4 / 305.82     # gross I (cm4 -> mm4) / distance to top fibre
    h_s = homogenised_stresses(BEAM_RECT, [], 1.0, 1.0)
    W = h_s['I'] / (YT - h_s['yc'])
    check('D.12', 'Beam W top fibre (gross)', 'A19.9.2.1.1', '28339.35', W / 1000, 'cm3')
    ffl = fctm_fl(C35['fctm'], 550)
    check('D.13', 'Beam As,min = W fctm,fl/(z fyd), z=0.9d=432', 'A19.9.2.1.1(1) (9.1)', '5.09', beam_As_min(W, 432.0, ffl, fyd) / 100, 'cm2')
    check('D.14', 'Beam As,min with z = 0.8 h (code text)', 'A19.9.2.1.1(1)', None, beam_As_min(W, 440.0, ffl, fyd) / 100, 'cm2',
          note='-1.8 % vs CYPE')
    check('D.15', 'Beam As,max = 0.04 Ac', 'A19.9.2.1.1(3)', None, 0.04 * 365000 / 100, 'cm2')
    check('D.16', 'Long. bar spacing <= 350 mm (bottom bars)', 'A19.9.2.3(4)', '150', 150.0, 'mm')

    # -------------------------------------------------------------------------- K SLS
    n_mod = Es / C35['Ecm']
    bars_x = [(x, y, a) for x, y, a in BEAM_TOP + BEAM_BOT]
    hs = homogenised_stresses(BEAM_RECT, bars_x, n_mod, 86e6)
    check('K.01', 'sigma_ct bottom, M_qp = +86 kN*m (P570)', 'A19.7.1(2)', '-2.20', hs['bot'], 'MPa', note=f'n = Es/Ecm = {n_mod:.2f}')
    check('K.02', 'sigma_c top, M_qp = +86', 'A19.7.1(2)', '2.73', hs['top'], 'MPa', status='INFO',
          note='figure strains imply Ec ~33.5 GPa and bars ~50 mm from the faces; no single linear homogenised '
               'section fits all 8 printed values; governing bottom value matches')
    check('K.03', 'sigma_s bottom bars, M_qp = +86', 'A19.7.1(2)', '-10.27', hs['bars'][-1], 'MPa')
    check('K.04', 'fct,eff = fctm used by engineer (3.2)', 'A19.7.1(2)', '3.2', C35['fctm'], 'MPa')
    Mcr_sag = C35['fctm'] * hs['I'] / (hs['yc'] - YB) / 1e6
    Mcr_hog = C35['fctm'] * hs['I'] / (YT - hs['yc']) / 1e6
    check('K.05', 'Mcr sagging (homogenised, fctm)', 'A19.7.1(2)', None, Mcr_sag, 'kN*m')
    check('K.06', 'Mcr hogging (homogenised, fctm)', 'A19.7.1(2)', None, Mcr_hog, 'kN*m')
    check('K.07', 'wmax XS3, RC, quasi-permanent comb.', 'CE Tabla 27.2', None, 0.1, 'mm')
    check('K.08', 'fT,lim = L/250, L = 3.05 m', 'CYPE (CTE/A19.7.4.1)', '12.20', 3050 / 250, 'mm')
    check('K.09', 'fA,lim = L/500', 'CYPE (CTE/A19.7.4.1)', '6.10', 3050 / 500, 'mm')
    check('K.10', 'Cantil tie: 4O20 x 400 MPa (P607)', 'hand check', '502', 4 * 314 * 400 / 1000, 'kN')
    check('K.11', 'Cantil tie force 410/4.65 (P607 prints 83)', 'hand check', '83', 410 / 4.65, 'kN', status='INFO',
          note='arithmetic slip in A10: 410/4.65 = 88.2 kN (still << 502)')
    # crack-width demonstration (no CYPE target): sagging at M_qp with psi2 = 0.8 (ROM 2.0-11)
    for M in (151.9e6,):
        cr = cracked_elastic(BEAM_RECT, bars_x, n_mod, M)
        sig_s = cr['sig_s'][-1]
        d_b = YT - (YB + 70)
        hcef = min(2.5 * (550 - d_b), (550 - cr['x']) / 3, 550 / 2)
        Aceff = 800 * hcef
        wk = crack_width(sig_s, Es, C35['Ecm'], C35['fctm'], 7 * A(20), Aceff, c=60.0, phi=20.0, kt=0.4)
        check('K.12', f'Cracked x (M = {M / 1e6:.1f}, sagging)', 'A19.7.3.4', None, cr['x'], 'mm')
        check('K.13', 'sigma_s bottom bars (cracked)', 'A19.7.3.4', None, sig_s, 'MPa')
        check('K.14', 'hc,ef = min(2.5(h-d), (h-x)/3, h/2)', 'A19.7.3.2(3)', None, hcef, 'mm')
        check('K.15', 'sr,max = k3 c + k1 k2 k4 phi/rho_p,eff', 'A19.7.3.4(3) (7.11)', None, wk['sr_max'], 'mm')
        check('K.16', 'eps_sm - eps_cm (kt = 0.4)', 'A19.7.3.4(2) (7.9)', None, wk['d_eps'] * 1000, 'permil')
        check('K.17', 'wk (limit 0.1 mm XS3)', 'A19.7.3.4 (7.8)', None, wk['wk'], 'mm')
    # smallest crack width once cracked: M = Mcr,sag (the (7.9) 0.6 sigma_s/Es floor governs)
    cr = cracked_elastic(BEAM_RECT, bars_x, n_mod, Mcr_sag * 1e6)
    wk_cr = crack_width(cr['sig_s'][-1], Es, C35['Ecm'], C35['fctm'], 7 * A(20),
                        800 * min(2.5 * (550 - (YT - (YB + 70))), (550 - cr['x']) / 3, 550 / 2), c=60.0, phi=20.0, kt=0.4)
    check('K.18', f'wk at M = Mcr,sag = {Mcr_sag:.1f} (first cracking)', 'A19.7.3.4 (7.8)', None, wk_cr['wk'], 'mm',
          note='> 0.1 mm: with 7O20 the XS3 limit can only be met by staying uncracked')

    # -------------------------------------------------------------------------- P psi2 sensitivity
    xlsx = '/home/user/MarshalSp/sap2000/resultados_sap/SAP27_Element_Forces_Frames.xlsx'
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        ws = wb['Element Forces - Frames']
        from collections import defaultdict
        M = defaultdict(dict)
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0] and row[0].startswith('VT') and row[2] in ('PP', 'CM', 'Qa'):
                M[(row[0], round(row[1], 4))][row[2]] = row[10]
        inner = [k for k in M if k[0].split('_')[0] in ('VT2', 'VT3', 'VT4', 'VT5', 'VT6')]   # VT1/VT7 = end frames (L)
        vt4 = [k for k in inner if k[0].startswith('VT4_')]                                    # VT4 = CYPE Portico 6
        for i, psi2 in enumerate((0.3, 0.5, 0.8)):
            mq = [M[k]['PP'] + M[k]['CM'] + psi2 * M[k]['Qa'] for k in inner if len(M[k]) == 3]
            msag, mhog = max(mq), min(mq)
            m4 = max(M[k]['PP'] + M[k]['CM'] + psi2 * M[k]['Qa'] for k in vt4 if len(M[k]) == 3)
            sig = homogenised_stresses(BEAM_RECT, bars_x, n_mod, msag * 1e6)['bot']
            check(f'P.{i + 1:02d}', f'Inner beams M_qp,max sag (psi2={psi2}) SAP', 'A18 6.5.3 (6.16b)', '86.66' if psi2 == 0.3 else None,
                  msag, 'kN*m', tol=0.06, status='INFO',
                  note=f'max of VT2-VT6 (VT4 = Portico 6: {m4:.2f}); hog {mhog:.1f}; sigma_ct = {-sig:.2f} MPa vs fctm {C35["fctm"]:.2f} -> '
                       f'{"UNCRACKED" if -sig <= C35["fctm"] else "CRACKED: check wk <= 0.1 mm"}')
        # psi2 at which the governing inner beam starts to crack (M_qp linear in psi2)
        k_gov = max((k for k in inner if len(M[k]) == 3), key=lambda k: M[k]['PP'] + M[k]['CM'] + 0.8 * M[k]['Qa'])
        m0, mq1 = M[k_gov]['PP'] + M[k_gov]['CM'], M[k_gov]['Qa']
        psi_cr = (Mcr_sag - m0) / mq1
        psi_cr_fl = (Mcr_sag * fctm_fl(C35['fctm'], 550) / C35['fctm'] - m0) / mq1
        check('P.06', 'psi2,Qa at which inner beams crack (SAP)', 'A19.7.1(2)', None, psi_cr, '-', status='INFO',
              note=f'{k_gov[0]} x={k_gov[1]}: M_qp = {m0:.1f} + {mq1:.1f} psi2; Mcr {Mcr_sag:.1f} (fctm); '
                   f'{psi_cr_fl:.2f} with fct,eff = fctm,fl')
        # piles: corner tension stress (gross section, biaxial) for psi2,Qa = 0.8 and psi2,TB in {0, 0.5}
        P = defaultdict(dict)
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row[0] and row[0].startswith('PIL') and row[2] in ('PP', 'CM', 'Qa', 'TB1'):
                P[(row[0], round(row[1], 3))][row[2]] = (row[5], row[9], row[10])
        for i, psi_tb in enumerate((0.0, 0.5)):
            worst = None
            for key, dd in P.items():
                if len(dd) < 4:
                    continue
                f = lambda j: dd['PP'][j] + dd['CM'][j] + 0.8 * dd['Qa'][j] + psi_tb * dd['TB1'][j]
                s_t = f(0) * 1e3 / 160000.0 + (abs(f(1)) + abs(f(2))) * 1e6 / (400.0 ** 3 / 6)   # SAP P<0 = compression
                if worst is None or s_t > worst[0]:
                    worst = (s_t, key)
            check(f'P.{i + 4:02d}', f'Piles max corner sigma_ct, psi2 Qa=0.8, TB={psi_tb}', 'A19.7.1(2)', None, worst[0], 'MPa',
                  status='INFO', note=f'{worst[1][0]} z={worst[1][1]} m; fctm(HA-50) = {C50["fctm"]:.2f} -> '
                  + ('UNCRACKED' if worst[0] <= C50['fctm'] else 'CRACKED: wk check needed (0.1 mm)'))
    except Exception as exc:   # pragma: no cover
        print(f'   [psi2 sensitivity skipped: {exc}]', file=sys.stderr)


if __name__ == '__main__':
    run()
    print_results()
    if '--json' in sys.argv:
        path = sys.argv[sys.argv.index('--json') + 1]
        with open(path, 'w') as f:
            json.dump(RESULTS, f, indent=1)
        print(f'JSON written to {path}')
    sys.exit(1 if any(r['status'] == 'FAIL' for r in RESULTS) else 0)
