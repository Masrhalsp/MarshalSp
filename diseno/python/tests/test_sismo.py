"""Tests of diseno/python/sismo (independent modal / response-spectrum analysis and accidental checks) and of
the seismic block of run_diseno.py (--sismo), plus a guard that the static outputs are unchanged.

Run:  python3 -m pytest diseno/python/tests/test_sismo.py -q      (~40 s)
"""

from __future__ import annotations

import hashlib
import json
import math
import sys
from pathlib import Path

import numpy as np
import pytest

DISENO = Path(__file__).resolve().parents[1]
for _p in (DISENO / "sismo", DISENO / "codigo", DISENO):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import comprobaciones as C  # noqa: E402
import modal as S  # noqa: E402

tm = S.tm
AC = tm.SEISMIC["ac_g"]

# sha256 of the static (non-seismic) outputs as delivered before the seismic work: they must stay
# byte-identical (the seismic results go to new files only).
STATIC_OUTPUTS = {
    "output/Diseno_Codigo_Estructural.xlsx": "e4b5d2a98e65cccf29d54557801fb59e69c160b8a9c9168a175c3bbd932c4500",
    "output/diseno_final.json": "cc7a214081bec2ff309cede343ec3cc2501689815946e43e5785304799c173ce",
    "output/diseno_final.md": "83d2e47b1d87545f1d3c7e58cb36630f28463bc80cd0345576d0619ee95a4a96",
    "output/pilotes.json": "3802e36aa023ef1010d3f331b2b190733d8fe86ae6f97b48ff0f98f4e0917d85",
    "output/pilotes.md": "4b25a4fd1abfe9f7e3b9b9c73b16826a8072cb225637a681f98594a602f9ae08",
    "output/vigas.json": "db2396a09959ec5b7614102d15a586cb29b7c52b42f6294d149826f47c098b81",
    "output/vigas.md": "33ddbeed83f9c61ba297417a7beebf74222117f9ce82a4eb7e124773148ea531",
    "figuras/d0_diagramas_pivotes.png": "2c0a2e469ef0915122556d222d8431d64e18c3f328e22f7ca7f0025a6fded172",
    "figuras/d1_pilotes_eta.png": "46c841cb0688476fec0b985b57e304d5e71a95fe794b32c555578fde25483cac",
    "figuras/d2_pilotes_As.png": "6774127af571adbd0f009949fc34d8205d1fb788f1d6198e762c2dc6bbc76e71",
    "figuras/d3_vigas_flexion.png": "7f914e219895a9405cd432609cc500c4a01a161bf769819a3862dda90e0a6277",
    "figuras/d4_fisuracion_psi2.png": "f7fb63708aa5c5dcc95d76fdc3c391e42be3560b8ce939ff976431d1dfa4df30",
    "figuras/d5_seccion_T_actual_propuesta.png": "0a0caba9d8879f0d35b6a68b4b7efaf234543124945a173cc6c519cc32975acf",
    "figuras/d6_seccion_pilote.png": "911516d3cbda1ede18d02466194c91c27e690f389eb604c164e061e56b06c6a0",
    "figuras/d7_paridad_cype.png": "10fd4816dcaea7f5827185edb666ff29a67cb376663d016e087355f6fdd8c8a4",
    "figuras/d8_interaccion_pilote.png": "341724389bdae3cee0606c2846a200272cf7c4b9cf42d565dcb6c153b3776416",
}


@pytest.fixture(scope="module")
def rs():
    st = S.build_structure()
    modes = S.solve_modes(st)
    return S.response_spectrum(st, modes, missing_mass=("Z",))


# --------------------------------------------------------------------------------------------
# spectrum
# --------------------------------------------------------------------------------------------
@pytest.mark.parametrize("T, alpha", [(0.0, 1.0), (0.11, 1.75), (0.22, 2.5), (0.5, 2.5), (0.88, 2.5),
                                      (1.76, 1.25), (3.52, 0.625)])
def test_spectrum_alpha(T, alpha):
    assert tm.seismic_alpha(T) == pytest.approx(alpha, rel=1e-9)
    assert S.spectral_acc(T, "H") == pytest.approx(AC * alpha * S.G, rel=1e-9)
    assert S.spectral_acc(T, "V") == pytest.approx(0.7 * AC * alpha * S.G, rel=1e-9)


def test_spectrum_parameters():
    assert AC == pytest.approx(0.15 * 1.91 * 0.85, rel=1e-6)          # 0.2435 g
    assert (tm.SEISMIC["TA"], tm.SEISMIC["TB"]) == (0.22, 0.88)
    plateau = max(v for _, v in tm.seismic_spectrum("H"))
    assert plateau == pytest.approx(2.5 * AC, rel=1e-5)                 # 0.609 (Rover)


# --------------------------------------------------------------------------------------------
# mass, modes, participation
# --------------------------------------------------------------------------------------------
def test_mass_equals_load_totals(rs):
    st = rs.st
    tot = tm.load_totals(st.model)
    W = sum(f * tot[p] for p, f in tm.SEISMIC["mass_source"]["patterns"].items())      # PP + CM + 0.8 Qa
    assert (st.mass_total + st.mass_base) * S.G == pytest.approx(W, rel=1e-9)
    for d in S.DIRS:                                                    # translational mass in X, Y and Z
        assert rs.modes.mtot[d] == pytest.approx(st.mass_total, rel=1e-9)


def test_mass_source_3720_kN(rs):
    """Mass source = loads PP (member self weight via SelfWtMult 1 + slab 4.10) + CM + 0.8·Qa; no element
    self mass on top (trasmallo SEISMIC['mass_source']['self_mass'] is False): 3720 kN = 379.4 t."""
    st = rs.st
    assert tm.SEISMIC["mass_source"]["self_mass"] is False
    assert (st.mass_total + st.mass_base) * S.G == pytest.approx(3720.2, abs=0.5)
    assert st.mass_total + st.mass_base == pytest.approx(379.35, abs=0.05)
    brk = st.info["weight_breakdown_kN"]
    tot = tm.load_totals(st.model)
    assert brk["self"] + brk["PP_area"] == pytest.approx(tot["PP"], rel=1e-9)      # self weight counted once
    assert brk["CM"] == pytest.approx(tot["CM"], rel=1e-9)
    assert brk["Qa"] == pytest.approx(0.8 * tot["Qa"], rel=1e-9)


def test_participation(rs):
    m = rs.modes
    assert len(m.omega) == tm.SEISMIC["n_modes"]
    assert np.all(np.diff(m.T) <= 1e-12)
    for d in S.DIRS:
        r = m.ratio(d)
        assert np.all(r >= 0) and r.sum() <= 1 + 1e-9
    assert m.ratio("X").sum() > 0.99 and m.ratio("Y").sum() > 0.99
    assert m.ratio("X")[0] > 0.99 and m.ratio("Y")[1] > 0.99            # mode 1 along X, mode 2 along Y
    assert m.ratio("Z").sum() < 0.5                                      # local slab modes: missing mass needed


def test_eigsh_vs_dense_guyan(rs):
    dense = S.dense_check(rs.st, 6)
    assert np.allclose(rs.modes.T[:6], dense, rtol=1e-5)


def test_eigen_residual(rs):
    st, m = rs.st, rs.modes
    for n in (0, 1, 2, 10):
        r = st.K @ m.phi[:, n] - m.omega[n] ** 2 * (st.M @ m.phi[:, n])
        assert np.linalg.norm(r) <= 1e-6 * np.linalg.norm(st.K @ m.phi[:, n])


# --------------------------------------------------------------------------------------------
# cross-checks: Rayleigh, hand estimate, equivalent static
# --------------------------------------------------------------------------------------------
def test_rayleigh_and_hand_period(rs):
    st, m = rs.st, rs.modes
    for d, n in (("X", 0), ("Y", 1)):
        k = S.static_push(st, d)
        assert 2 * math.pi * math.sqrt(st.mass_total / k) == pytest.approx(m.T[n], rel=0.01)
    h = S.hand_period(st)
    assert h["k_pile_fixed_guided"] == pytest.approx(12 * 33550e3 * 0.4 ** 4 / 12 / 6.70 ** 3, rel=1e-9)
    assert h["TX"] == pytest.approx(m.T[0], rel=0.20)                    # fixed-head lower bound of TX
    assert h["TX"] < m.T[0]
    assert h["TY"] == pytest.approx(m.T[1], rel=0.05)


def test_equivalent_static_base_shear(rs):
    for i, d in enumerate(("X", "Y")):
        e = S.equivalent_static(rs.st, rs.modes, d)
        assert e["Sa_g"] == pytest.approx(2.5 * AC, rel=1e-9)            # T1 on the plateau
        assert rs.base[d][i] == pytest.approx(e["V"], rel=0.02)
        assert rs.base[d][i] == pytest.approx(math.sqrt(sum((rs.modes.meff[d] * rs.sa[d]) ** 2)), rel=1e-5)


def test_missing_mass_restores_vertical_shear(rs):
    st = rs.st
    V0 = st.mass_total * S.spectral_acc(0.0, "V")                        # rigid response bound
    assert 0.8 * V0 < rs.base["Z"][2] < 2.0 * V0


# --------------------------------------------------------------------------------------------
# member-force recovery vs PyNite's own analysis
# --------------------------------------------------------------------------------------------
def test_recovery_matches_pynite():
    st = S.build_structure()
    fe, model = st.fe, st.model
    fe.add_node_load("D10_5", "FY", -50.0, case="CHK")
    fe.add_node_load("D10_5", "FX", 20.0, case="CHK")
    fe.add_node_load("D30_12", "FZ", -15.0, case="CHK")
    fe.add_load_combo("CHK", {"CHK": 1.0})
    fe.analyze_linear(check_stability=False, log=False)
    D = np.zeros(6 * len(fe.nodes))
    for nd in fe.nodes.values():
        D[6 * nd.ID:6 * nd.ID + 6] = [nd.DX["CHK"], nd.DY["CHK"], nd.DZ["CHK"], nd.RX["CHK"], nd.RY["CHK"], nd.RZ["CHK"]]
    rec = S.Recovery(st)
    for p in ("P1", "P3", "P8"):
        ref = S.vp.pile_head_forces(fe, model, "CHK", p)
        ours = rec.pile(D, p, [6.70])[0]
        assert np.allclose(ours, [ref[k] for k in S.PILE_COMPS], atol=1e-6)
    for y, M, V, seg in S.vp.beam_line(fe, model, 2, "CHK"):
        frame = seg[:-3] if seg.endswith(("_ZI", "_ZJ")) else seg
        o = rec.beam_frame(D, frame, y)
        assert o[0] == pytest.approx(M, abs=1e-6)


def test_directional_rule():
    env = {"X": np.array([10.0]), "Y": np.array([20.0]), "Z": np.array([5.0])}
    assert S.directional(env, "X")[0] == pytest.approx(10 + 0.3 * 20 + 0.3 * 5)
    assert S.directional(env, "Y")[0] == pytest.approx(0.3 * 10 + 20 + 0.3 * 5)


# --------------------------------------------------------------------------------------------
# combinations and sign corners
# --------------------------------------------------------------------------------------------
def test_sign_combinations():
    for kind, comps in (("nm", (0, 1, 2)), ("v", (0, 3, 4))):
        cmb = C.pile_combos(kind)
        assert len(cmb) == 6 * 8
        for name, fac in cmb.items():
            static, seis = fac[:6], np.array(fac[6:]).reshape(3, 5)
            assert static[:2] == (1.0, 1.0) and static[2] in (0.0, 0.8) and static[3:] == (0.0, 0.0, 0.0)
            assert (static[2] == 0.8) == name.split("[")[0].endswith("Q")
            lead = "XYZ".index(name[3])
            for i in range(3):
                mag = np.abs(seis[i, list(comps)])
                assert np.allclose(mag, 1.0 if i == lead else 0.3)
                assert np.all(seis[i, [c for c in range(5) if c not in comps]] == 0)
            signs = np.sign(seis[lead, list(comps)])
            assert np.all(np.sign(seis[:, list(comps)]) == signs)       # same corner in every direction
        corners = {tuple(np.sign(np.array(f[6:11])[list(comps)])) for n, f in cmb.items() if n.startswith("SISXQ[")}
        assert len(corners) == 8


def test_pile_forces_ext_combination():
    z = np.array([0.0, 6.70])
    vals = np.zeros((2, 6, 5))
    vals[:, 0, 0] = 100.0                                               # PP: N = 100
    static = {"P1": C.pilotes.PileForces("P1", z, vals)}
    env = {"P1": {"z": z, **{d: np.tile([10.0, 20.0, 30.0, 4.0, 5.0, 0.0], (2, 1)) * (k + 1)
                              for k, d in enumerate(S.DIRS)}}}
    pf = C.pile_forces_ext(static, env)["P1"]
    fac = C.pile_combos("nm")["SISYQ[N+,Mx-,My+]"]
    r = pf.combine(fac)[0]
    # section axes: (N, MEd,x = My, MEd,y = Mx); Y leads: 1.0·2E + 0.3·(1E + 3E) = 3.2·E
    assert r[0] == pytest.approx(100 + 3.2 * 10)
    assert r[1] == pytest.approx(-3.2 * 30)
    assert r[2] == pytest.approx(3.2 * 20)
    assert r[3] == 0 and r[4] == 0


def test_accidental_materials():
    lay = C.pile_layouts_acc()["fuste"]
    assert lay.concrete.fcd == pytest.approx(50 / 1.3)
    assert lay.steel.fyd == pytest.approx(500.0)


# --------------------------------------------------------------------------------------------
# outputs
# --------------------------------------------------------------------------------------------
def test_static_outputs_unchanged():
    for rel, h in STATIC_OUTPUTS.items():
        assert hashlib.sha256((DISENO / rel).read_bytes()).hexdigest() == h, rel


@pytest.mark.skipif(not (DISENO / "output" / "sismo.json").exists(), reason="run diseno/python/sismo/run_sismo.py first")
def test_sismo_json_and_run_diseno_block(tmp_path):
    import run_diseno as R
    J = json.loads((DISENO / "output" / "sismo.json").read_text(encoding="utf-8"))
    s = J["summary"]
    assert s["governs"] is True and s["holds"] is False
    assert s["pile_eta_nm"] > s["pile_eta_nm_static"] and s["beam_eta_M"] > s["beam_eta_M_static"]
    assert len(J["modes"]) == 30 and len(J["piles"]) == 14
    out, fig = tmp_path / "out", tmp_path / "fig"
    assert R.main(["--sismo", "--reuse", "--out", str(out), "--fig-dir", str(fig)]) == 0
    from openpyxl import load_workbook
    wb = load_workbook(out / R.SISMO_XLSX.name)
    assert wb.sheetnames == ["Sismo"] and len(wb["Sismo"]._charts) == 1
    assert (fig / R.SISMO_FIG).read_bytes()[:4] == b"\x89PNG"
    assert sorted(p.name for p in out.iterdir()) == [R.SISMO_XLSX.name]   # no static output written
