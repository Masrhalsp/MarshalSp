"""Quick test of diseno/python/run_diseno.py: runs it with --reuse (the committed diseno/python/output/pilotes.json and
vigas.json) into a temporary directory and checks the workbook sheets / charts, the figures, the
final-design summary and its consistency with the final-design decisions F0-F8.

Run:  python3 -m pytest diseno/python/tests/test_run_diseno.py -q      (~20 s)
"""

from __future__ import annotations

import json
import struct
import sys
from pathlib import Path

import pytest

DISENO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(DISENO))

import run_diseno as R  # noqa: E402

pytestmark = pytest.mark.skipif(not ((R.OUT / "pilotes.json").exists() and (R.OUT / "vigas.json").exists()),
                                reason="diseno/python/output/pilotes.json / vigas.json not present (run the checks first)")

CHART_SHEETS = ("Pilotes_ELU", "Pilotes_As", "Pilotes_ELS", "Vigas_Flexion", "Vigas_Cortante", "Vigas_Fisuracion",
                "Comparacion_CYPE")


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    fig = tmp_path_factory.mktemp("fig")
    assert R.main(["--reuse", "--out", str(out), "--fig-dir", str(fig)]) == 0
    return out, fig


def test_workbook_sheets_and_charts(run):
    from openpyxl import load_workbook
    out, _ = run
    wb = load_workbook(out / R.XLSX.name)
    assert tuple(wb.sheetnames) == R.SHEETS
    for name in CHART_SHEETS:
        assert len(wb[name]._charts) >= 1, name
    ws = wb["Resumen"]
    assert ws["A5"].value.startswith("P1")
    verdicts = [ws.cell(i, 12).value for i in range(5, 5 + len(R.PILES) + len(R.BEAMS))]
    assert all(v and v.startswith("CUMPLE") for v in verdicts), verdicts


def test_figures_exist(run):
    _, fig = run
    for name in R.FIGURES:
        p = fig / name
        assert p.exists() and p.stat().st_size > 20_000, name
        with p.open("rb") as f:
            head = f.read(24)
        assert head[:8] == b"\x89PNG\r\n\x1a\n"
        w, h = struct.unpack(">II", head[16:24])
        assert w >= 800 and h >= 500, (name, w, h)


def test_final_design_summary(run):
    out, _ = run
    md = (out / R.MD.name).read_text()
    for s in ("8Ø20 alma + 4Ø16 alas", "c/100 en los 0.40 m superiores", "3Ø25 sobre apoyos", "4Ø25",
              "Puntos abiertos", "ψ2 = 0"):
        assert s in md, s
    D = json.loads((out / R.JSON_OUT.name).read_text())
    assert D["warnings"] == []
    assert [s["verdict"].startswith("CUMPLE") for s in D["schedule"]] == [True] * 4
    pk = D["pile_key"]
    assert pk["nm2_max"]["CE-ROM"]["eta"] == pytest.approx(0.987, abs=1e-3)       # F1
    assert pk["nm2_max"]["CE-ROM"]["pile"] == "P3"
    assert pk["nm2_max"]["CYPE"]["eta"] == pytest.approx(0.920, abs=1e-3)
    assert sorted(pk["sigma_c_fail"]) == ["P1", "P16", "P2", "P3"]
    assert pk["As_req_max"]["As_req_cm2"] == pytest.approx(58.43, abs=0.01)
    inner = D["proposal"]["inner"]["bottom"]
    assert inner["proposed"] == "8Ø20 alma + 4Ø16 alas" and inner["wk_mm"] < 0.1   # F3
    cur = D["curves"]
    assert cur["threshold_provided"] == pytest.approx(0.583, abs=2e-3)
    assert cur["threshold_proposed"] > 0.8
    assert max(cur["layouts"]["proposed"]["wk_mm"][:41]) < 0.1                      # ψ2 <= 0.8
    end = {b["member"]: b for b in D["beams"] if b["group"] == "end"}
    assert end["VT1"]["eta_final"] == pytest.approx(0.973, abs=1e-3)               # F4
    assert D["proposal"]["edge"]["top"]["proposed"] == "3Ø25"                       # F5


def test_interaction_and_confinement(run):
    out, _ = run
    D = json.loads((out / R.JSON_OUT.name).read_text())
    inter = D["interaction"]
    p2 = inter["points"][-1]
    assert p2["check"] == "nm2" and p2["eta"] == pytest.approx(0.987, abs=1e-3)
    # the meridian passes through the capacity point of the governing 2nd-order ray
    import numpy as np
    m = np.array(inter["meridian"])
    assert np.interp(p2["NRd"], m[:, 0], m[:, 1]) == pytest.approx(p2["MRd"], rel=0.01)
    c = {x["s_mm"]: x for x in D["confinement"]}
    assert c[100.0]["sigma2_MPa"] > c[150.0]["sigma2_MPa"]
    assert c[100.0]["limit_0.6fckc"] > 30.0 * D["pile_key"]["sigma_c_max"]
