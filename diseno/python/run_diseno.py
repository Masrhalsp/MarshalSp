"""Final design of one 40 m module of the Muelle de Trasmallo (Port of Cullera) to the Código
Estructural (CE 2021, Anejo 19 = EN 1992-1-1 with the Spanish values): one command that runs the
pile and beam checks and writes the design workbook, the final-design summary and the figures.

    python3 diseno/python/run_diseno.py            # runs pilotes.run() and vigas.run() (~90 s), then writes
    python3 diseno/python/run_diseno.py --reuse    # reuses diseno/python/output/pilotes.json and vigas.json (~20 s)
    python3 diseno/python/run_diseno.py --reuse --no-figs

Outputs
    diseno/python/output/Diseno_Codigo_Estructural.xlsx   workbook with native Excel charts (sheets Resumen,
                                                   Pilotes_ELU, Pilotes_As, Pilotes_ELS, Vigas_Flexion,
                                                   Vigas_Cortante, Vigas_Fisuracion, Propuesta,
                                                   Comparacion_CYPE, Supuestos)
    diseno/python/output/diseno_final.md                  final design: reinforcement schedule + open points
    diseno/python/output/diseno_final.json                the data behind the two (incl. the ψ2 curves)
    diseno/python/figuras/d1..d8_*.png      figures for the reports (150 dpi)

Nothing is recomputed here except what the figures and the ψ2 sensitivity need (the ψ2,Qa curves of
the inner beams with the provided and the proposed bottom steel, the N-M interaction curve of the
pile and an indicative confinement estimate); every verdict number comes from pilotes.json /
vigas.json.  The final-design decisions F0-F8 of the design review are stated in DECISIONS and
applied in ``collect()``: final verdict = Mode.CODIGO + ROM basis + XS3 wmax 0.1 mm (quasi-permanent).
"""

from __future__ import annotations

import argparse
import json
import math
import sys
import time
from pathlib import Path

import numpy as np

DISENO = Path(__file__).resolve().parent
CODIGO = DISENO / "codigo"
for _p in (str(CODIGO), str(DISENO)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

OUT = DISENO / "output"
FIG = DISENO / "figuras"
XLSX = OUT / "Diseno_Codigo_Estructural.xlsx"
MD = OUT / "diseno_final.md"
JSON_OUT = OUT / "diseno_final.json"

SHEETS = ("Resumen", "Pilotes_ELU", "Pilotes_As", "Pilotes_ELS", "Vigas_Flexion", "Vigas_Cortante",
          "Vigas_Fisuracion", "Propuesta", "Comparacion_CYPE", "Supuestos")
FIGURES = ("d1_pilotes_eta.png", "d2_pilotes_As.png", "d3_vigas_flexion.png", "d4_fisuracion_psi2.png",
           "d5_seccion_T_actual_propuesta.png", "d6_seccion_pilote.png", "d7_paridad_cype.png",
           "d8_interaccion_pilote.png")

# ---------------------------------------------------------------------------------------------
# members
# ---------------------------------------------------------------------------------------------
SEA = ("P1", "P3", "P5", "P7", "P13", "P15", "P16")          # sea piles, axes 1..7
LAND = ("P2", "P4", "P6", "P8", "P14", "P17", "P18")         # land piles, axes 1..7
AXIS_OF = {p: i + 1 for i, p in enumerate(SEA)} | {p: i + 1 for i, p in enumerate(LAND)}
PILES = ("P1", "P2", "P3", "P4", "P5", "P6", "P7", "P8", "P13", "P14", "P15", "P16", "P17", "P18")
BEAMS = ("VT1", "VT2", "VT3", "VT4", "VT5", "VT6", "VT7", "VBM", "VBT")
INNER = ("VT2", "VT3", "VT4", "VT5", "VT6")
BEAM_NAME = {"VT1": "Eje 1 (pórtico 3)", "VT2": "Eje 2 (pórtico 4)", "VT3": "Eje 3 (pórtico 5)",
             "VT4": "Eje 4 (pórtico 6)", "VT5": "Eje 5 (pórtico 7)", "VT6": "Eje 6 (pórtico 8)",
             "VT7": "Eje 7 (pórtico 9)", "VBM": "Borde mar (pórtico 1)", "VBT": "Borde tierra (pórtico 10)"}
BEAM_SHORT = {"VT1": "Eje 1", "VT2": "Eje 2", "VT3": "Eje 3", "VT4": "Eje 4", "VT5": "Eje 5", "VT6": "Eje 6",
              "VT7": "Eje 7", "VBM": "Borde\nmar", "VBT": "Borde\ntierra"}
GROUP_KEY = {"inner": "inner (ejes 2-6, pórticos 4-8)", "end": "end (ejes 1 y 7, pórticos 3 y 9)",
             "edge": "edge (pórticos 1 y 10)"}
LEDGE_SENS_KEY = GROUP_KEY["end"] + " + torsión del ala (sensibilidad)"

# (basis, mode) variants of the piles and cases of the beams, in the order CYPE parity -> final
VARIANTS = (("CYPE", "cype", "CYPE"), ("CYPE", "codigo", "CE-CYPE"), ("ROM", "codigo", "CE-ROM"))
VARIANT_LABEL = {"CYPE": "CYPE/cype (paridad)", "CE-CYPE": "CYPE/codigo", "CE-ROM": "ROM/codigo (veredicto)"}
QP_CASE = {"CYPE": "QP ψ2,Qa=0.3, ψ2,TB=0", "ROM": "QP ψ2,Qa=0.8, ψ2,TB=0"}

PARITY_EXACT = "datos de CYPE (reproducción exacta)"
PARITY_SAP = "fuerzas SAP frente a resultados CYPE (≤ 5 %)"
PILE_TIE_ZONE_M = 0.40          # F1: top 0.40 m of every pile (A19.9.5.3(4): larger section dimension)
PILE_TIE_S_NEW = 100.0          # mm

# F0-F8: final-design decisions of the design review (stated identically in the workbook and the md)
DECISIONS = [
    ("F0", "Bases de cálculo / design bases",
     "'CYPE' = hipótesis del Anejo 10 (ψ0,Qa 0.7; ψ2,Qa 0.3, valor de edificación del CTE; tiro de bolardo "
     "tratado como viento, ψ2 0) para la paridad; 'ROM' = recomendada para un puerto: ROM 2.0-11 Tabla 4.6.4.1 "
     "(ψ0,Qa 1.0; ψ2,Qa 0.8; bolardo ψ2 0, con 0.5 como sensibilidad). Veredicto final = Mode.CODIGO (sección "
     "estricta al Código) + base ROM + XS3 wmax 0.1 mm (CE Art. 27 Tabla 27.2) en combinación cuasipermanente."),
    ("F1", "Pilotes 40x40 HA-50 (12Ø25 + 3 cercos Ø10/150)",
     "ELU cumple en todas las bases (η N-M 2º orden máx 0.987 en cabeza de P3, ROM; 0.920 paridad CYPE frente a "
     "0.913 de CYPE), cortante 0.814, disposiciones (As 58.9 ≤ 0.04Ac = 64 cm²), sin fisurar en cuasipermanente "
     "(bolardo ψ2 0). σc ≤ 0.6fck (A19.7.2(2), XS) superada en cabeza de P1, P2, P3, P16 con Ecm para toda la "
     "carga característica (máx 1.109); con Ec,eff para la parte sostenida todas cumplen (máx 0.966). A19.7.2(2) "
     "admite 'otras medidas' como el confinamiento con armadura transversal -> FINAL: se mantiene 12Ø25 y se "
     "reduce la separación de los cercos a Ø10 c/100 en los 0.40 m superiores de cada pilote (zona extrema, "
     "A19.9.5.3(4)); veredicto CUMPLE con esa medida. As,nec = 58.43 cm² (cabeza P3, ROM) frente a 58.91 dispuesta (0.99): "
     "12Ø25 es necesaria y suficiente; los pilotes de tierra necesitarían ~39-46 cm² (16Ø20) pero se mantiene "
     "un pilote prefabricado uniforme."),
    ("F2", "Tiro de bolardo ψ2 = 0",
     "Se mantiene ψ2 = 0 como en el proyecto original (el tiro de amarre no es cuasipermanente en un muelle "
     "pesquero); con 0.5 fisurarían seis pilotes (wk hasta 0.153 mm > 0.1) y ninguna armadura dentro de As,max "
     "lo resuelve -> documentado como hipótesis."),
    ("F3", "Vigas transversales interiores (ejes 2-6, T invertida)",
     "ELU flexión (máx 0.88), cortante incl. armadura de suspensión de las alas (máx 0.72 con cot θ óptimo), "
     "disposiciones y flecha cumplen. La fisuración NO cumple con ψ2 0.8 de la ROM (M_qp 137-152 kN·m > Mcr "
     "125 kN·m; wk 0.155-0.172 mm > 0.1 mm; el 7Ø20 inferior dispuesto sólo cumple para ψ2,Qa ≤ 0.58-0.70). "
     "FINAL: armadura inferior 8Ø20 en el alma + 4Ø16 en las alas (33.18 cm², wk 0.096 mm) en lugar de 5Ø20 + "
     "2Ø20 (21.99 cm²); superior 5Ø20 y estribos (1 cerco Ø10 + 1 rama Ø10 + cerco de alas, 3 ramas Ø10 c/100) "
     "sin cambio. Con ψ2 0.3 de CYPE el armado dispuesto cumple (como obtuvo CYPE)."),
    ("F4", "Vigas extremas (ejes 1 y 7, L invertida 80x55+15x30)",
     "CUMPLE con el armado dispuesto (determinante: cordón traccionado por torsión 0.973, flexión 0.873). "
     "Sensibilidad: si las alveoplacas apoyan en la única ala a ~0.42 m del eje del pilote (no modelado ni en SAP "
     "ni en CYPE: 'la alveoplaca apoya en el eje de los pilotes'), la torsión de equilibrio ~56-61 kN·m en las "
     "caras de los pilotes lleva el cordón a 1.18 -> verificar el detalle de apoyo; si se confirma, superior "
     "4Ø25 en lugar de 5Ø20."),
    ("F5", "Vigas de borde 25x30 (2Ø16 sup./inf., Ø8/150)",
     "En nuestro modelo SAP comparten la carga del forjado con la lámina ortótropa (V ~2x y M +13-21 % frente a "
     "CYPE) -> la fisuración no cumple (wk 0.31 mm en apoyos, momento negativo) y el cortante con θ = 45° es "
     "1.04-1.08 (cumple con cot θ óptimo). FINAL (condicionado al reparto de SAP): superior 3Ø25 sobre los "
     "apoyos (wk 0.067 mm); alternativa: viga de borde más canto; si sólo aplica la carga por bandas de CYPE, "
     "las barras dispuestas bastan en ELU."),
    ("F6", "Decalaje / shift rule (información)",
     "Con cot θ = 2, la regla de decalaje (A19.6.2.3(7)/9.2.1.3) en la cara mar del pilote limitada al momento "
     "del nudo (eje del pilote) elevaría el cordón superior a 1.03-1.18; recomendación: no cortar las barras "
     "superiores dentro de al = z·cot θ/2 ≈ 0.43 m de las caras de los pilotes y anclarlas por completo sobre "
     "los pilotes; con cot θ = 1 (CYPE) el efecto es menor. No forma parte del veredicto."),
    ("F7", "Comprobación EC2 propia de SAP2000",
     "La ratio de los pilotes en SAP será ~0.71 (SAP usa M0e = 0.6M02 + 0.4M01 ≥ 0.4M02 y espera un análisis "
     "P-Delta): no comparable con nuestro 2º orden 0.93-0.99; rige nuestro veredicto. SAP no comprueba wk XS3, "
     "la convención ν1/fywd del CE, As,min española ni la armadura de suspensión de las alas."),
    ("F8", "Anejo 10: qué hizo CYPE bien / mal",
     "Resultados de CYPE reproducidos exactamente (capacidades 0.00 %); su veredicto de fisuración depende de "
     "ψ2 0.3 (valor de edificación); no comprobó la armadura de suspensión de las alas, el límite de tensión "
     "característica ni la torsión del ala de las vigas extremas; el 'Aprov.' del listado = 1.014 × η del rayo "
     "(sin explicar); error aritmético en el texto del Anejo 10: 410/4.65 = 88 kN (no 83)."),
]

# palette (dataviz reference palette; validated: first three slots all-pairs, light surface)
C_ROM, C_CECY, C_CYPE = "#2a78d6", "#eb6834", "#1baf7a"     # slot 1 blue, 2 orange, 3 aqua
C_LIMIT = "#d03b3b"                                          # status critical: limits
INK, INK2, MUTED, GRID, AXISC = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
CONCRETE = "#f0efec"
CASE_COLOR = {"CYPE": C_CYPE, "CE-CYPE": C_CECY, "CE-ROM": C_ROM}


# =============================================================================================
# inputs
# =============================================================================================
def load_results(reuse: bool, out: Path = OUT, verbose: bool = True) -> tuple[dict, dict]:
    """pilotes.json / vigas.json: with ``reuse`` read from diseno/python/output (if present), otherwise
    recomputed and written to ``out`` exactly as ``python3 diseno/python/codigo/pilotes.py`` / ``vigas.py`` do."""
    out.mkdir(parents=True, exist_ok=True)
    if reuse and (OUT / "pilotes.json").exists():
        P = json.loads((OUT / "pilotes.json").read_text(encoding="utf-8"))
    else:
        import pilotes
        from seccion import Mode
        P = pilotes.run("ROM", Mode.CODIGO, verbose=verbose)
        (out / "pilotes.json").write_text(json.dumps(P, indent=1, ensure_ascii=False), encoding="utf-8")
        (out / "pilotes.md").write_text(pilotes.to_markdown(P), encoding="utf-8")
    if reuse and (OUT / "vigas.json").exists():
        V = json.loads((OUT / "vigas.json").read_text(encoding="utf-8"))
    else:
        import vigas
        V = vigas.run(verbose=verbose)
        (out / "vigas.json").write_text(json.dumps(V, indent=1, ensure_ascii=False), encoding="utf-8")
        vigas.write_md(V, out / "vigas.md")
    return P, V


def _r(x, nd=3):
    return None if x is None else round(float(x), nd)


def _max(*xs):
    v = [x for x in xs if x is not None]
    return max(v) if v else None


# =============================================================================================
# computations needed by the figures / sensitivity (not verdict numbers)
# =============================================================================================
def _beam_from_layout(rows: list[dict], member_key: str):
    """BeamReinforcement of a proposal layout (vigas.json 'layout' rows), on the geometry of the
    provided reinforcement of ``member_key``."""
    import armado_vigas as AV
    groups, st = [], None
    for r in rows:
        if r["face"] == "stirrups":
            st = AV.Stirrups(float(r["phi"]), int(r["n"]), float(r["s_mm"]), "proposal")
        else:
            groups.append(AV.BarGroup(r["group"], float(r["phi"]), tuple(float(x) for x in r["x_mm"]),
                                      float(r["y_soffit_mm"]), r["face"], bool(r.get("active_cype", True)),
                                      bool(r.get("active_codigo", True)), r.get("src", "")))
    base = AV.provided(int(member_key[2:])) if member_key.startswith("VT") else None
    return base, groups, st


def psi2_curves(V: dict, step: float = 0.02) -> dict:
    """σct/fctm (homogenised section) and wk of the inner beams (axes 2-6, sagging) as functions of
    ψ2,Qa (bollard ψ2 = 0), with the provided bottom steel (7Ø20) and the proposed one (8Ø20 + 4Ø16):
    envelope over the five beams, per-member values at ψ2 0.3 / 0.8 and the thresholds."""
    import vigas as VG
    from seccion import KNM
    members, reinf = VG.build_members()
    lay = V["proposal"][GROUP_KEY["inner"]]["layout"]
    _, groups, st = _beam_from_layout(lay, "VT2")
    psis = [round(x, 4) for x in np.arange(0.0, 1.0 + 1e-9, step)]
    out = {"psi": psis, "layouts": {}, "members": {}}
    for tag in ("provided", "proposed"):
        env_s, env_w = np.zeros(len(psis)), np.zeros(len(psis))
        for m in members:
            if m.key not in INNER:
                continue
            r0 = reinf[m.key]
            r = r0 if tag == "provided" else VG._apply_layout(
                r0, r0.__class__(r0.member, r0.label, r0.geom, list(groups), st, notes=r0.notes))
            a, b = VG.region_limits(m)
            vals = VG.member_values(m, a, b)

            def Mq(psi):
                return max(VG.E.combine(v, (1.0, 1.0, psi, 0.0, 0.0, 0.0), ("M",))["M"] for _, v in vals)
            s_list, w_list = [], []
            for i, psi in enumerate(psis):
                ck = VG.crack_check(r, Mq(psi) * KNM)
                s_list.append(ck["sigma_ct"] / ck["fctm"])
                w_list.append(ck["wk"])
            env_s = np.maximum(env_s, s_list)
            env_w = np.maximum(env_w, w_list)
            thr = VG.psi2_threshold(m, r, +1)
            at = {}
            for psi in (0.3, 0.8):
                ck = VG.crack_check(r, Mq(psi) * KNM)
                at[str(psi)] = {"M_qp": _r(Mq(psi), 2), "sigma_ct_fctm": _r(ck["sigma_ct"] / ck["fctm"]),
                                "wk_mm": _r(ck["wk"], 4), "cracked": bool(ck["cracked"])}
            out["members"].setdefault(m.key, {})[tag] = {
                "As_bottom_cm2": _r(r.As("bottom") / 100, 2), "layout_bottom": r.describe("bottom"),
                "Mcr_kNm": _r(VG.mcr(r, True) / KNM, 2), "psi2_crack": thr["psi2_crack_fctm"],
                "psi2_wk_fail": thr["psi2_wk_fail"], "at": at}
        out["layouts"][tag] = {"sigma_ct_fctm": [_r(x, 4) for x in env_s], "wk_mm": [_r(x, 4) for x in env_w]}
    thr_prov = [out["members"][k]["provided"]["psi2_wk_fail"] for k in INNER]
    thr_prop = [out["members"][k]["proposed"]["psi2_wk_fail"] for k in INNER]
    out["threshold_provided"] = min(x for x in thr_prov if x is not None)
    out["threshold_proposed"] = min((x for x in thr_prop if x is not None), default=None)
    return out


def pile_interaction(P: dict) -> dict:
    """N-M interaction of the 40x40 pile (12Ø25, Mode.CODIGO): uniaxial curve (bending about x)
    and the meridian of the N-Mx-My surface at the moment direction of the governing 2nd-order point
    (resultant moment), plus the governing ULS points (1st and 2nd order) and all the ULS rows."""
    from scipy.optimize import brentq, least_squares
    from armado_pilote import PROVIDED
    from seccion import KN, KNM, Mode
    sec = PROVIDED["fuste"].section(Mode.CODIGO, cell=4.0)
    # pivot A (t < 1) with the unlimited steel strain of Mode.CODIGO: the realistic planes are at t -> 1
    ts = np.concatenate([1.0 - np.geomspace(1.0, 1e-6, 500), np.linspace(1.0, 3.0, 300)])
    uni = np.array([sec.ult(math.pi / 2, t) for t in ts])
    rows = P["tables"]["uls_nm"]
    rows = [r for r in rows if r.get("eta") is not None]
    g2 = max((r for r in rows if r["check"] == "nm2"), key=lambda r: r["eta"])
    g1 = max((r for r in rows if r["check"] == "nm1"), key=lambda r: r["eta"])
    g2_1 = next(r for r in rows if r["check"] == "nm1" and r["pile"] == g2["pile"] and r["position"] == g2["position"])
    ang = math.atan2(abs(g2["MEd_y"]), abs(g2["MEd_x"]))
    Nmin, Nmax = uni[:, 0].min(), uni[:, 0].max()
    Ns = np.linspace(Nmin * 0.995, Nmax * 0.995, 90)
    mer = []
    t0 = brentq(lambda t: sec.ult(math.pi / 2, t)[0] - Ns[0], 0.0, 3.0)
    x0 = [math.pi / 2 - ang, t0]
    for N in Ns:
        def res(p):
            R = sec.ult(p[0], p[1])
            return [(R[0] - N) / 1e5, math.atan2(R[2], R[1]) - ang]
        sol = least_squares(res, x0, bounds=([-2 * math.pi, 0.0], [2 * math.pi, 3.0]), xtol=1e-12, ftol=1e-12)
        R = sec.ult(*sol.x)
        if np.linalg.norm(res(sol.x)) < 1e-4:
            mer.append((R[0] / KN, math.hypot(R[1], R[2]) / KNM))
            x0 = list(sol.x)
    pts = []
    for r in (g1, g2_1, g2):
        pts.append({"pile": r["pile"], "position": r["position"], "check": r["check"], "combo": r["combo"],
                    "N": r["N"], "M": math.hypot(r["MEd_x"], r["MEd_y"]), "MEd_x": r["MEd_x"], "MEd_y": r["MEd_y"],
                    "NRd": r["NRd"], "MRd": math.hypot(r["MRd_x"], r["MRd_y"]), "eta": r["eta"]})
    cloud = [{"N": r["N"], "M": math.hypot(r["MEd_x"], r["MEd_y"]), "check": r["check"]} for r in rows]
    return {"uniaxial": [(float(u[0] / KN), float(abs(u[1]) / KNM)) for u in uni],
            "meridian": mer, "meridian_angle_deg": math.degrees(ang), "points": pts, "cloud": cloud}


def confinement_estimate(s_mm: float) -> dict:
    """Indicative (not in the verdict): effective lateral pressure of the pile ties and confined
    strength A19.3.1.9 (3.24)/(3.25), with the effectiveness factor α = αn·αs of the confinement
    literature (EC8-1 5.4.3.2.2(8), MC2010 7.2.3.1.8): b0 = tie centreline, 12 bars all tied."""
    from armado_pilote import PROVIDED
    from seccion import bar_area
    lay = PROVIDED["fuste"]
    t = lay.ties
    b0 = lay.b - 2 * lay.cover - t.phi
    bi = lay.step
    alpha_n = 1 - lay.n_bars * bi ** 2 / (6 * b0 * b0)
    alpha_s = (1 - s_mm / (2 * b0)) ** 2
    rho = t.legs * bar_area(t.phi) / (s_mm * b0)
    fyd = lay.steel.fyd
    sigma2 = alpha_n * alpha_s * rho * fyd
    fck = lay.concrete.fck
    fckc = fck * (1 + 5 * sigma2 / fck) if sigma2 <= 0.05 * fck else fck * (1.125 + 2.5 * sigma2 / fck)
    return {"s_mm": s_mm, "b0_mm": b0, "alpha_n": alpha_n, "alpha_s": alpha_s, "rho_w": rho, "sigma2_MPa": sigma2,
            "fck_c": fckc, "limit_0.6fckc": 0.6 * fckc}


# =============================================================================================
# data model of the final design
# =============================================================================================
def _pile_verdict(s: dict, var: dict, basis: str, mode: str) -> tuple[str, float]:
    uls = _max(var.get(f"nm1[{basis},{mode}]"), var.get(f"nm2[{basis},{mode}]"), var.get(f"shear[{basis},{mode}]"))
    crack_ok = s.get(f"sls_crack[{QP_CASE[basis]}]_ok", s["sls_crack_ok"])
    ok = uls <= 1.0 and s["detailing_ok"] and crack_ok and s["sls_sigma_c"] <= 1.0 and s["sls_sigma_s"] <= 1.0
    return ("CUMPLE" if ok else "NO CUMPLE"), uls


PILE_CHECK_NAME = {"nm1": "N-M 1er orden", "nm2": "N-M 2º orden", "nm_stations": "N-M estaciones intermedias",
                   "shear": "cortante", "detailing": "disposiciones (As ≤ 0.04Ac)", "sls_crack": "fisuración QP",
                   "sls_sigma_s": "σs ≤ 0.8fyk (característica)", "sls_sigma_c": "σc ≤ 0.6fck (Ecm)",
                   "sls_sigma_c_Ec_eff": "σc ≤ 0.6fck (Ec,eff)"}
BEAM_ULS = ("flexión ELU", "cortante", "torsión")
BEAM_SLS = ("fisuración ELS", "tensiones ELS", "flecha")
BEAM_FAMS = ("flexión ELU", "cortante", "torsión", "disposiciones", "fisuración ELS", "tensiones ELS", "flecha")


def _fam_eta(case_summary: dict, fams) -> float | None:
    return _max(*[case_summary[f]["eta"] for f in fams if f in case_summary and not case_summary[f].get("info")])


def _governing(case_summary: dict) -> tuple[str, float]:
    best = max(((f, x) for f, x in case_summary.items() if isinstance(x, dict) and not x.get("info")),
               key=lambda fx: fx[1]["eta"])
    return best[0], best[1]["eta"]


def _gov_text(keys, rows: dict) -> str:
    k = max(keys, key=lambda m: rows[m]["eta_final"])
    return f"{rows[k]['gov_final']} ({k})"


def _group_of(member: str) -> str:
    return "inner" if member in INNER else ("end" if member in ("VT1", "VT7") else "edge")


def collect(P: dict, V: dict, curves: dict | None, confinement: dict) -> dict:
    """Everything the workbook, the md and the figures show, from the two result files and F0-F8."""
    S, var = P["summary"], {r["pile"]: r for r in P["tables"]["variants"]}
    D: dict = {"decisions": DECISIONS, "warnings": []}

    # ---- piles ------------------------------------------------------------------------------
    piles = []
    for p in PILES:
        s, v = S[p], var[p]
        row = {"pile": p, "side": "mar" if p in SEA else "tierra", "axis": AXIS_OF[p]}
        for b, m, case in VARIANTS:
            vd, uls = _pile_verdict(s, v, b, m)
            row[f"uls[{case}]"] = uls
            row[f"nm2[{case}]"] = v.get(f"nm2[{b},{m}]")
            row[f"nm1[{case}]"] = v.get(f"nm1[{b},{m}]")
            row[f"shear[{case}]"] = v.get(f"shear[{b},{m}]")
            row[f"verdict[{case}]"] = vd
        if row["verdict[CE-ROM]"] != s["verdict"]:
            D["warnings"].append(f"{p}: derived ROM/codigo verdict {row['verdict[CE-ROM]']} != pilotes.json {s['verdict']}")
        row["sls[CE-ROM]"] = _max(s["sls_crack"], s["sls_sigma_c"], s["sls_sigma_s"])
        keys = ("nm1", "nm2", "nm_stations", "shear", "detailing", "sls_crack", "sls_sigma_s")
        gk = max(keys, key=lambda k: s[k] if s[k] is not None else -1)
        row["eta_final"], row["gov_final"] = s[gk], PILE_CHECK_NAME[gk]
        row["sigma_c"], row["sigma_c_eff"] = s["sls_sigma_c"], s["sls_sigma_c_Ec_eff"]
        confined = s["sls_sigma_c"] > 1.0
        ok_final = (s["nm1"] <= 1 and s["nm2"] <= 1 and s["shear"] <= 1 and s["detailing_ok"] and s["sls_crack_ok"]
                    and s["sls_sigma_s"] <= 1 and (s["sls_sigma_c_Ec_eff"] <= 1 or not confined))
        row["verdict_final"] = ("CUMPLE (confinamiento)" if confined else "CUMPLE") if ok_final else "NO CUMPLE"
        row["final_layout"] = f"12Ø25; cercos Ø10 c/150, c/{PILE_TIE_S_NEW:g} en {PILE_TIE_ZONE_M:.2f} m superiores"
        row["note"] = (f"σc/0.6fck = {s['sls_sigma_c']:.3f} (Ecm) / {s['sls_sigma_c_Ec_eff']:.3f} (Ec,eff): "
                       "confinamiento A19.7.2(2)" if confined else "")
        piles.append(row)
    D["piles"] = piles
    nm2 = {c: max((r[f"nm2[{c}]"] or 0, r["pile"]) for r in piles) for _, _, c in VARIANTS}
    D["pile_key"] = {
        "nm2_max": {c: {"eta": x[0], "pile": x[1]} for c, x in nm2.items()},
        "shear_max": max(S[p]["shear"] for p in PILES),
        "sigma_c_max": max(S[p]["sls_sigma_c"] for p in PILES),
        "sigma_c_eff_max": max(S[p]["sls_sigma_c_Ec_eff"] for p in PILES),
        "sigma_c_fail": [p for p in PILES if S[p]["sls_sigma_c"] > 1],
        "crack_max": max(S[p]["sls_crack"] for p in PILES),
        "nm2_row": max((r for r in P["tables"]["uls_nm"] if r["check"] == "nm2" and r.get("eta") is not None),
                       key=lambda r: r["eta"]),
        "As_req_max": max(P["tables"]["as_required"], key=lambda r: r["As_req_cm2"]),
        "As_prov": P["meta"]["provided"]["fuste"]["As_cm2"],
        "tb05_fail": P["proposal"]["sensitivity"]["failing_piles"],
        "tb05_max_wk": max(r["wk"] for r in P["tables"]["sls_quasi_permanent"]
                           if r["case"] == "QP ψ2,Qa=0.8, ψ2,TB=0.5"),
        "kg_per_m": P["proposal"]["provided"]["kg_per_m"],
    }
    kg_tie_m = next(o["kg_ties"] for o in P["tables"]["options"] if o.get("provided"))
    D["pile_key"]["extra_tie_kg_per_pile"] = kg_tie_m * PILE_TIE_ZONE_M * (150.0 / PILE_TIE_S_NEW - 1.0)
    D["confinement"] = confinement
    if abs(D["pile_key"]["nm2_max"]["CE-ROM"]["eta"] - 0.987) > 0.0015:
        D["warnings"].append("pile nm2 max differs from F1 (0.987)")

    # ---- beams ------------------------------------------------------------------------------
    beams = []
    prop = V["proposal"]
    for k in BEAMS:
        cs = V["summary"][k]["cases"]
        g = _group_of(k)
        row = {"member": k, "name": BEAM_NAME[k], "group": g, "section": V["meta"]["members"][k]["section"]}
        for case in ("CYPE", "CE-CYPE", "CE-ROM"):
            row[f"uls[{case}]"] = _fam_eta(cs[case], BEAM_ULS)
            row[f"sls[{case}]"] = _fam_eta(cs[case], BEAM_SLS)
            row[f"verdict[{case}]"] = cs[case]["verdict"]
            for f in BEAM_FAMS:
                x = cs[case].get(f)
                row[f"{f}[{case}]"] = None if x is None or x.get("info") else x["eta"]
        fin = prop[GROUP_KEY[g]]["verification"][k] if g != "end" else cs["CE-ROM"]
        row["final"] = {f: (fin[f]["eta"] if f in fin and not fin[f].get("info") else None) for f in BEAM_FAMS}
        row["final_info"] = {f: fin[f]["eta"] for f in fin if isinstance(fin[f], dict) and fin[f].get("info")}
        gf, ge = _governing(fin)
        row["eta_final"], row["gov_final"] = ge, gf
        row["verdict_final"] = fin["verdict"]
        pg = prop[GROUP_KEY[g]]
        row["changed"] = any(pg[f]["change"] for f in ("top", "bottom", "stirrups"))
        beams.append(row)
    D["beams"] = beams

    # ---- reinforcement schedule (final design) ------------------------------------------------
    pin, pend, pedg = prop[GROUP_KEY["inner"]], prop[GROUP_KEY["end"]], prop[GROUP_KEY["edge"]]
    sens = prop[LEDGE_SENS_KEY]
    # unrounded tension-chord eta of the ledge-torsion sensitivity (VT1, CE-ROM: 1.1845); the proposal's
    # 'provided_eta' is already rounded to 3 decimals (1.185) and would print as 1.19 with 2 decimals (F4: 1.18)
    sens_eta = max(V["summary"][k]["cases"]["CE-ROM"]["torsión ala [sensib.]"]["eta"] for k in ("VT1", "VT7"))
    b_in = {r["member"]: r for r in beams}
    pk = D["pile_key"]
    D["schedule"] = [
        {"element": "Pilotes P1-P18 (14 ud., 40x40 HA-50 prefabricados)",
         "top": "12Ø25 (4 por cara, ρ = 3.68 %)", "bottom": "-", "side": "-",
         "transverse": f"3 cercos Ø10 (2eØ10+1eØ10) c/150 en el fuste; c/{PILE_TIE_S_NEW:g} en los "
                       f"{PILE_TIE_ZONE_M:.2f} m superiores; arranque 1eØ8",
         "change": f"cercos Ø10 c/150 -> c/{PILE_TIE_S_NEW:g} en {PILE_TIE_ZONE_M:.2f} m superiores (confinamiento, "
                   f"+{pk['extra_tie_kg_per_pile']:.1f} kg/pilote)",
         "eta": max(r["eta_final"] for r in piles),
         "gov": f"N-M 2º orden, {pk['nm2_row']['position'].lower()} {pk['nm2_row']['pile']} ({pk['nm2_row']['combo']})",
         "verdict": "CUMPLE (con cercos de confinamiento)"},
        {"element": "Vigas interiores ejes 2-6 (T invertida 50x55+15x30+15x30)",
         "top": pin["top"]["proposed"] + f" ({pin['top']['proposed_As_cm2']:.2f} cm²)",
         "bottom": f"{pin['bottom']['proposed']} ({pin['bottom']['proposed_As_cm2']:.2f} cm²)",
         "side": "2Ø10 + 2Ø10 a 235 mm (alas / piel)",
         "transverse": "cerco Ø10 + rama Ø10 (3 ramas en el alma) + cerco de alas Ø10, c/100",
         "change": f"inferior {pin['bottom']['provided']} ({pin['bottom']['provided_As_cm2']:.2f} cm²) -> "
                   f"{pin['bottom']['proposed']} ({pin['bottom']['proposed_As_cm2']:.2f} cm²); "
                   f"{pin['weight_long_kg_m']['provided']:.1f} -> {pin['weight_long_kg_m']['proposed']:.1f} kg/m long.",
         "eta": max(b_in[k]["eta_final"] for k in INNER), "gov": _gov_text(INNER, b_in) +
         f" (wk = {pin['bottom']['wk_mm']:.3f} mm, ψ2 0.8)", "verdict": "CUMPLE"},
        {"element": "Vigas extremas ejes 1 y 7 (L invertida 80x55+15x30)",
         "top": f"{pend['top']['provided']} ({pend['top']['provided_As_cm2']:.2f} cm²)",
         "bottom": f"5Ø20 alma + 1Ø20 ala ({pend['bottom']['provided_As_cm2']:.2f} cm²)",
         "side": "1Ø20 sup. ala + 2Ø10 a 235 mm",
         "transverse": "cerco Ø10 + rama Ø10 (3 ramas) + cerco de ala Ø10, c/100",
         "change": f"ninguno (si las alveoplacas apoyan en el ala: superior {sens['top']['proposed']} en lugar de "
                   f"{sens['top']['provided']}, η {sens_eta:.2f} -> {sens['top']['proposed_eta']:.2f})",
         "eta": max(b_in[k]["eta_final"] for k in ("VT1", "VT7")),
         "gov": _gov_text(("VT1", "VT7"), b_in) + ": cordón traccionado M/MRd + ΣAsl,T/As", "verdict": "CUMPLE"},
        {"element": "Vigas de borde pórticos 1 y 10 (25x30)",
         "top": f"{pedg['top']['proposed']} sobre apoyos ({pedg['top']['proposed_As_cm2']:.2f} cm²)",
         "bottom": f"{pedg['bottom']['proposed']} ({pedg['bottom']['proposed_As_cm2']:.2f} cm²)", "side": "-",
         "transverse": pedg["stirrups"]["proposed"].replace("2 ramas ", "cerco ").replace(" (1e)", ""),
         "change": f"superior {pedg['top']['provided']} -> {pedg['top']['proposed']} sobre apoyos (condicionado al "
                   f"reparto de carga de SAP, wk {pedg['top']['wk_mm']:.3f} mm)",
         "eta": max(b_in[k]["eta_final"] for k in ("VBM", "VBT")),
         "gov": _gov_text(("VBM", "VBT"), b_in), "verdict": "CUMPLE (condicionado, F5)"},
    ]
    D["proposal"] = {"inner": pin, "end": pend, "edge": pedg, "end_sensitivity": sens,
                     "end_sensitivity_eta": sens_eta}
    D["curves"] = curves

    # ---- parity -----------------------------------------------------------------------------
    par = []
    for r in P["cype_comparison"]:
        st = "INFO" if r.get("info") else ("PASS" if r.get("ok") else "FAIL")
        kind = PARITY_SAP if "SAP" in r["quantity"] else PARITY_EXACT
        par.append({"element": "Pilotes", "kind": kind, "quantity": r["quantity"], "ours": r["ours"],
                    "cype": r["cype"], "diff_pct": r["diff_pct"], "status": st, "src": r.get("src", ""),
                    "note": r.get("note", "")})
    for r in V["cype_comparison"]:
        kind = "fuerzas SAP (comprobación cruzada)" if "SAP" in r["quantity"] else PARITY_EXACT
        par.append({"element": "Vigas", "kind": kind, "quantity": r["quantity"], "ours": r["ours"], "cype": r["cype"],
                    "diff_pct": r["diff_pct"], "status": r["status"], "src": r.get("src", ""),
                    "note": r.get("note", "")})
    D["parity"] = par
    return D


# =============================================================================================
# workbook
# =============================================================================================
def write_xlsx(D: dict, P: dict, V: dict, path: Path = XLSX) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEAD = PatternFill("solid", fgColor="1F4E78")
    GOOD = PatternFill("solid", fgColor="E2EFDA")
    WARN = PatternFill("solid", fgColor="FFF2CC")
    BAD = PatternFill("solid", fgColor="F8CBAD")
    INFO = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="D0D0D0")
    BORDER = Border(bottom=thin)

    def eta_fill(v):
        if not isinstance(v, (int, float)):
            return None
        return GOOD if v <= 0.95 else WARN if v <= 1.0 else BAD

    def verdict_fill(v):
        if not isinstance(v, str):
            return None
        if v.startswith("NO"):
            return BAD
        if "(" in v:
            return WARN
        return GOOD if v.startswith("CUMPLE") or v == "PASS" else INFO if v == "INFO" else None

    def title(ws, text, sub=None):
        ws["A1"] = text
        ws["A1"].font = Font(bold=True, size=13)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = Font(italic=True, color="52514E")

    def table(ws, top, headers, rows, eta_cols=(), verdict_cols=(), fmt=None, widths=None, wrap_cols=()):
        for j, h in enumerate(headers, 1):
            c = ws.cell(top, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[top].height = 32
        for i, r in enumerate(rows, top + 1):
            for j, v in enumerate(r, 1):
                if isinstance(v, (list, tuple, dict)):
                    v = json.dumps(v, ensure_ascii=False)
                c = ws.cell(i, j, v)
                c.border = BORDER
                if isinstance(v, float):
                    c.number_format = (fmt or {}).get(j, "0.000")
                if j in wrap_cols:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            for j in eta_cols:
                f = eta_fill(ws.cell(i, j).value)
                if f:
                    ws.cell(i, j).fill = f
            for j in verdict_cols:
                f = verdict_fill(ws.cell(i, j).value)
                if f:
                    ws.cell(i, j).fill = f
        for j, h in enumerate(headers, 1):
            w = (widths or {}).get(j)
            if w is None:
                w = min(max(10, len(str(h)) * 0.9 + 2), 28)
            cur = ws.column_dimensions[get_column_letter(j)].width or 0
            ws.column_dimensions[get_column_letter(j)].width = max(cur if cur > 13 else 0, w)
        return top + len(rows)

    def color_series(s, color, kind="bar", width_pt=2.0, dash=None, marker=None):
        if kind == "bar":
            s.graphicalProperties.solidFill = color.lstrip("#")
            s.graphicalProperties.line.solidFill = "FFFFFF"
        else:
            s.graphicalProperties.line.solidFill = color.lstrip("#")
            s.graphicalProperties.line.width = int(width_pt * 12700)
            if dash:
                s.graphicalProperties.line.dashStyle = dash
            s.marker.symbol = marker or "none"
            if marker:
                s.marker.size = 6
                s.marker.graphicalProperties.solidFill = color.lstrip("#")
                s.marker.graphicalProperties.line.solidFill = color.lstrip("#")
            s.smooth = False

    def axes_on(ch):
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.legend.position = "b"

    def bar_chart(ws, title_, ytitle, cats_ref, cols, colors, first, last, anchor, limit_col=None,
                  limit_title=None, extra_lines=(), ymax=None, width=22, height=9):
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "clustered"
        ch.gapWidth = 60
        ch.title = title_
        ch.y_axis.title = ytitle
        ch.height, ch.width = height, width
        for c in cols:
            ch.add_data(Reference(ws, min_col=c, min_row=first - 1, max_row=last), titles_from_data=True)
        ch.set_categories(cats_ref)
        for s, colr in zip(ch.series, colors):
            color_series(s, colr)
        lines = ([(limit_col, C_LIMIT, "dash")] if limit_col else []) + list(extra_lines)
        if lines:
            ln = LineChart()
            for c, colr, dash in lines:
                ln.add_data(Reference(ws, min_col=c, min_row=first - 1, max_row=last), titles_from_data=True)
            ln.set_categories(cats_ref)
            for s, (c, colr, dash) in zip(ln.series, lines):
                color_series(s, colr, "line", 1.75, dash)
            ch += ln
        if ymax:
            ch.y_axis.scaling.min = 0
            ch.y_axis.scaling.max = ymax
        axes_on(ch)
        ws.add_chart(ch, anchor)
        return ch

    wb = Workbook()
    fin = next(c for c in V["meta"]["cases"] if c["final"])
    curves = D["curves"]

    # ---- Resumen -----------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    title(ws, "Diseño a Código Estructural (CE 2021, Anejo 19) — Muelle de Trasmallo, módulo 40 m (Puerto de Cullera)",
          "Esfuerzos SAP2000 v27.1; comprobaciones diseno/python/codigo (pilotes.py, vigas.py); veredicto final: "
          "Mode.CODIGO + base ROM (ψ0,Qa 1.0, ψ2,Qa 0.8, bolardo ψ2 0) + XS3 wmax 0.1 mm (cuasipermanente). "
          f"Generado por diseno/python/run_diseno.py el {time.strftime('%Y-%m-%d')}.")
    hdr = ["Elemento", "Tipo / armado Anejo 10", "η ELU CYPE/cype (paridad)", "η ELU CYPE/codigo",
           "η ELU ROM/codigo", "η ELS ROM/codigo (armado actual)", "Veredicto CYPE/cype", "Veredicto CYPE/codigo",
           "Veredicto ROM/codigo (armado actual)", "Diseño final (armado)", "η determinante final",
           "Veredicto final", "Comprobación determinante / nota"]
    rows = []
    for r in D["piles"]:
        rows.append([f"{r['pile']} ({r['side']}, eje {r['axis']})", "Pilote 40x40 HA-50, 12Ø25 + 3cØ10/150",
                     r["uls[CYPE]"], r["uls[CE-CYPE]"], r["uls[CE-ROM]"], r["sls[CE-ROM]"], r["verdict[CYPE]"],
                     r["verdict[CE-CYPE]"], r["verdict[CE-ROM]"], r["final_layout"], r["eta_final"],
                     r["verdict_final"], r["gov_final"] + ("; " + r["note"] if r["note"] else "")])
    for b in D["beams"]:
        lay = {"inner": "T inv. 50x55+15x30+15x30: 5Ø20 sup. / 5Ø20+2Ø20 inf. / 3r Ø10/100",
               "end": "L inv. 80x55+15x30: 5Ø20 sup. / 5Ø20+1Ø20 inf. / 3r Ø10/100",
               "edge": "25x30: 2Ø16 sup. / 2Ø16 inf. / cØ8/150"}[b["group"]]
        final_lay = {"inner": "inf. 8Ø20 alma + 4Ø16 alas (33.18 cm²); resto sin cambio",
                     "end": "sin cambio", "edge": "sup. 3Ø25 sobre apoyos (condicionado, F5)"}[b["group"]]
        rows.append([f"{b['member']} {b['name']}", lay, b["uls[CYPE]"], b["uls[CE-CYPE]"], b["uls[CE-ROM]"],
                     b["sls[CE-ROM]"], b["verdict[CYPE]"], b["verdict[CE-CYPE]"], b["verdict[CE-ROM]"], final_lay,
                     b["eta_final"], b["verdict_final"] + (" (condicionado)" if b["group"] == "edge" else ""),
                     b["gov_final"]])
    top = 4
    ws.cell(top - 1, 1, "1. Veredicto por elemento (η = solicitación / capacidad)").font = Font(bold=True, size=11)
    last = table(ws, top, hdr, rows, eta_cols=(3, 4, 5, 6, 11), verdict_cols=(7, 8, 9, 12),
                 widths={1: 26, 2: 40, 10: 44, 13: 60}, wrap_cols=(2, 10, 13))
    ws.freeze_panes = ws.cell(top + 1, 2)
    r0 = last + 2
    ws.cell(r0, 1, "Leyenda").font = Font(bold=True)
    legend = [
        (GOOD, "η ≤ 0.95 / CUMPLE"), (WARN, "0.95 < η ≤ 1.00 / CUMPLE con medida o condicionado"),
        (BAD, "η > 1.00 / NO CUMPLE"),
        (None, "CYPE/cype = hipótesis del Anejo 10 (ψ0,Qa 0.7, ψ2,Qa 0.3) con las convenciones de CYPECAD (paridad); "
               "CYPE/codigo = mismas cargas, sección estricta al Código; ROM/codigo = ROM 2.0-11 (ψ0,Qa 1.0, "
               "ψ2,Qa 0.8), sección estricta: VEREDICTO."),
        (None, "Pilotes: veredicto por base = ELU (N-M 1er/2º orden, cortante) + disposiciones + fisuración QP con el "
               "ψ2 de la base + σc ≤ 0.6fck y σs ≤ 0.8fyk (característica ELS01-08, igual en ambas bases). Vigas: casos "
               "CYPE / CE-CYPE / CE-ROM de vigas.py."),
        (None, "Diseño final: η determinante de la verificación con el armado propuesto (vigas interiores y de borde) "
               "o con el dispuesto (vigas extremas, pilotes). Pilotes: el límite σc ≤ 0.6fck se cubre con el "
               "confinamiento de los cercos (A19.7.2(2) 'otras medidas'), ver F1."),
    ]
    for i, (f, t) in enumerate(legend, r0 + 1):
        c = ws.cell(i, 1, "" if f is None else "  ")
        if f:
            c.fill = f
        ws.cell(i, 2, t)
    r0 = r0 + len(legend) + 2
    ws.cell(r0, 1, "2. Decisiones del diseño final (F0-F8)").font = Font(bold=True, size=11)
    ncol = len(hdr)
    for j, h in ((1, "Id"), (2, "Tema"), (3, "Decisión")):
        c = ws.cell(r0 + 1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    ws.merge_cells(start_row=r0 + 1, start_column=3, end_row=r0 + 1, end_column=ncol)
    width_chars = sum(ws.column_dimensions[get_column_letter(j)].width or 13 for j in range(3, ncol + 1))
    for i, (fid, tema, text) in enumerate(D["decisions"], r0 + 2):
        ws.cell(i, 1, fid).alignment = Alignment(vertical="top")
        ws.cell(i, 2, tema).alignment = Alignment(vertical="top", wrap_text=True)
        c = ws.cell(i, 3, text)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=ncol)
        ws.row_dimensions[i].height = 14 * max(2, math.ceil(len(text) / (width_chars * 1.1)) + 1)
    last = r0 + 1 + len(D["decisions"])
    r0 = last + 2
    ws.cell(r0, 1, "3. Cuadro de armado final").font = Font(bold=True, size=11)
    table(ws, r0 + 1, ["Elemento", "Longitudinal superior", "Longitudinal inferior", "Lateral / piel", "Transversal",
                       "Cambio frente al Anejo 10", "η determinante", "Comprobación determinante", "Veredicto"],
          [[s["element"], s["top"], s["bottom"], s["side"], s["transverse"], s["change"], s["eta"], s["gov"],
            s["verdict"]] for s in D["schedule"]], eta_cols=(7,), verdict_cols=(9,), wrap_cols=(1, 2, 3, 5, 6, 8))

    # ---- Pilotes_ELU -------------------------------------------------------------------------
    ws = wb.create_sheet("Pilotes_ELU")
    title(ws, "Pilotes: ELU N-Mx-My (1er y 2º orden, A19.5.8.8 / A19.6.1) y cortante (A19.6.2) — base ROM, "
              "Mode.CODIGO", "η N-M: 'mismas excentricidades' (rayo); cortante: veredicto = VRd,c o cercos "
                             "(cot θ óptimo) y VRd,max; 'Aprov.' CYPE = listado del Anejo 10 (SAP forces: ours).")
    nm = {(r["pile"], r["position"], r["check"]): r for r in P["tables"]["uls_nm"]}
    sh = {}
    for r in P["tables"]["shear"]:
        sh[(r["pile"], r["position"], r["key"])] = r
    ap = {(r["pile"], r["position"]): r for r in P["tables"].get("aprov_listado", [])}
    rows = []
    for p in PILES:
        for pos in ("Cabeza", "Pie", "Arranque"):
            n1, n2 = nm.get((p, pos, "nm1")), nm.get((p, pos, "nm2"))
            vc, vm, vv = sh.get((p, pos, "vrdc")), sh.get((p, pos, "vrdmax")), sh.get((p, pos, "verdict"))
            a = ap.get((p, pos), {})
            n2 = n2 if n2 and n2.get("eta") is not None else {}
            rows.append([p, "mar" if p in SEA else "tierra", pos, n1["combo"], n1["N"], n1["MEd_x"], n1["MEd_y"],
                         n1["eta"], n2.get("combo", "N.P."), n2.get("MEd_x"), n2.get("MEd_y"), n2.get("eta"),
                         vv["combo"], vv["demand"],
                         vc["capacity"] if vc else None, vm["capacity"] if vm else None, vv["capacity"], vv["eta"],
                         _r(a.get("cype_NM", None) / 100, 3) if a.get("cype_NM") is not None else None,
                         _r(a.get("cype_Q", None) / 100, 3) if a.get("cype_Q") is not None else None])
    hdr = ["Pilote", "Lado", "Sección", "Combo 1er", "N (kN)", "MEd,x 1er (kN·m)", "MEd,y 1er", "η N-M 1er",
           "Combo 2º", "MEd,x 2º (kN·m)", "MEd,y 2º", "η N-M 2º", "Combo V", "VEd (kN)", "VRd,c (kN)",
           "VRd,max (kN)", "VRd veredicto (kN)", "η cortante", "CYPE Aprov. N,M", "CYPE Aprov. Q"]
    last = table(ws, 4, hdr, rows, eta_cols=(8, 12, 18), fmt={5: "0.0", 6: "0.0", 7: "0.0", 10: "0.0", 11: "0.0",
                                                             14: "0.0", 15: "0.0", 16: "0.0", 17: "0.0"})
    ws.freeze_panes = "D5"
    r0 = last + 3
    ws.cell(r0 - 1, 1, "η N-M 2º orden máximo por pilote y base / modo (tabla del gráfico)").font = Font(bold=True)
    rows = []
    for p in SEA + LAND:
        r = next(x for x in D["piles"] if x["pile"] == p)
        rows.append([f"{p} ({'m' if p in SEA else 't'}{AXIS_OF[p]})", r["nm2[CYPE]"], r["nm2[CE-CYPE]"],
                     r["nm2[CE-ROM]"], r["nm1[CE-ROM]"], r["shear[CE-ROM]"], 1.0])
    last2 = table(ws, r0, ["Pilote", "η 2º CYPE/cype", "η 2º CYPE/codigo", "η 2º ROM/codigo", "η 1er ROM/codigo",
                           "η cortante ROM/codigo", "Límite η = 1"], rows, eta_cols=(2, 3, 4, 5, 6))
    ws.cell(last2 + 1, 1, "Celdas vacías: sin comprobación de 2º orden en ese modo (λ ≤ λlim en todas las "
                          "combinaciones, pilotes traccionados).")
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=last2)
    bar_chart(ws, "η N-M 2º orden por pilote (m = mar, t = tierra; eje 1-7)", "η", cats, [2, 3, 4],
              [C_CYPE, C_CECY, C_ROM], r0 + 1, last2, "V4", limit_col=7, ymax=1.1)
    bar_chart(ws, "η cortante por pilote (ROM/codigo)", "η", cats, [6], [C_ROM], r0 + 1, last2, "V24",
              limit_col=7, ymax=1.1)

    # ---- Pilotes_As --------------------------------------------------------------------------
    ws = wb.create_sheet("Pilotes_As")
    title(ws, "Pilotes: armadura longitudinal necesaria frente a dispuesta (N-M 1er y 2º orden, base ROM)",
          "Misma disposición de 12 barras, área escalada; As,máx = 0.04·Ac = 64 cm² (A19.9.5.2(3)); "
          "práctica = disposición comercial mínima que cubre As,nec en ese pilote.")
    asr = {}
    for r in P["tables"]["as_required"]:
        asr.setdefault(r["pile"], {})[r["zone"]] = r
    rows = []
    for p in SEA + LAND:
        f, a = asr[p]["fuste"], asr[p]["arranque"]
        g = f if f["As_req_cm2"] >= a["As_req_cm2"] else a
        rows.append([f"{p} ({'m' if p in SEA else 't'}{AXIS_OF[p]})", f["position"], f["combo"], f["As_req_cm2"],
                     a["As_req_cm2"], g["As_req_cm2"], g["As_prov_cm2"], g["As_req_cm2"] / g["As_prov_cm2"],
                     g["eta_prov"], g["practical"], g["practical_As_cm2"], 64.0])
    hdr = ["Pilote", "Sección fuste", "Combo", "As,nec fuste (cm²)", "As,nec arranque (cm²)", "As,nec máx (cm²)",
           "As dispuesta (cm²)", "As,nec / As,disp", "η con As dispuesta", "Disposición práctica",
           "As práctica (cm²)", "As,máx 0.04Ac (cm²)"]
    last = table(ws, 4, hdr, rows, eta_cols=(8, 9), fmt={4: "0.00", 5: "0.00", 6: "0.00", 7: "0.00", 11: "0.00",
                                                         12: "0.0"}, widths={10: 22})
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    bar_chart(ws, "As necesaria por pilote frente a dispuesta (12Ø25 = 58.9 cm²) y As,máx", "cm²", cats, [6],
              [C_ROM], 5, last, "N4", extra_lines=[(7, INK2, None), (12, C_LIMIT, "dash")], ymax=70)
    r0 = last + 3
    ws.cell(r0 - 1, 1, "Disposiciones alternativas evaluadas (pilote uniforme, envolvente de los 14 pilotes)").font = \
        Font(bold=True)
    rows = [[o["layout"], o["ties"], o["As_cm2"], o["kg_per_m"], o.get("nm"), o.get("shear"), o.get("detailing"),
             o.get("sls_crack"), o.get("sls_crack_tb05"), o.get("sls_sigma_c"), o["status"],
             ", ".join(o.get("fails") or [])]
            for o in P["tables"]["options"]]
    table(ws, r0, ["Disposición", "Cercos", "As (cm²)", "kg/m", "η N-M", "η cortante", "η disposiciones",
                   "η fisuración QP", "η fisuración ψ2,TB 0.5", "σc/0.6fck (Ecm)", "Estado", "No cumple"], rows,
          eta_cols=(5, 6, 7, 8, 9, 10), fmt={3: "0.00", 4: "0.00"}, widths={1: 18, 11: 26})

    # ---- Pilotes_ELS -------------------------------------------------------------------------
    ws = wb.create_sheet("Pilotes_ELS")
    title(ws, "Pilotes: ELS — fisuración cuasipermanente (XS3, wmax 0.1 mm) y tensiones características",
          "Fisuración: η = σct/fctm si no fisura (sección homogeneizada), wk/0.1 si fisura. σc/0.6fck con Ecm "
          "(veredicto de pilotes.py) y con Ec,eff = Ecm/(1 + φ·Mqp/Mk) para la parte sostenida (información).")
    qp = {}
    for r in P["tables"]["sls_quasi_permanent"]:
        k = (r["pile"], r["case"])
        if k not in qp or r["eta"] > qp[k]["eta"]:
            qp[k] = r
    ch_ = {}
    for r in P["tables"]["sls_characteristic"]:
        if r["check"].startswith("σc"):
            k = r["pile"]
            if k not in ch_ or r["eta"] > ch_[k]["eta"]:
                ch_[k] = r
    rows = []
    for p in SEA + LAND:
        s = P["summary"][p]
        q = qp[(p, QP_CASE["ROM"])]
        c = ch_[p]
        rows.append([f"{p} ({'m' if p in SEA else 't'}{AXIS_OF[p]})", q["position"], q["N"], q["sig_ct"], q["fctm"],
                     q["sig_ct"] / q["fctm"], "sí" if q["cracked"] else "no", q["wk"], c["position"], c["combo"],
                     c["demand"], c["sig_c_Ec_eff"], s["sls_sigma_c"], s["sls_sigma_c_Ec_eff"], s["sls_sigma_s"],
                     s["sls_sigma_c_qp"],
                     P["summary"][p][f"sls_crack[{QP_CASE['CYPE']}]"],
                     P["summary"][p]["sls_crack[QP ψ2,Qa=0.3, ψ2,TB=0.5]"], s["sls_crack"],
                     P["summary"][p]["sls_crack[QP ψ2,Qa=0.8, ψ2,TB=0.5]"], 1.0])
    hdr = ["Pilote", "Sección QP", "N QP (kN)", "σct (MPa)", "fctm (MPa)", "σct/fctm (ψ2 0.8)", "Fisura",
           "wk (mm)", "Sección σc", "Combo σc", "σc Ecm (MPa)", "σc Ec,eff (MPa)", "σc/0.6fck Ecm",
           "σc/0.6fck Ec,eff", "σs/0.8fyk", "σc,qp/0.45fck", "η fis. ψ2 0.3, TB 0", "η fis. ψ2 0.3, TB 0.5",
           "η fis. ψ2 0.8, TB 0 (veredicto)", "η fis. ψ2 0.8, TB 0.5", "Límite"]
    last = table(ws, 4, hdr, rows, eta_cols=(6, 13, 14, 15, 16, 17, 18, 19, 20),
                 fmt={3: "0.0", 4: "0.00", 5: "0.00", 8: "0.000", 11: "0.00", 12: "0.00"})
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    bar_chart(ws, "σc/0.6fck (característica) por pilote: Ecm frente a Ec,eff", "η", cats, [13, 14],
              [C_CECY, C_ROM], 5, last, "B22", limit_col=21, ymax=1.2)
    bar_chart(ws, "Fisuración QP: sensibilidad al ψ2 (bolardo 0 / 0.5)", "η", cats, [17, 19, 20],
              [C_CYPE, C_ROM, C_CECY], 5, last, "N22", limit_col=21, ymax=1.8)
    ws.cell(last + 1, 1, f"P1, P2, P3, P16: σc > 0.6fck con Ecm -> se mantiene 12Ø25 y se reduce la separación de "
                         f"los cercos a Ø10 c/{PILE_TIE_S_NEW:g} en los {PILE_TIE_ZONE_M:.2f} m superiores (confinamiento, "
                         "A19.7.2(2)); ver Propuesta y F1. ψ2,TB = 0.5: fisuran seis pilotes (F2).")

    # ---- Vigas_Flexion -----------------------------------------------------------------------
    ws = wb.create_sheet("Vigas_Flexion")
    title(ws, f"Vigas: flexión ELU (A19.6.1) y armadura necesaria / dispuesta — caso final {fin['name']}",
          "As,uls = armadura de tracción de cálculo (simple); As,req = max(As,uls, As,min CE 9.1). Filas "
          "'eje pilote (info)' = momento en el eje del nudo rígido, información (F6).")
    asrow = {(r["member"], r["section"], r["face"]): r for r in V["tables"]["as_required"] if r["case"] == fin["name"]}
    rows = []
    for r in V["tables"]["bending"]:
        if r["case"] != fin["name"]:
            continue
        a = asrow.get((r["member"], r["section"], r["face"]), {})
        rows.append([r["member"], r["section"], r["pos_m"], r["face"], r["combo"], r["demand"], r["capacity"],
                     r["eta"], a.get("As_uls_cm2"), a.get("As_min_cm2"), a.get("As_req_cm2"), r["As_prov_cm2"],
                     a.get("req_over_prov"), "sí" if r["in_verdict"] else "info"])
    hdr = ["Viga", "Sección", "y / x (m)", "Cara", "Combo", "MEd (kN·m)", "MRd (kN·m)", "η", "As,uls (cm²)",
           "As,min (cm²)", "As,req (cm²)", "As,disp (cm²)", "As,req/As,disp", "En veredicto"]
    last = table(ws, 4, hdr, rows, eta_cols=(8, 13), fmt={3: "0.000", 6: "0.0", 7: "0.0", 9: "0.00", 10: "0.00",
                                                          11: "0.00", 12: "0.00"}, widths={2: 34})
    for i in range(5, last + 1):
        if ws.cell(i, 14).value == "info":
            for j in (8, 13):
                ws.cell(i, j).fill = INFO
    ws.freeze_panes = "C5"
    r0 = last + 3
    ws.cell(r0 - 1, 1, "Resumen por viga (tabla del gráfico): η flexión por caso y As,req/As,disp máximos (caso final)"
            ).font = Font(bold=True)
    rows = []
    for b in D["beams"]:
        k = b["member"]
        top = max((x for x in asrow.values() if x["member"] == k and x["face"] == "top" and x["in_verdict"]),
                  key=lambda x: x["req_over_prov"])
        bot = max((x for x in asrow.values() if x["member"] == k and x["face"] == "bottom" and x["in_verdict"]),
                  key=lambda x: x["req_over_prov"])
        rows.append([k, b["flexión ELU[CYPE]"], b["flexión ELU[CE-CYPE]"], b["flexión ELU[CE-ROM]"],
                     top["As_req_cm2"], top["As_prov_cm2"], bot["As_req_cm2"], bot["As_prov_cm2"], 1.0])
    last2 = table(ws, r0, ["Viga", "η flexión CYPE/cype", "η flexión CYPE/codigo", "η flexión ROM/codigo",
                           "As,req sup. (cm²)", "As,disp sup. (cm²)", "As,req inf. (cm²)", "As,disp inf. (cm²)",
                           "Límite"], rows, eta_cols=(2, 3, 4), fmt={5: "0.00", 6: "0.00", 7: "0.00", 8: "0.00"})
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=last2)
    bar_chart(ws, "η flexión ELU por viga y base / modo", "η", cats, [2, 3, 4], [C_CYPE, C_CECY, C_ROM], r0 + 1,
              last2, "P4", limit_col=9, ymax=1.1)
    bar_chart(ws, "Armadura necesaria frente a dispuesta (caso final, armado actual)", "cm²", cats, [5, 6, 7, 8],
              [C_ROM, "#86b6ef", C_CECY, "#f3a582"], r0 + 1, last2, "P23")

    # ---- Vigas_Cortante ----------------------------------------------------------------------
    ws = wb.create_sheet("Vigas_Cortante")
    title(ws, f"Vigas: cortante (A19.6.2), suspensión de las alveoplacas (A19.6.2.1(9)) y torsión (A19.6.3) — "
              f"caso final {fin['name']}", "VRd,max con V en la cara; VRd,s y Asw/s con V a d de la cara; "
                                          "cot θ óptimo en [0.5, 2]; fywd = 400 MPa con ν1 = 0.6.")
    st = [r for r in V["tables"]["shear_torsion"] if r["case"] == fin["name"]]
    idx: dict = {}
    for r in st:
        idx.setdefault((r["member"], r["section"]), []).append(r)
    rows = []
    for (m, sec), rs in idx.items():
        def f(pred):
            return next((x for x in rs if pred(x["check"])), None)
        vmax = f(lambda c: c.startswith("VEd(cara) <= VRd,max"))
        vs = f(lambda c: c.startswith("VEd(d) <= VRd,s"))
        hg = f(lambda c: c.startswith("estribos V + suspensión"))
        tmax = f(lambda c: c.startswith("TEd/TRd,max") and "sensibilidad" not in c)
        leg = f(lambda c: c.startswith("estribos V + T") and "sensibilidad" not in c)
        chord = f(lambda c: c.startswith("M/MRd + ΣAsl,T") and "sensibilidad" not in c)
        chs = f(lambda c: c.startswith("M/MRd + ΣAsl,T") and "sensibilidad" in c)
        if vs is None and vmax is None:
            continue
        rows.append([m, sec, vmax["combo"] if vmax else None, vmax["demand"] if vmax else None,
                     vmax["capacity"] if vmax else None, vmax["eta"] if vmax else None,
                     vs["demand"] if vs else None, vs.get("VRd_c") if vs else None, vs["capacity"] if vs else None,
                     vs["eta"] if vs else None, vs.get("cot_theta") if vs else None,
                     vs.get("Asw_s_req_cm2_m") if vs else None, hg.get("Asw_hanger_cm2_m") if hg else None,
                     hg["demand"] if hg else (vs.get("Asw_s_req_cm2_m") if vs else None),
                     vs.get("Asw_s_prov_cm2_m") if vs else None, hg["eta"] if hg else (vs["eta"] if vs else None),
                     tmax["demand"] if tmax else None, tmax["eta"] if tmax else None, leg["eta"] if leg else None,
                     chord["eta"] if chord else None,
                     "sí" if (chord and chord["in_verdict"]) else ("info" if chord else ""),
                     chs["eta"] if chs else None])
    hdr = ["Viga", "Sección", "Combo", "VEd cara (kN)", "VRd,max (kN)", "η VRd,max", "VEd(d) (kN)", "VRd,c (kN)",
           "VRd,s (kN)", "η VRd,s", "cot θ", "Asw/s nec. V (cm²/m)", "Asw/s suspensión (cm²/m)",
           "Asw/s nec. total (cm²/m)", "Asw/s disp. (cm²/m)", "η estribos (V + susp.)", "TEd (kN·m)",
           "η TEd/TRd,max + VEd/VRd,max", "η estribos V + T (rama)", "η cordón M + T", "Torsión en veredicto",
           "η cordón sensib. ala (F4)"]
    last = table(ws, 4, hdr, rows, eta_cols=(6, 10, 16, 18, 19, 20, 22),
                 fmt={4: "0.0", 5: "0.0", 7: "0.0", 8: "0.0", 9: "0.0", 11: "0.00", 12: "0.00", 13: "0.00",
                      14: "0.00", 15: "0.00", 17: "0.0"}, widths={2: 30})
    for i in range(5, last + 1):
        if ws.cell(i, 21).value == "info":
            for j in (18, 19, 20):
                ws.cell(i, j).fill = INFO
    ws.freeze_panes = "C5"
    r0 = last + 3
    ws.cell(r0 - 1, 1, "Resumen por viga (tabla del gráfico): η cortante por caso; torsión (vigas extremas)").font = \
        Font(bold=True)
    rows = [[b["member"], b["cortante[CYPE]"], b["cortante[CE-CYPE]"], b["cortante[CE-ROM]"], b["torsión[CE-ROM]"],
             b["final_info"].get("torsión ala [sensib.]"), 1.0] for b in D["beams"]]
    last2 = table(ws, r0, ["Viga", "η cortante CYPE/cype (θ 45°)", "η cortante CYPE/codigo", "η cortante ROM/codigo",
                           "η torsión ROM/codigo (en veredicto)", "η cordón sensib. ala (F4)", "Límite"], rows,
                  eta_cols=(2, 3, 4, 5, 6))
    cats = Reference(ws, min_col=1, min_row=r0 + 1, max_row=last2)
    bar_chart(ws, "η cortante por viga y base / modo (CYPE: θ = 45°, CE: cot θ óptimo + suspensión)", "η", cats,
              [2, 3, 4], [C_CYPE, C_CECY, C_ROM], r0 + 1, last2, "X4", limit_col=7, ymax=1.2)
    ws.cell(last2 + 1, 1, "Torsión: en veredicto sólo para las vigas extremas (ejes 1 y 7, torsión de equilibrio); "
                          "en las interiores y de borde es de compatibilidad (A19.6.3.1(2)), información.")

    # ---- Vigas_Fisuracion --------------------------------------------------------------------
    ws = wb.create_sheet("Vigas_Fisuracion")
    title(ws, "Vigas: fisuración cuasipermanente (XS3, wmax = 0.1 mm, CE Art. 27 Tabla 27.2) — ψ2,Qa 0.3 (CYPE) "
              "y 0.8 (ROM)", "QP = G + ψ2,Qa·Qa (bolardo ψ2 = 0). ψ2 umbral = valor de ψ2,Qa a partir del cual "
                             "fisura (σct > fctm) / falla wk > 0.1 mm con el armado dispuesto.")
    sls = {}
    for r in V["tables"]["sls"]:
        if r["check"].startswith("wk <= 0.1") and "sensibilidad" not in r["check"]:
            sls[(r["member"], r["section"], r["case"])] = r
    sct = {}
    for r in V["tables"]["sls"]:
        if r["check"].startswith("sigma_ct <= fctm") and "sensibilidad" not in r["check"]:
            sct[(r["member"], r["section"], r["case"])] = r
    thr = {(x["member"], x["face"]): x for x in V["tables"]["psi2_threshold"]}
    rows = []
    for k in BEAMS:
        for face, sec in (("bottom", "sag"), ("top", "hog")):
            t = thr[(k, face)]
            w3, w8 = sls.get((k, sec, "CE-CYPE")), sls.get((k, sec, "CE-ROM"))
            s3, s8 = sct.get((k, sec, "CE-CYPE")), sct.get((k, sec, "CE-ROM"))
            pw = None
            if curves and k in curves["members"] and face == "bottom":
                pw = curves["members"][k]["proposed"]["at"]["0.8"]["wk_mm"]
            elif k in ("VBM", "VBT") and face == "top":
                pw = D["proposal"]["edge"]["top"]["wk_mm"]
            rows.append([k, "inferior (M+)" if face == "bottom" else "superior (M-)", t["M_qp_03"], t["M_qp_08"],
                         t["Mcr_kNm"], s3["demand"] / s3["capacity"] if s3 else None,
                         s8["demand"] / s8["capacity"] if s8 else None, w3["demand"] if w3 else None,
                         w8["demand"] if w8 else None, pw, 0.1, t["psi2_crack_fctm"], t["psi2_crack_fctm_fl"],
                         t["psi2_wk_fail"]])
    hdr = ["Viga", "Cara traccionada", "M_qp ψ2 0.3 (kN·m)", "M_qp ψ2 0.8 (kN·m)", "Mcr (kN·m)",
           "σct/fctm ψ2 0.3", "σct/fctm ψ2 0.8", "wk ψ2 0.3 (mm)", "wk ψ2 0.8 (mm) actual", "wk ψ2 0.8 (mm) propuesta",
           "wmax (mm)", "ψ2 umbral fisura (fctm)", "ψ2 umbral fisura (fctm,fl)", "ψ2 umbral wk > 0.1"]
    last = table(ws, 4, hdr, rows, eta_cols=(6, 7), fmt={3: "0.0", 4: "0.0", 5: "0.0", 8: "0.000", 9: "0.000",
                                                         10: "0.000", 11: "0.0", 12: "0.000", 13: "0.000", 14: "0.000"})
    for i in range(5, last + 1):
        for j in (8, 9, 10):
            v = ws.cell(i, j).value
            if isinstance(v, (int, float)):
                ws.cell(i, j).fill = GOOD if v <= 0.1 else BAD
    ws.cell(last + 1, 1, "wk = 0: la sección no fisura (σct ≤ fctm). Propuesta: vigas interiores 8Ø20 alma + 4Ø16 "
                         "alas; vigas de borde 3Ø25 sup. sobre apoyos (F3, F5).")
    if curves:
        r0 = last + 4
        ws.cell(r0 - 1, 1, "Vigas interiores (ejes 2-6): envolvente de σct/fctm y wk frente a ψ2,Qa (tabla del "
                           "gráfico)").font = Font(bold=True)
        rows = [[psi, curves["layouts"]["provided"]["sigma_ct_fctm"][i], curves["layouts"]["proposed"]["sigma_ct_fctm"][i],
                 curves["layouts"]["provided"]["wk_mm"][i], curves["layouts"]["proposed"]["wk_mm"][i], 0.1]
                for i, psi in enumerate(curves["psi"])]
        last2 = table(ws, r0, ["ψ2,Qa", "σct/fctm 7Ø20 (actual)", "σct/fctm 8Ø20+4Ø16 (propuesta)",
                               "wk 7Ø20 (mm)", "wk 8Ø20+4Ø16 (mm)", "wmax (mm)"], rows,
                      fmt={1: "0.00", 2: "0.000", 3: "0.000", 4: "0.0000", 5: "0.0000", 6: "0.0"})
        ch = ScatterChart()
        ch.title = "Vigas interiores: wk frente a ψ2,Qa (envolvente ejes 2-6)"
        ch.style = 13
        ch.x_axis.title = "ψ2,Qa"
        ch.y_axis.title = "wk (mm)"
        ch.height, ch.width = 9, 18
        xs = Reference(ws, min_col=1, min_row=r0 + 1, max_row=last2)
        for col, colr, dash in ((4, INK2, None), (5, C_ROM, None), (6, C_LIMIT, "dash")):
            s = Series(Reference(ws, min_col=col, min_row=r0, max_row=last2), xs, title_from_data=True)
            color_series(s, colr, "line", 2.0, dash)
            ch.series.append(s)
        ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, 1
        axes_on(ch)
        ws.add_chart(ch, f"P{r0}")
        ch = ScatterChart()
        ch.title = "Vigas interiores: σct/fctm (sección homogeneizada) frente a ψ2,Qa"
        ch.style = 13
        ch.x_axis.title = "ψ2,Qa"
        ch.y_axis.title = "σct/fctm"
        ch.height, ch.width = 9, 18
        lim_col = 7
        ws.cell(r0, lim_col, "Fisuración σct = fctm")
        for i in range(r0 + 1, last2 + 1):
            ws.cell(i, lim_col, 1.0)
        for col, colr, dash in ((2, INK2, None), (3, C_ROM, None), (lim_col, C_LIMIT, "dash")):
            s = Series(Reference(ws, min_col=col, min_row=r0, max_row=last2), xs, title_from_data=True)
            color_series(s, colr, "line", 2.0, dash)
            ch.series.append(s)
        ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, 1
        axes_on(ch)
        ws.add_chart(ch, f"P{r0 + 20}")

    # ---- Propuesta ---------------------------------------------------------------------------
    ws = wb.create_sheet("Propuesta")
    title(ws, "Propuesta de armado final (base ROM, Mode.CODIGO, XS3): dispuesto (Anejo 10) frente a propuesto",
          "η de cada comprobación después del cambio (verificación de vigas.py con el armado propuesto; pilotes: "
          "armado dispuesto + confinamiento de cabeza).")
    import armado_vigas as AV
    geo = {"inner": AV.geometry_T(), "end": AV.geometry_L(+1), "edge": AV.geometry_R()}
    stw = {"inner": AV.Stirrups(10.0, 3, 100.0), "end": AV.Stirrups(10.0, 3, 100.0), "edge": AV.Stirrups(8.0, 2, 150.0)}
    rows = []
    pk = D["pile_key"]
    for p in PILES:
        s = P["summary"][p]
        rows.append([p, "pilote", "12Ø25 + 3cØ10 c/150", f"12Ø25 + 3cØ10 c/150; c/{PILE_TIE_S_NEW:g} en "
                     f"{PILE_TIE_ZONE_M:.2f} m sup.", pk["As_prov"], pk["As_prov"], pk["kg_per_m"],
                     pk["kg_per_m"], f"+{pk['extra_tie_kg_per_pile']:.1f} kg/pilote (cercos)",
                     _max(s["nm1"], s["nm2"], s["nm_stations"]), s["shear"], None, s["detailing"], s["sls_crack"],
                     _max(s["sls_sigma_s"], s["sls_sigma_c_Ec_eff"]), None,
                     next(r for r in D["piles"] if r["pile"] == p)["verdict_final"],
                     f"σc/0.6fck Ecm {s['sls_sigma_c']:.3f} -> confinamiento" if s["sls_sigma_c"] > 1 else ""])
    for b in D["beams"]:
        g = b["group"]
        pg = D["proposal"][g]
        fl = b["final"]
        for face in ("top", "bottom"):
            pf = pg[face]
            wl = pg["weight_long_kg_m"]
            rows.append([b["member"], "sup." if face == "top" else "inf.", pf["provided"],
                         pf["proposed"] + (" (sobre apoyos)" if g == "edge" and face == "top" and pf["change"] else ""),
                         pf["provided_As_cm2"], pf.get("proposed_As_cm2", pf["provided_As_cm2"]),
                         wl["provided"] if face == "top" else None, wl["proposed"] if face == "top" else None,
                         "CAMBIO" if pf["change"] else "sin cambio",
                         fl["flexión ELU"], fl["cortante"], fl["torsión"], fl["disposiciones"], fl["fisuración ELS"],
                         fl["tensiones ELS"], fl["flecha"], b["verdict_final"] + (" (cond.)" if g == "edge" else ""),
                         "; ".join(f"{k} {v:.3f}" for k, v in b["final_info"].items())])
        stv = pg["stirrups"]
        rows.append([b["member"], "estribos", stv["provided"], stv["proposed"], None, None,
                     round(stw[g].weight(geo[g]), 2), round(stw[g].weight(geo[g]), 2),
                     "CAMBIO" if stv["change"] else "sin cambio", None, stv.get("provided_eta"), None, None, None,
                     None, None, "", ""])
    hdr = ["Elemento", "Cara / zona", "Dispuesto (Anejo 10)", "Propuesto (final)", "As disp. (cm²)", "As prop. (cm²)",
           "kg/m disp. (fila sup.: toda la long.; estribos: alma)", "kg/m prop.", "Cambio", "η flexión / N-M",
           "η cortante", "η torsión", "η disposiciones", "η fisuración", "η tensiones", "η flecha", "Veredicto final",
           "Información / sensibilidad"]
    last = table(ws, 4, hdr, rows, eta_cols=tuple(range(10, 17)), verdict_cols=(17,),
                 fmt={5: "0.00", 6: "0.00", 7: "0.00", 8: "0.00"}, widths={3: 24, 4: 30, 9: 26, 18: 50})
    for i in range(5, last + 1):
        if ws.cell(i, 9).value == "CAMBIO":
            ws.cell(i, 9).fill = WARN
            ws.cell(i, 4).font = Font(bold=True)
    ws.freeze_panes = "C5"
    r0 = last + 2
    c = D["confinement"]
    notes = [
        "Pilotes (F1): σc ≤ 0.6fck se cubre con 'otras medidas' (A19.7.2(2)): cercos Ø10 c/100 en los 0.40 m "
        "superiores. Estimación indicativa del confinamiento (A19.3.1.9 con α = αn·αs): "
        + "; ".join(f"s = {x['s_mm']:.0f} mm: σ2 = {x['sigma2_MPa']:.2f} MPa, fck,c = {x['fck_c']:.1f} MPa, "
                    f"0.6·fck,c = {x['limit_0.6fckc']:.1f} MPa" for x in c)
        + f" frente a σc,máx = {pk['sigma_c_max'] * 30:.1f} MPa.",
        "Vigas interiores (F3): kg/m longitudinal 32.1 -> 40.8; el resto de comprobaciones mejora o no cambia.",
        "Vigas extremas (F4): sin cambio; sensibilidad del apoyo de las alveoplacas en el ala: superior 4Ø25 "
        f"(η cordón {D['proposal']['end_sensitivity']['top']['provided_eta']:.3f} -> "
        f"{D['proposal']['end_sensitivity']['top']['proposed_eta']:.3f}).",
        "Vigas de borde (F5): 3Ø25 sup. sobre apoyos condicionado al reparto de carga del modelo SAP (lámina "
        "ortótropa); kg/m calculado con la barra en toda la longitud (cota superior).",
    ]
    for i, t in enumerate(notes):
        ws.cell(r0 + i, 1, t)

    # ---- Comparacion_CYPE --------------------------------------------------------------------
    ws = wb.create_sheet("Comparacion_CYPE")
    title(ws, "Paridad con CYPECAD (Anejo 10): valores propios frente a los impresos por CYPE",
          "PASS = dentro de la tolerancia de validate_codigo.py (0.5 % o media unidad del último dígito impreso); "
          "INFO = alternativa del Código o valor informativo (no es de paridad).")
    rows = [[r["element"], r["quantity"], r["ours"], r["cype"], r["diff_pct"], r["status"], r["kind"], r["src"],
             r["note"]] for r in D["parity"]]
    last = table(ws, 4, ["Elemento", "Magnitud", "Propio", "CYPE", "Dif. %", "Estado", "Tipo", "Fuente", "Nota"],
                 rows, verdict_cols=(6,), fmt={3: "0.000", 4: "0.000", 5: "0.00"},
                 widths={2: 60, 7: 30, 8: 36, 9: 40})
    ws.freeze_panes = "C5"
    npass = sum(1 for r in D["parity"] if r["status"] == "PASS")
    n_el = {e: sum(1 for r in D["parity"] if r["element"] == e and r["status"] == "PASS") for e in ("Pilotes", "Vigas")}
    ws.cell(3, 1, f"{npass} filas de paridad PASS de {len(D['parity'])} (pilotes {n_el['Pilotes']}, vigas "
                  f"{n_el['Vigas']}); FAIL: {sum(1 for r in D['parity'] if r['status'] == 'FAIL')}.")
    # chart data block: PASS rows with both values > 0 (log axes)
    cc = 13
    ws.cell(4, cc, "|CYPE| datos CYPE")
    ws.cell(4, cc + 1, "Datos de CYPE (reproducción exacta)")
    ws.cell(4, cc + 2, "|CYPE| fuerzas SAP")
    ws.cell(4, cc + 3, "Fuerzas SAP (≤ 5 %)")
    ws.cell(4, cc + 4, "1:1 x")
    ws.cell(4, cc + 5, "1:1")

    def _pts(kind):
        return [(abs(r["cype"]), abs(r["ours"])) for r in D["parity"] if r["status"] == "PASS" and r["kind"] == kind
                and isinstance(r["cype"], (int, float)) and abs(r["cype"]) > 1e-6 and abs(r["ours"]) > 1e-6]
    pp, vv = _pts(PARITY_EXACT), _pts(PARITY_SAP)
    for i, (a, b) in enumerate(pp, 5):
        ws.cell(i, cc, a)
        ws.cell(i, cc + 1, b)
    for i, (a, b) in enumerate(vv, 5):
        ws.cell(i, cc + 2, a)
        ws.cell(i, cc + 3, b)
    lo = min(min(x) for x in pp + vv)
    hi = max(max(x) for x in pp + vv)
    ws.cell(5, cc + 4, lo)
    ws.cell(5, cc + 5, lo)
    ws.cell(6, cc + 4, hi)
    ws.cell(6, cc + 5, hi)
    ch = ScatterChart()
    ch.title = "Paridad: propio frente a CYPE (valores absolutos, escala log)"
    ch.style = 13
    ch.x_axis.title = "|CYPE|"
    ch.y_axis.title = "|propio|"
    ch.height, ch.width = 12, 16
    for (xc, yc, n, colr) in ((cc, cc + 1, len(pp), C_ROM), (cc + 2, cc + 3, len(vv), C_CECY)):
        s = Series(Reference(ws, min_col=yc, min_row=4, max_row=4 + n), Reference(ws, min_col=xc, min_row=5,
                                                                              max_row=4 + n), title_from_data=True)
        s.marker.symbol = "circle"
        s.marker.size = 6
        s.marker.graphicalProperties.solidFill = colr.lstrip("#")
        s.marker.graphicalProperties.line.solidFill = "FFFFFF"
        s.graphicalProperties.line.noFill = True
        ch.series.append(s)
    s = Series(Reference(ws, min_col=cc + 5, min_row=4, max_row=6), Reference(ws, min_col=cc + 4, min_row=5, max_row=6),
               title_from_data=True)
    color_series(s, INK2, "line", 1.25)
    ch.series.append(s)
    ch.x_axis.scaling.logBase = 10
    ch.y_axis.scaling.logBase = 10
    axes_on(ch)
    ws.add_chart(ch, "T4")

    # ---- Supuestos ---------------------------------------------------------------------------
    ws = wb.create_sheet("Supuestos")
    title(ws, "Supuestos, bases y cláusulas", "Decisiones F0-F8, hipótesis de pilotes.py y vigas.py, materiales y "
                                              "cláusulas del Código Estructural (Anejo 19 = EN 1992-1-1 + AN España).")
    rows = [[d[0], "Decisión final", d[1] + ": " + d[2]] for d in D["decisions"]]
    rows += [[f"P{i}", "Pilotes (pilotes.py)", a] for i, a in enumerate(P["meta"]["assumptions"], 1)]
    rows += [[f"V{i}", "Vigas (vigas.py)", a] for i, a in enumerate(V["meta"]["assumptions"], 1)]
    rows += [["M1", "Materiales vigas", "HA-35/F/20/XS3 (árido calizo): " + ", ".join(
        f"{k} = {v}" for k, v in V["meta"]["materials"].items())],
             ["M2", "Materiales pilotes", "HA-50 prefabricado, B500SD, recubrimiento 50 mm (XS3), Ecm (caliza) "
                                          "33550 MPa, fctm 4.07 MPa; φef = 1.75 (dato de CYPE)."],
             ["M3", "Combinaciones", "ELU ELU01-22 (CYPE) / ELR01-22 (ROM); característica ELS01-08 (todas ψ = 1); "
                                     "cuasipermanente QP = G + ψ2,Qa·Qa (bolardo ψ2 = 0; 0.5 sensibilidad)."],
             ["M4", "Listado CYPE", P["meta"]["cype_listing_note"]]]
    rows += [[f"C-P {k}", "Cláusula pilotes", v] for k, v in P["meta"]["clauses"].items()]
    rows += [[f"C-V {k}", "Cláusula vigas", v] for k, v in V["meta"]["clauses"].items()]
    rows += [["R1", "Referencias", "Código Estructural (RD 470/2021) Anejo 19 y Art. 27 Tabla 27.2; ROM 2.0-11 "
                                   "Tabla 4.6.4.1; ROM 0.2-90; Anejo 10 del proyecto (CYPECAD); "
                                   "diseno/python/investigacion/codigo_estructural_spec.md (C1-C21) y sap_design_spec.md."]]
    last = table(ws, 4, ["Id", "Tema", "Texto"], rows, widths={1: 10, 2: 22, 3: 160}, wrap_cols=(3,))
    for i in range(5, last + 1):
        n = len(str(ws.cell(i, 3).value or ""))
        ws.row_dimensions[i].height = 15 * max(1, math.ceil(n / 150))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =============================================================================================
# markdown
# =============================================================================================
def write_md(D: dict, P: dict, V: dict, path: Path = MD) -> Path:
    pk = D["pile_key"]
    pin, pend, pedg, sens = (D["proposal"][k] for k in ("inner", "end", "edge", "end_sensitivity"))
    bm = {b["member"]: b for b in D["beams"]}
    cur = D["curves"]
    L = [
        "# Diseño final — Muelle de Trasmallo, módulo 40 m (Código Estructural 2021)",
        "",
        "*Final design of the 40 m module to the Spanish Código Estructural (CE 2021, Anejo 19 = EN 1992-1-1 with "
        "the Spanish values). Generated by `python3 diseno/python/run_diseno.py` from `diseno/python/output/pilotes.json` and "
        "`vigas.json` (SAP2000 v27.1 forces). Workbook: `diseno/python/output/Diseno_Codigo_Estructural.xlsx`; figures: "
        "`diseno/python/figuras/d1-d8`.*",
        "",
        "**Base del veredicto / verdict basis (F0):** Mode.CODIGO (sección estricta al Código) + base ROM "
        "(ROM 2.0-11 Tabla 4.6.4.1: ψ0,Qa = 1.0, ψ2,Qa = 0.8, tiro de bolardo ψ2 = 0) + XS3, wmax = 0.1 mm "
        "(CE Art. 27 Tabla 27.2) en combinación cuasipermanente. La base 'CYPE' (hipótesis del Anejo 10: ψ0,Qa 0.7, "
        "ψ2,Qa 0.3) se usa para la paridad con CYPECAD, reproducida exactamente (capacidades 0.00 %).",
        "",
        "## 1. Cuadro de armado final / reinforcement schedule",
        "",
        "| Elemento | Long. superior | Long. inferior | Lateral / piel | Transversal | Cambio frente al Anejo 10 | "
        "η det. | Comprobación determinante | Veredicto |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for s in D["schedule"]:
        L.append(f"| {s['element']} | {s['top']} | {s['bottom']} | {s['side']} | {s['transverse']} | {s['change']} | "
                 f"{s['eta']:.3f} | {s['gov']} | **{s['verdict']}** |")
    c100 = next(x for x in D["confinement"] if x["s_mm"] == PILE_TIE_S_NEW)
    L += [
        "",
        "η = solicitación / capacidad del caso final (ROM/codigo). Pilotes: 14 ud. iguales (prefabricados); "
        "vigas interiores 5 ud.; extremas 2 ud.; de borde 2 ud. por módulo.",
        "",
        "## 2. Justificación por elemento / rationale",
        "",
        f"**Pilotes (F1, F2).** ELU cumple en las tres bases/modos: η N-M 2º orden máx "
        f"{pk['nm2_max']['CE-ROM']['eta']:.3f} (cabeza {pk['nm2_max']['CE-ROM']['pile']}, ROM/codigo), "
        f"{pk['nm2_max']['CE-CYPE']['eta']:.3f} (CYPE/codigo) y {pk['nm2_max']['CYPE']['eta']:.3f} (paridad CYPE; "
        f"CYPE imprime 0.913 con sus propias fuerzas, reproducido exactamente); cortante {pk['shear_max']:.3f}; "
        f"As = {pk['As_prov']:.2f} cm² ≤ 0.04·Ac = 64 cm². As,nec = {pk['As_req_max']['As_req_cm2']:.2f} cm² "
        f"({pk['As_req_max']['pile']} {pk['As_req_max']['position']}, {pk['As_req_max']['combo']}) frente a "
        f"{pk['As_prov']:.2f} dispuesta ({pk['As_req_max']['As_req_cm2'] / pk['As_prov']:.2f}): 12Ø25 es necesaria y "
        "suficiente; los pilotes de tierra necesitarían ~39-46 cm² (16Ø20), pero se mantiene un pilote prefabricado "
        f"uniforme. Sin fisurar en la combinación cuasipermanente (η máx {pk['crack_max']:.3f}, σct < fctm). "
        f"La tensión característica σc ≤ 0.6·fck (A19.7.2(2), XS) se supera con Ecm en la cabeza de {', '.join(pk['sigma_c_fail'])} "
        f"(máx {pk['sigma_c_max']:.3f}); con Ec,eff para la parte sostenida todas cumplen (máx "
        f"{pk['sigma_c_eff_max']:.3f}). A19.7.2(2) admite 'otras medidas', como el confinamiento con armadura "
        f"transversal: **se reduce la separación de los cercos a Ø10 c/{PILE_TIE_S_NEW:g} en los "
        f"{PILE_TIE_ZONE_M:.2f} m superiores de cada pilote** (zona extrema, A19.9.5.3(4)); +{pk['extra_tie_kg_per_pile']:.1f} kg de acero por pilote. "
        f"Estimación indicativa (A19.3.1.9 con α = αn·αs): σ2 ≈ {c100['sigma2_MPa']:.1f} MPa, fck,c ≈ "
        f"{c100['fck_c']:.0f} MPa, 0.6·fck,c ≈ {c100['limit_0.6fckc']:.1f} MPa > σc,máx ≈ {pk['sigma_c_max'] * 30:.1f} MPa. "
        f"Tiro de bolardo ψ2 = 0 (F2): con 0.5 fisurarían {len(pk['tb05_fail'])} pilotes "
        f"({', '.join(pk['tb05_fail'])}; wk hasta {pk['tb05_max_wk']:.3f} mm > 0.1) y ninguna armadura con "
        "As ≤ As,máx lo resuelve: hipótesis documentada.",
        "",
        f"**Vigas interiores, ejes 2-6 (F3).** ELU flexión máx {max(bm[k]['flexión ELU[CE-ROM]'] for k in INNER):.2f}, "
        f"cortante con suspensión de las alveoplacas máx {max(bm[k]['cortante[CE-ROM]'] for k in INNER):.2f} "
        "(cot θ óptimo), disposiciones y flecha cumplen. Con ψ2,Qa = 0.8 el momento cuasipermanente (137-152 kN·m) "
        "supera Mcr = 125 kN·m y wk = 0.155-0.172 mm > 0.1 mm con el 7Ø20 inferior dispuesto, que sólo cumple para "
        "ψ2,Qa ≤ 0.58-0.70" + (f" (umbral mínimo {cur['threshold_provided']:.3f}, ejes 2 y 6)" if cur else "") +
        f". **Armadura inferior: {pin['bottom']['proposed']} ({pin['bottom']['proposed_As_cm2']:.2f} cm²) en lugar de "
        f"5Ø20 + 2Ø20 ({pin['bottom']['provided_As_cm2']:.2f} cm²)**: wk = {pin['bottom']['wk_mm']:.3f} mm con ψ2 0.8"
        + ((f"; cumple hasta ψ2,Qa ≈ {cur['threshold_proposed']:.2f}" if cur["threshold_proposed"]
            else "; cumple para cualquier ψ2,Qa ≤ 1") if cur else "")
        + ". Superior 5Ø20 y estribos (cerco Ø10 + rama Ø10 + cerco de alas Ø10, c/100) sin cambio. Con el ψ2 0.3 "
          "de CYPE el armado dispuesto cumple (como obtuvo CYPE).",
    ]
    L += [
        "",
        f"**Vigas extremas, ejes 1 y 7 (F4).** CUMPLE con el armado dispuesto: determinante el cordón traccionado por "
        f"torsión {max(bm[k]['torsión[CE-ROM]'] for k in ('VT1', 'VT7')):.3f} (flexión "
        f"{max(bm[k]['flexión ELU[CE-ROM]'] for k in ('VT1', 'VT7')):.3f}). Sensibilidad no incluida en el veredicto: "
        "si las alveoplacas apoyan en la única ala a ~0.42 m del eje del pilote (ni SAP ni CYPE lo modelan: 'la "
        "alveoplaca apoya en el eje de los pilotes'), la torsión de equilibrio de ~56-61 kN·m en las caras de los "
        f"pilotes lleva el cordón a {D['proposal']['end_sensitivity_eta']:.2f}; si se confirma el detalle, superior "
        f"{sens['top']['proposed']} en lugar de 5Ø20 (η {sens['top']['proposed_eta']:.3f}).",
        "",
        f"**Vigas de borde 25x30 (F5).** En el modelo SAP comparten la carga del forjado con la lámina ortótropa "
        "(V ~2x y M +13-21 % frente a CYPE): fisuración en apoyos wk ≈ 0.31 mm > 0.1 mm (momento negativo) y cortante "
        f"con θ = 45° 1.04-1.08 (cumple con cot θ óptimo: {max(bm[k]['cortante[CE-ROM]'] for k in ('VBM', 'VBT')):.2f}). "
        f"**Superior {pedg['top']['proposed']} sobre los apoyos** (wk {pedg['top']['wk_mm']:.3f} mm), condicionado a "
        "ese reparto de carga; alternativa: viga de borde de más canto. Si sólo aplica la carga por bandas de CYPE, "
        "las barras dispuestas (2Ø16) bastan en ELU.",
        "",
        "**Decalaje (F6, información).** Con cot θ = 2 la regla de decalaje (A19.6.2.3(7) / 9.2.1.3) en la cara mar "
        "del pilote, limitada al momento del nudo (eje del pilote), elevaría el cordón superior a 1.03-1.18: no cortar "
        "las barras superiores dentro de al = z·cot θ/2 ≈ 0.43 m de las caras de los pilotes y anclarlas por completo "
        "sobre los pilotes (con cot θ = 1, como CYPE, el efecto es menor).",
        "",
        "## 3. Cuantías / steel quantities",
        "",
        "| Elemento | Dispuesto (Anejo 10) | Propuesto | kg/m long. dispuesto | kg/m long. propuesto |",
        "|---|---|---|---|---|",
        f"| Pilote (fuste) | 12Ø25 + 3cØ10/150 | igual + cØ10/100 en 0.40 m sup. | {pk['kg_per_m']:.1f} (long. + cercos) | "
        f"{pk['kg_per_m']:.1f} + {pk['extra_tie_kg_per_pile']:.1f} kg/pilote |",
        f"| Viga interior | 5Ø20 / 5Ø20+2Ø20 | 5Ø20 / {pin['bottom']['proposed']} | {pin['weight_long_kg_m']['provided']:.1f} | "
        f"{pin['weight_long_kg_m']['proposed']:.1f} |",
        f"| Viga extrema | 5Ø20 / 5Ø20+1Ø20 | sin cambio | {pend['weight_long_kg_m']['provided']:.1f} | "
        f"{pend['weight_long_kg_m']['proposed']:.1f} |",
        f"| Viga de borde | 2Ø16 / 2Ø16 | 3Ø25 sup. (apoyos) / 2Ø16 | {pedg['weight_long_kg_m']['provided']:.1f} | "
        f"{pedg['weight_long_kg_m']['proposed']:.1f} (3Ø25 en toda la longitud, cota superior) |",
        "",
        "## 4. Veredicto por base / modo (armado del Anejo 10)",
        "",
        "| Elemento | CYPE/cype (paridad) | CYPE/codigo | ROM/codigo (armado actual) | Final |",
        "|---|---|---|---|---|",
    ]
    fails = {c: [r["pile"] for r in D["piles"] if r[f"verdict[{c}]"] != "CUMPLE"] for c in ("CYPE", "CE-CYPE", "CE-ROM")}
    L.append("| Pilotes | " + " | ".join(("NO CUMPLE " + ", ".join(fails[c]) + " (σc ≤ 0.6fck, Ecm)") if fails[c]
                                        else "CUMPLE" for c in ("CYPE", "CE-CYPE", "CE-ROM"))
             + " | CUMPLE (cercos de confinamiento) |")
    for k in BEAMS:
        b = bm[k]
        L.append(f"| {k} {b['name']} | {b['verdict[CYPE]']} | {b['verdict[CE-CYPE]']} | {b['verdict[CE-ROM]']} | "
                 f"{b['verdict_final']}{' (condicionado)' if b['group'] == 'edge' else ''} (η {b['eta_final']:.3f}, "
                 f"{b['gov_final']}) |")
    L += [
        "",
        "Pilotes: σc ≤ 0.6fck no depende de la base (combinaciones características ELS01-08); CYPE no comprueba "
        "este límite (F8). Vigas de borde: con el reparto de carga de SAP no cumplen fisuración ni con ψ2 0.3 "
        "(umbral ψ2 ≈ 0.25) y en el caso CYPE/cype el cortante con θ = 45° es 1.04-1.08.",
        "",
        "## 5. Puntos abiertos / open points",
        "",
        f"1. **Confinamiento de la cabeza de los pilotes (F1).** Detallar los cercos Ø10 c/{PILE_TIE_S_NEW:g} en los "
        f"{PILE_TIE_ZONE_M:.2f} m superiores en el plano del prefabricado; la justificación cuantitativa del "
        "confinamiento como 'otra medida' de A19.7.2(2) es indicativa (A19.3.1.9); alternativa: HA-60 (fck ≥ 55.5 MPa).",
        f"2. **ψ2 del tiro de bolardo = 0 (F2).** Confirmar con la Autoridad Portuaria (registros de amarre). Con 0.5 "
        f"fisuran {len(pk['tb05_fail'])} pilotes (wk hasta {pk['tb05_max_wk']:.3f} mm) y no hay solución con "
        "As ≤ 0.04·Ac; alternativa: aceptar wk ≤ 0.2 mm para ese caso.",
        "3. **Apoyo de las alveoplacas en las vigas extremas (F4).** Verificar el detalle de apoyo (eje del pilote "
        "o ala); si apoyan en el ala: superior 4Ø25 en lugar de 5Ø20.",
        "4. **Reparto de carga en las vigas de borde (F5).** La propuesta 3Ø25 sup. depende del modelo del forjado "
        "(lámina ortótropa de SAP frente a bandas de CYPE); confirmar el modelo o dar más canto a la viga de borde.",
        "5. **Corte de barras superiores (F6).** No cortar dentro de al ≈ 0.43 m de las caras de los pilotes; "
        "anclaje completo sobre los pilotes.",
        "6. **Nueva armadura inferior de las vigas interiores (F3).** 8Ø20 en el alma con separación libre 31 mm "
        "(≥ dg + 5 = 25 mm): comprobar el hormigonado y el anclaje de las barras en los nudos con los pilotes; las "
        "4Ø16 de las alas quedan dentro del cerco de alas.",
        "7. **Comprobación EC2 de SAP2000 (F7).** La ratio prevista de los pilotes en SAP (~0.71, estimación; pendiente de los resultados de SAP) no es comparable con nuestro "
        "2º orden (0.93-0.99): rige nuestro veredicto; SAP no comprueba wk XS3, ν1/fywd del CE, As,min española ni la "
        "armadura de suspensión (diseno/sap/GUIA_SAP_DISENO.md).",
        "8. **Anejo 10 (F8).** El 'Aprov.' del listado de CYPE = 1.014 × η del rayo (sin explicar); errata "
        "aritmética 410/4.65 = 88 kN (no 83); CYPE no comprobó la suspensión de las alas, σc característica ni la "
        "torsión del ala de las vigas extremas.",
        "9. **φef = 1.75** (dato de CYPE) en el 2º orden de los pilotes es conservador (A19.5.8.4 con ψ2,bolardo = 0 "
        "daría φef ≈ 0); no se ha reducido.",
        "",
        "## 6. Ficheros / files",
        "",
        "- `diseno/python/output/Diseno_Codigo_Estructural.xlsx`: Resumen, Pilotes_ELU, Pilotes_As, Pilotes_ELS, "
        "Vigas_Flexion, Vigas_Cortante, Vigas_Fisuracion, Propuesta, Comparacion_CYPE, Supuestos (gráficos nativos).",
        "- `diseno/python/output/pilotes.md|json`, `diseno/python/output/vigas.md|json`: comprobaciones detalladas.",
        "- `diseno/python/figuras/d1_pilotes_eta.png` ... `d8_interaccion_pilote.png`.",
        "- `diseno/sap/`: modelo y guía para la comprobación en SAP2000 (información, F7).",
    ]
    if D["warnings"]:
        L += ["", "## Avisos del generador", ""] + [f"- {w}" for w in D["warnings"]]
    path.write_text("\n".join(x for x in L) + "\n", encoding="utf-8")
    return path


# =============================================================================================
# figures
# =============================================================================================
def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9.5, "axes.titlesize": 10.5, "axes.titleweight": "bold",
        "axes.labelsize": 9.5, "axes.edgecolor": AXISC, "axes.labelcolor": INK2, "axes.titlecolor": INK,
        "xtick.color": INK2, "ytick.color": INK2, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "axes.grid.axis": "y", "grid.color": GRID, "grid.linewidth": 0.8, "grid.linestyle": "-",
        "legend.frameon": False, "legend.fontsize": 9, "figure.dpi": 150, "savefig.dpi": 150,
        "savefig.bbox": "tight", "savefig.facecolor": "white", "figure.facecolor": "white",
        "lines.solid_capstyle": "round", "lines.solid_joinstyle": "round"})
    return plt


def fig_d1(D: dict, path: Path) -> None:
    plt = _mpl()
    fig, axs = plt.subplots(1, 2, figsize=(11, 4.3), sharey=True)
    w = 0.26
    for ax, group, ttl in ((axs[0], SEA, "Pilotes lado mar (sea piles)"),
                           (axs[1], LAND, "Pilotes lado tierra (land piles)")):
        rows = [next(r for r in D["piles"] if r["pile"] == p) for p in group]
        x = np.arange(len(rows))
        for j, case in enumerate(("CYPE", "CE-CYPE", "CE-ROM")):
            vals = [r[f"nm2[{case}]"] for r in rows]
            xs = x + (j - 1) * w
            ax.bar([a for a, v in zip(xs, vals) if v is not None], [v for v in vals if v is not None], w * 0.92,
                   color=CASE_COLOR[case], label=VARIANT_LABEL[case], zorder=2, edgecolor="white", linewidth=0.6)
            for a, v in zip(xs, vals):
                if v is None:
                    ax.text(a, 0.02, "n/a", rotation=90, ha="center", va="bottom", fontsize=7, color=MUTED)
        ax.set_xticks(x, [f"{r['pile']}\neje {r['axis']}" for r in rows])
        ax.set_title(ttl, loc="left")
        ax.set_ylim(0, 1.15)
        ax.set_xlim(-0.6, len(rows) - 0.4)
        ax.axhline(1.0, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)), zorder=3)
    axs[0].set_ylabel("η N-M 2º orden (2nd order)")
    top = max(D["piles"], key=lambda r: r["nm2[CE-ROM]"] or 0)
    i = SEA.index(top["pile"]) if top["pile"] in SEA else None
    if i is not None:
        axs[0].annotate(f"{top['nm2[CE-ROM]']:.3f}", (i + w, 1.0), xytext=(0, 5),
                        textcoords="offset points", ha="center", va="bottom", fontsize=8.5, color=INK)
    axs[1].text(len(LAND) - 0.45, 1.012, "límite η = 1", color=INK2, fontsize=8.5, va="bottom", ha="right")
    h, l_ = axs[0].get_legend_handles_labels()
    fig.legend(h, l_, loc="upper center", ncol=3, bbox_to_anchor=(0.5, 1.04))
    fig.text(0.5, -0.03, "n/a = N.P. en ese modo: tracción o λ ≤ λlim en todas las combinaciones (2º orden "
                         "despreciable). Máximo de cabeza / pie / arranque por pilote.", ha="center", fontsize=8,
             color=MUTED)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_d2(D: dict, P: dict, path: Path) -> None:
    plt = _mpl()
    asr = {}
    for r in P["tables"]["as_required"]:
        asr[r["pile"]] = max(asr.get(r["pile"], 0), r["As_req_cm2"])
    order = list(SEA) + list(LAND)
    x = np.concatenate([np.arange(7), np.arange(7) + 7.8])
    vals = [asr[p] for p in order]
    fig, ax = plt.subplots(figsize=(10.5, 4.3))
    ax.bar(x, vals, 0.55, color=C_ROM, zorder=2, label="As necesaria (N-M 1er y 2º orden, ROM/codigo)")
    prov = D["pile_key"]["As_prov"]
    ax.axhline(prov, color=INK2, lw=1.8, zorder=3, label=f"As dispuesta 12Ø25 = {prov:.1f} cm²")
    ax.axhline(64.0, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)), zorder=3, label="As,máx = 0.04·Ac = 64 cm²")
    ax.axhline(50.27, color=MUTED, lw=1.2, ls=(0, (2, 2)), zorder=3, label="16Ø20 = 50.3 cm² (opción pilotes tierra)")
    imax = int(np.argmax(vals))
    ax.annotate(f"{vals[imax]:.2f}", (x[imax], vals[imax]), xytext=(0, 3), textcoords="offset points", ha="center",
                fontsize=8.5, color=INK)
    ax.set_xticks(x, [f"{p}\neje {AXIS_OF[p]}" for p in order])
    ax.set_ylim(0, 72)
    ax.set_ylabel("As longitudinal (cm²)")
    ax.text(3, 68.5, "lado mar (sea)", ha="center", color=INK2, fontsize=9)
    ax.text(10.8, 68.5, "lado tierra (land)", ha="center", color=INK2, fontsize=9)
    ax.set_title("Pilotes: armadura longitudinal necesaria frente a dispuesta (required vs provided)", loc="left")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=2)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_d3(D: dict, path: Path) -> None:
    plt = _mpl()
    fig, axs = plt.subplots(1, 3, figsize=(12.5, 4.4), sharey=True, gridspec_kw={"width_ratios": [9, 9, 2.6]})
    w = 0.26
    rows = D["beams"]
    x = np.arange(len(rows))
    for ax, fam, ttl in ((axs[0], "flexión ELU", "Flexión ELU (bending)"), (axs[1], "cortante", "Cortante (shear)")):
        for j, case in enumerate(("CYPE", "CE-CYPE", "CE-ROM")):
            ax.bar(x + (j - 1) * w, [b[f"{fam}[{case}]"] for b in rows], w * 0.92, color=CASE_COLOR[case],
                   label=VARIANT_LABEL[case], zorder=2, edgecolor="white", linewidth=0.6)
        ax.set_xticks(x, [BEAM_SHORT[b["member"]] for b in rows], fontsize=8.5)
        ax.set_title(ttl, loc="left")
        ax.axhline(1.0, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)), zorder=3)
        ax.set_xlim(-0.6, len(rows) - 0.4)
    ends = [b for b in rows if b["member"] in ("VT1", "VT7")]
    xe = np.arange(len(ends))
    for j, case in enumerate(("CYPE", "CE-CYPE", "CE-ROM")):
        axs[2].bar(xe + (j - 1) * w, [b[f"torsión[{case}]"] for b in ends], w * 0.92, color=CASE_COLOR[case], zorder=2,
                   edgecolor="white", linewidth=0.6)
    axs[2].set_xticks(xe, [BEAM_SHORT[b["member"]] for b in ends], fontsize=8.5)
    axs[2].set_title("Torsión\n(ejes 1, 7)", loc="left")
    axs[2].axhline(1.0, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)), zorder=3)
    axs[2].set_xlim(-0.6, 1.6)
    axs[0].set_ylabel("η = MEd/MRd, VEd/VRd, ...")
    axs[0].set_ylim(0, 1.2)
    vbt = next(b for b in rows if b["member"] == "VBT")
    axs[1].annotate(f"{vbt['cortante[CYPE]']:.2f} (θ 45°)", (len(rows) - 1 - w, vbt["cortante[CYPE]"]),
                    xytext=(-4, 4), textcoords="offset points", ha="right", fontsize=8, color=INK)
    axs[0].text(len(rows) - 0.45, 1.01, "límite η = 1", color=INK2, fontsize=8.5, va="bottom", ha="right")
    h, l_ = axs[0].get_legend_handles_labels()
    fig.legend(h, l_, loc="upper center", ncol=3, bbox_to_anchor=(0.5, 1.05))
    fig.text(0.5, -0.04, "Armado del Anejo 10. CYPE/cype: θ = 45°, sin suspensión; CE: cot θ óptimo en [0.5, 2] + "
                         "armadura de suspensión de las alveoplacas (A19.6.2.1(9)). Torsión en veredicto sólo en las "
                         "vigas extremas.", ha="center", fontsize=8, color=MUTED)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_d4(D: dict, path: Path) -> None:
    plt = _mpl()
    cur = D["curves"]
    psi = np.array(cur["psi"])
    fig, axs = plt.subplots(1, 2, figsize=(11, 4.3))
    lab = {"provided": "7Ø20 inf. (Anejo 10, 21.99 cm²)", "proposed": "8Ø20 + 4Ø16 inf. (propuesta, 33.18 cm²)"}
    col = {"provided": INK2, "proposed": C_ROM}
    for tag in ("provided", "proposed"):
        axs[0].plot(psi, cur["layouts"][tag]["sigma_ct_fctm"], color=col[tag], lw=2, label=lab[tag], zorder=3)
        w = np.array(cur["layouts"][tag]["wk_mm"], float)
        wm = np.where(w > 0, w, np.nan)
        axs[1].plot(psi, wm, color=col[tag], lw=2, label=lab[tag], zorder=3)
        axs[1].plot(psi[w <= 0], np.zeros((w <= 0).sum()), color=col[tag], lw=2, zorder=3)
    axs[0].axhline(1.0, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)))
    axs[0].text(0.01, 1.02, "σct = fctm (fisura / cracks)", color=INK2, fontsize=8.5, va="bottom")
    axs[1].axhline(0.1, color=C_LIMIT, lw=1.4, ls=(0, (5, 3)))
    axs[1].text(0.01, 0.103, "wmax = 0.1 mm (XS3)", color=INK2, fontsize=8.5, va="bottom")
    axs[0].set_ylim(0, 1.35)
    axs[1].set_ylim(0, 0.24)
    for ax, ypos, va in ((axs[0], 0.03, "bottom"), (axs[1], 0.97, "top")):
        for v, t in ((0.3, "ψ2 0.3\n(CYPE/CTE)"), (0.8, "ψ2 0.8\n(ROM)")):
            ax.axvline(v, color=MUTED, lw=1, zorder=1)
            ax.text(v + 0.012, ax.get_ylim()[1] * ypos, t, color=INK2, fontsize=8, va=va)
        ax.set_xlim(0, 1)
        ax.set_xlabel("ψ2,Qa (sobrecarga de uso; bolardo ψ2 = 0)")
    # crack onset: wk jumps from 0 to the first cracked value (drawn as a thin vertical step)
    for tag in ("provided", "proposed"):
        w = np.array(cur["layouts"][tag]["wk_mm"], float)
        k = int(np.argmax(w > 0)) if (w > 0).any() else None
        if k:
            axs[1].plot([psi[k], psi[k]], [0, w[k]], color=col[tag], lw=1, alpha=0.8, zorder=2)
    thr = cur["threshold_provided"]
    axs[1].plot([thr], [0], "o", ms=7, color=INK2, mec="white", mew=1.5, zorder=4, clip_on=False)
    axs[1].annotate(f"7Ø20 fisura en ψ2 = {thr:.3f} (ejes 2 y 6)\ny wk salta a ≈ 0.14 mm", (thr, 0.0),
                    xytext=(0.03, 0.045), textcoords="data", ha="left", fontsize=8.5, color=INK,
                    arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8, shrinkB=4))
    if cur.get("threshold_proposed"):
        tp = cur["threshold_proposed"]
        axs[1].plot([tp], [0.1], "o", ms=7, color=C_ROM, mec="white", mew=1.5, zorder=4)
        axs[1].annotate(f"propuesta:\nwk > 0.1 mm\nsi ψ2 > {tp:.2f}", (tp, 0.1), xytext=(0.645, 0.012),
                        textcoords="data", ha="left", fontsize=8.5, color=INK,
                        arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8, shrinkB=4))
    w8 = {t: np.interp(0.8, psi, np.array(cur["layouts"][t]["wk_mm"], float)) for t in ("provided", "proposed")}
    axs[1].annotate(f"{w8['provided']:.3f} mm", (0.8, w8["provided"]), xytext=(-6, 8), textcoords="offset points",
                    ha="right", fontsize=8.5, color=INK)
    axs[1].annotate(f"{w8['proposed']:.3f} mm", (0.8, w8["proposed"]), xytext=(6, -16), textcoords="offset points",
                    ha="left", fontsize=8.5, color=INK)
    axs[0].set_ylabel("σct / fctm (sección homogeneizada)")
    axs[1].set_ylabel("wk (mm)")
    axs[0].set_title("Criterio de no fisuración (A19.7.1(2))", loc="left")
    axs[1].set_title("Abertura de fisura wk (A19.7.3.4)", loc="left")
    h, l_ = axs[0].get_legend_handles_labels()
    fig.legend(h, l_, loc="upper center", ncol=2, bbox_to_anchor=(0.5, 1.05))
    fig.text(0.5, -0.03, "Vigas interiores ejes 2-6 (T invertida), momento positivo cuasipermanente QP = G + ψ2,Qa·Qa; "
                         "envolvente de las cinco vigas. wk = 0 mientras la sección no fisura.", ha="center",
             fontsize=8, color=MUTED)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def _dim(ax, p0, p1, text, off, orient="h", fs=8):
    """Dimension line between p0 and p1 offset by ``off`` (mm)."""
    kw = dict(arrowprops=dict(arrowstyle="<->", color=MUTED, lw=0.8, shrinkA=0, shrinkB=0))
    if orient == "h":
        y = p0[1] + off
        ax.annotate("", (p0[0], y), (p1[0], y), **kw)
        for x in (p0[0], p1[0]):
            ax.plot([x, x], [p0[1], y + (6 if off > 0 else -6)], color=MUTED, lw=0.6)
        ax.text((p0[0] + p1[0]) / 2, y + (8 if off > 0 else -10), text, ha="center", va="bottom" if off > 0 else "top",
                fontsize=fs, color=INK2)
    else:
        x = p0[0] + off
        ax.annotate("", (x, p0[1]), (x, p1[1]), **kw)
        for y in (p0[1], p1[1]):
            ax.plot([p0[0], x + (6 if off > 0 else -6)], [y, y], color=MUTED, lw=0.6)
        ax.text(x + (8 if off > 0 else -8), (p0[1] + p1[1]) / 2, text, ha="left" if off > 0 else "right",
                va="center", fontsize=fs, color=INK2, rotation=90)


def fig_d5(D: dict, V: dict, path: Path) -> None:
    plt = _mpl()
    from matplotlib.patches import Circle, FancyBboxPatch, Polygon
    prov = [r for r in V["tables"]["reinforcement_provided"] if r["member"] == "VT2"]
    prop = [r for r in D["proposal"]["inner"]["layout"] if r["member"] == "VT2"]
    pin = D["proposal"]["inner"]
    fig, axs = plt.subplots(1, 2, figsize=(12, 5.6))
    wk_prov = max(D["curves"]["members"][k]["provided"]["at"]["0.8"]["wk_mm"] for k in INNER) if D["curves"] else 0.172
    panels = ((axs[0], prov, "Anejo 10 (actual / current)",
               [f"inferior 5Ø20 + 2Ø20 = {pin['bottom']['provided_As_cm2']:.2f} cm²",
                f"wk (ψ2 0.8) = {wk_prov:.3f} mm > 0.1 mm  ✗"]),
              (axs[1], prop, "Propuesta final (proposed)",
               [f"inferior 8Ø20 alma + 4Ø16 alas = {pin['bottom']['proposed_As_cm2']:.2f} cm²",
                f"wk (ψ2 0.8) = {pin['bottom']['wk_mm']:.3f} mm ≤ 0.1 mm  ✓"]))
    for ax, rows, ttl, notes in panels:
        ax.add_patch(Polygon([(-400, 0), (400, 0), (400, 300), (250, 300), (250, 550), (-250, 550), (-250, 300),
                              (-400, 300)], closed=True, fc=CONCRETE, ec=INK2, lw=1.3, zorder=1))
        ax.add_patch(FancyBboxPatch((-195, 55), 390, 440, boxstyle="round,pad=0,rounding_size=14", fill=False,
                                    ec=INK2, lw=1.1, zorder=2))
        ax.add_patch(FancyBboxPatch((-345, 55), 690, 190, boxstyle="round,pad=0,rounding_size=14", fill=False,
                                    ec=INK2, lw=1.1, zorder=2))
        ax.plot([0, 0], [55, 495], color=INK2, lw=1.1, zorder=2)
        for r in rows:
            if r["face"] == "stirrups":
                continue
            changed = ttl.startswith("Propuesta") and r["face"] == "bottom"
            old_bottom = ttl.startswith("Anejo") and r["face"] == "bottom"
            for xb in r["x_mm"]:
                ax.add_patch(Circle((xb, r["y_soffit_mm"]), r["phi"] / 2, fc=C_ROM if changed else INK2,
                                    ec="white", lw=0.8, zorder=4))
            if old_bottom:
                pass
        # labels
        top = next(r for r in rows if r["face"] == "top")
        ax.annotate(top["group"].replace(" sup.", " superior"), (180, 480), xytext=(330, 600), fontsize=8.5,
                    color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8), ha="left")
        bots = [r for r in rows if r["face"] == "bottom"]
        txt = " + ".join(r["group"].replace(" inf.", "").replace("I1Ø20+D1Ø20 alas", "2Ø20 alas") for r in bots)
        ax.annotate(txt, (180, 60), xytext=(260, -150), fontsize=8.5, color=INK, ha="center",
                    arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
        ax.annotate("2Ø10 + 2Ø10 (a 235)", (-335, 235), xytext=(-560, 420), fontsize=8.5, color=INK, ha="left",
                    arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
        ax.annotate("estribos Ø10 c/100:\ncerco alma + rama\n+ cerco de alas", (-195, 400), xytext=(-560, 560),
                    fontsize=8.5, color=INK, ha="left", va="bottom",
                    arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
        _dim(ax, (-400, 0), (400, 0), "800", -60)
        _dim(ax, (-250, 550), (250, 550), "500", 45)
        _dim(ax, (400, 0), (400, 550), "550", 90, "v")
        _dim(ax, (-400, 0), (-400, 300), "300", -45, "v")
        _dim(ax, (250, 300), (400, 300), "150", 40, fs=7.5)
        ax.text(0, -250, "\n".join(notes), ha="center", va="top", fontsize=9, color=INK)
        ax.set_title(ttl, loc="center")
        ax.set_xlim(-600, 620)
        ax.set_ylim(-360, 700)
        ax.set_aspect("equal")
        ax.axis("off")
    fig.suptitle("Vigas interiores ejes 2-6: sección T invertida 50x55+15x30+15x30 (HA-35, rec. 50 mm; cotas en mm)",
                 fontsize=10.5, fontweight="bold", color=INK, y=0.99)
    fig.text(0.5, 0.01, "Superior 5Ø20, armadura lateral y estribos sin cambio. En azul: armadura inferior propuesta.",
             ha="center", fontsize=8.5, color=MUTED)
    fig.tight_layout(rect=(0, 0.03, 1, 0.96))
    fig.savefig(path)
    plt.close(fig)


def fig_d6(path: Path) -> None:
    plt = _mpl()
    from matplotlib.patches import Circle, FancyBboxPatch, Rectangle
    from armado_pilote import PROVIDED
    lay = PROVIDED["fuste"]
    fig, (ax, ax2) = plt.subplots(1, 2, figsize=(11, 5.6), gridspec_kw={"width_ratios": [1.35, 1]})
    ax.add_patch(Rectangle((-200, -200), 400, 400, fc=CONCRETE, ec=INK2, lw=1.3, zorder=1))
    tc = 200 - lay.cover - lay.ties.phi / 2
    ax.add_patch(FancyBboxPatch((-tc, -tc), 2 * tc, 2 * tc, boxstyle="round,pad=0,rounding_size=12", fill=False,
                                ec=INK2, lw=1.1, zorder=2))
    xi = 42.5 + 12.5 + 5
    ax.add_patch(FancyBboxPatch((-xi, -tc + 4), 2 * xi, 2 * tc - 8, boxstyle="round,pad=0,rounding_size=10",
                                fill=False, ec=INK2, lw=1.0, zorder=2))
    ax.add_patch(FancyBboxPatch((-tc + 4, -xi), 2 * tc - 8, 2 * xi, boxstyle="round,pad=0,rounding_size=10",
                                fill=False, ec=INK2, lw=1.0, zorder=2))
    for x, y in lay.coords():
        ax.add_patch(Circle((x, y), lay.phi / 2, fc=INK2, ec="white", lw=0.8, zorder=4))
    _dim(ax, (-200, -200), (200, -200), "400", -45)
    _dim(ax, (200, -200), (200, 200), "400", 45, "v")
    _dim(ax, (-200, 200), (-150, 200), "50", 30, fs=7.5)
    ax.annotate("12Ø25 (4 por cara)\nAs = 58.9 cm², ρ = 3.68 %", (-127.5, 127.5), xytext=(-330, 300), fontsize=8.5,
                color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
    ax.annotate("3 cercos Ø10 (2eØ10+1eØ10):\nperimetral + 2 interiores", (xi, -90), xytext=(80, -330),
                fontsize=8.5, color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
    ax.text(-330, -300, "HA-50, B500SD\nrec. 50 mm (XS3)", fontsize=8.5, color=INK2, va="top")
    ax.set_xlim(-360, 330)
    ax.set_ylim(-380, 380)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("Sección del pilote 40x40 (cotas en mm)", loc="center")
    # elevation (schematic, horizontal scale exaggerated)
    H, Z = 6.70, PILE_TIE_ZONE_M
    ax2.add_patch(Rectangle((-0.2, 0), 0.4, H, fc=CONCRETE, ec=INK2, lw=1.2, zorder=1))
    ax2.add_patch(Rectangle((-0.75, H), 1.5, 0.55, fc="#e1e0d9", ec=INK2, lw=1.2, zorder=1))
    ax2.text(0, H + 0.275, "viga (beam)", ha="center", va="center", fontsize=8.5, color=INK2)
    for xb in (-0.1275, 0.1275):
        ax2.plot([xb, xb], [-0.3, H + 0.45], color=INK2, lw=1.6, zorder=3)
    zs = list(np.arange(0.075, H - Z, 0.15))
    for z in zs:
        ax2.plot([-0.145, 0.145], [z, z], color=MUTED, lw=1.0, zorder=2)
    for z in np.arange(H - Z + 0.05, H, PILE_TIE_S_NEW / 1000):
        ax2.plot([-0.145, 0.145], [z, z], color=C_ROM, lw=1.8, zorder=4)
    ax2.plot([-0.9, 0.9], [0, 0], color=INK2, lw=1.0)
    for xh in np.arange(-0.9, 0.9, 0.12):
        ax2.plot([xh, xh - 0.08], [0, -0.12], color=MUTED, lw=0.7)
    ax2.annotate(f"cercos Ø10 c/{PILE_TIE_S_NEW:g} en {Z:.2f} m\n(confinamiento, propuesta F1)", (0.145, H - Z / 2),
                 xytext=(0.45, H - 1.3), fontsize=8.5, color=INK, arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
    ax2.annotate("cercos Ø10 c/150\n(Anejo 10)", (0.145, 3.0), xytext=(0.45, 3.0), fontsize=8.5, color=INK,
                 va="center", arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8))
    ax2.text(0.45, 0.35, "empotramiento / arranque\n(z = 0, cerco Ø8)", fontsize=8.5, color=INK2, va="bottom")
    ax2.annotate("", (-0.35, 0), (-0.35, H), arrowprops=dict(arrowstyle="<->", color=MUTED, lw=0.8))
    ax2.text(-0.4, H / 2, "6.70 m (longitud libre)", rotation=90, ha="right", va="center", fontsize=8.5, color=INK2)
    ax2.set_xlim(-1.0, 1.55)
    ax2.set_ylim(-0.5, H + 0.8)
    ax2.axis("off")
    ax2.set_title("Alzado esquemático (no a escala)", loc="center")
    fig.suptitle("Pilotes prefabricados P1-P18: armado final (12Ø25 sin cambio; cercos Ø10 c/100 en cabeza)",
                 fontsize=10.5, fontweight="bold", color=INK, y=0.99)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path)
    plt.close(fig)


def fig_d7(D: dict, path: Path) -> None:
    plt = _mpl()
    fig, ax = plt.subplots(figsize=(6.4, 6.0))
    allv, stats = [], []
    for elem, kind, colr, lab in (("Pilotes", PARITY_EXACT, C_ROM, "Pilotes, datos de CYPE"),
                                  ("Vigas", PARITY_EXACT, C_CECY, "Vigas, datos de CYPE"),
                                  ("Pilotes", PARITY_SAP, C_CYPE, "Pilotes, fuerzas SAP (listado 'Aprov.', P3)")):
        pts = [(abs(r["cype"]), abs(r["ours"]), r["diff_pct"]) for r in D["parity"] if r["element"] == elem and
               r["kind"] == kind and r["status"] == "PASS" and isinstance(r["cype"], (int, float))
               and abs(r["cype"]) > 1e-6 and abs(r["ours"]) > 1e-6]
        if not pts:
            continue
        allv += pts
        stats.append((lab, len(pts), max(abs(p[2]) for p in pts if p[2] is not None)))
        ax.scatter([p[0] for p in pts], [p[1] for p in pts], s=26, color=colr, edgecolor="white", linewidth=0.8,
                   label=f"{lab}: {len(pts)}", zorder=3)
    lo = min(min(p[0], p[1]) for p in allv) / 1.5
    hi = max(max(p[0], p[1]) for p in allv) * 1.5
    ax.plot([lo, hi], [lo, hi], color=INK2, lw=1.2, zorder=2, label="1:1")
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlim(lo, hi)
    ax.set_ylim(lo, hi)
    ax.grid(True, which="major", axis="both", color=GRID, lw=0.8)
    ax.set_xlabel("|valor CYPE (Anejo 10)|")
    ax.set_ylabel("|valor propio (diseno/python/codigo)|")
    ax.set_title("Paridad con CYPECAD (valores impresos en el Anejo 10)", loc="left")
    ax.text(0.03, 0.97, "|dif.| máx por grupo:\n" + "\n".join(f"  {lab}: {d:.2f} %" for lab, n, d in stats) +
            "\nDatos de CYPE: tolerancia 0.5 % o ½ unidad del\núltimo dígito impreso; fuerzas SAP: ≤ 5 %",
            transform=ax.transAxes, va="top", fontsize=8, color=INK2)
    ax.legend(loc="lower right")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def fig_d8(inter: dict, path: Path) -> None:
    plt = _mpl()
    fig, (ax0, ax) = plt.subplots(1, 2, figsize=(12, 5.4), gridspec_kw={"width_ratios": [1, 1.35]})
    u = np.array(inter["uniaxial"])
    m = np.array(inter["meridian"]) if inter["meridian"] else None
    cl = inter["cloud"]
    p1, p21, p2 = inter["points"]
    for a in (ax0, ax):
        a.plot(u[:, 1], u[:, 0], color=INK2, lw=2, zorder=3, label="Interacción uniaxial (MRd,x)")
        if m is not None:
            a.plot(m[:, 1], m[:, 0], color=C_ROM, lw=2, zorder=3,
                   label=f"N-Mx-My en la dirección del punto de 2º orden ({inter['meridian_angle_deg']:.1f}°)")
        a.axhline(0, color=AXISC, lw=0.8)
        a.grid(True, axis="both", color=GRID, lw=0.8)
        a.set_xlabel("M = √(MEd,x² + MEd,y²)  (kN·m)")
    ax0.set_ylabel("N (kN, compresión +)")
    zx, zy = (150, 520), (-700, 1500)
    ax0.set_xlim(0, max(u[:, 1].max() * 1.08, zx[1] + 15))
    ax0.add_patch(__import__("matplotlib.patches", fromlist=["Rectangle"]).Rectangle(
        (zx[0], zy[0]), zx[1] - zx[0], zy[1] - zy[0], fill=False, ec=INK2, lw=0.9, zorder=4))
    ax0.scatter([p["M"] for p in (p1, p21, p2)], [p["N"] for p in (p1, p21, p2)], s=10, color=INK, zorder=5)
    ax0.set_title("Diagrama completo (12Ø25, Mode.CODIGO)", loc="left")
    # zoom
    ax.scatter([c["M"] for c in cl if c["check"] == "nm1"], [c["N"] for c in cl if c["check"] == "nm1"], s=14,
               color=MUTED, alpha=0.5, lw=0, zorder=2, label="ELR 1er orden (todos los pilotes y secciones)")
    ax.scatter([c["M"] for c in cl if c["check"] == "nm2"], [c["N"] for c in cl if c["check"] == "nm2"], s=14,
               marker="s", color=MUTED, alpha=0.5, lw=0, zorder=2, label="ELR 2º orden")
    spec = ((p21, C_CECY, "o", (175, 1150), "left"), (p2, C_ROM, "o", (412, 230), "left"),
            (p1, C_CECY, "D", (360, -560), "left"))
    for p, colr, mk, xy, ha in spec:
        ax.plot([0, p["MRd"]], [0, p["NRd"]], color=colr, lw=0.9, alpha=0.6, zorder=2)
        ax.plot(p["M"], p["N"], mk, ms=8.5, color=colr, mec="white", mew=1.8, zorder=5)
        lab = (f"{p['pile']} {p['position'].lower()}, {p['combo']}, {'1er' if p['check'] == 'nm1' else '2º'} orden\n"
               f"N = {p['N']:.0f} kN, M = {p['M']:.0f} kN·m\nη = {p['eta']:.3f}")
        ax.annotate(lab, (p["M"], p["N"]), xytext=xy, textcoords="data", fontsize=8.5, color=INK, ha=ha,
                    va="center", arrowprops=dict(arrowstyle="-", color=MUTED, lw=0.8, shrinkB=5))
    ax.annotate("", (p2["M"], p2["N"]), (p21["M"], p21["N"]),
                arrowprops=dict(arrowstyle="->", color=INK2, lw=1.3, shrinkA=6, shrinkB=6))
    ax.text((p2["M"] + p21["M"]) / 2, p2["N"] + 45, "+ e2 + ei", ha="center", fontsize=8, color=INK2)
    ax.set_xlim(*zx)
    ax.set_ylim(*zy)
    ax.set_title("Detalle: puntos ELU determinantes (ROM/codigo)", loc="left")
    h, l_ = ax.get_legend_handles_labels()
    fig.legend(h, l_, loc="upper center", ncol=2, bbox_to_anchor=(0.5, 1.07), fontsize=8.5)
    fig.suptitle("Pilote 40x40 HA-50, 12Ø25: interacción N-M", fontsize=10.5, fontweight="bold", color=INK,
                 x=0.01, ha="left", y=1.10)
    fig.text(0.5, -0.03, "η = |S|/|R| a lo largo del rayo desde el origen ('mismas excentricidades'; líneas finas "
                         "hasta la capacidad). El 2º orden (A19.5.8.8, curvatura nominal) añade e2 + ei al punto de "
                         "1er orden de la misma combinación.", ha="center", fontsize=8, color=MUTED)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def write_figures(D: dict, P: dict, V: dict, inter: dict, fig_dir: Path = FIG) -> list[Path]:
    fig_dir.mkdir(parents=True, exist_ok=True)
    out = []
    jobs = [("d1_pilotes_eta.png", lambda p: fig_d1(D, p)), ("d2_pilotes_As.png", lambda p: fig_d2(D, P, p)),
            ("d3_vigas_flexion.png", lambda p: fig_d3(D, p)), ("d4_fisuracion_psi2.png", lambda p: fig_d4(D, p)),
            ("d5_seccion_T_actual_propuesta.png", lambda p: fig_d5(D, V, p)),
            ("d6_seccion_pilote.png", lambda p: fig_d6(p)), ("d7_paridad_cype.png", lambda p: fig_d7(D, p)),
            ("d8_interaccion_pilote.png", lambda p: fig_d8(inter, p))]
    for name, fn in jobs:
        if name == "d4_fisuracion_psi2.png" and not D["curves"]:
            continue
        fn(fig_dir / name)
        out.append(fig_dir / name)
    return out


# =============================================================================================
# main
# =============================================================================================
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--reuse", action="store_true", help="reuse diseno/python/output/pilotes.json and vigas.json if present")
    ap.add_argument("--no-figs", action="store_true", help="skip the PNG figures")
    ap.add_argument("--no-curves", action="store_true", help="skip the ψ2 curves (needs the SAP tables, ~10 s)")
    ap.add_argument("--out", default=str(OUT), help="output directory of the workbook / md / json")
    ap.add_argument("--fig-dir", default=str(FIG), help="output directory of the figures")
    a = ap.parse_args(argv)
    t0 = time.time()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    P, V = load_results(a.reuse, out)
    print(f"[{time.time() - t0:5.1f} s] results {'reused' if a.reuse else 'computed'}")
    curves = None if a.no_curves else psi2_curves(V)
    conf = [confinement_estimate(150.0), confinement_estimate(PILE_TIE_S_NEW)]
    D = collect(P, V, curves, conf)
    print(f"[{time.time() - t0:5.1f} s] final design collected ({len(D['warnings'])} warnings)")
    for w in D["warnings"]:
        print("  WARNING:", w)
    x = write_xlsx(D, P, V, out / XLSX.name)
    m = write_md(D, P, V, out / MD.name)
    inter = None
    figs = []
    if not a.no_figs:
        try:
            inter = pile_interaction(P)
            figs = write_figures(D, P, V, inter, Path(a.fig_dir))
        except ImportError as e:        # matplotlib missing
            print("figures skipped:", e)
    js = {k: v for k, v in D.items() if k not in ("parity",)}
    js["interaction"] = inter
    (out / JSON_OUT.name).write_text(json.dumps(js, indent=1, ensure_ascii=False, default=float), encoding="utf-8")
    print(f"[{time.time() - t0:5.1f} s] written {x}, {m}, {out / JSON_OUT.name}" + (f", {len(figs)} figures in "
                                                                                    f"{a.fig_dir}" if figs else ""))
    for s in D["schedule"]:
        print(f"  {s['element'][:52]:52s} η {s['eta']:.3f}  {s['verdict']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
