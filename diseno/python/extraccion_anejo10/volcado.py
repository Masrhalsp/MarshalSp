"""Text dumps of Anejo 10 (report/A10_CALC_ESTRUCTURAS.docx), step 1 of the CYPE-reference extraction.

Purpose
    python-docx's plain `Paragraph.text` / `Table.cell.text` loses what the CYPE listings need:
    Symbol-font glyphs (η, σ, ν, ρ, λ, ϕ, θ, α, ε, γ, ≤, ≥ are stored as w:sym or as private-use
    characters U+F0xx), the tables nested inside table cells ('Datos del pilar', 'Datos de la viga')
    and the position of the images.  The functions below walk the document XML in order and write
    four kinds of text dump.  Paragraph numbers P<n> are the index in python-docx
    Document.paragraphs (the numbering used by every `src` in diseno/ref/cype_design_reference.json).

    volcado_rango(doc, lo, hi)     'P<n>\\t<text>' lines + 'TABLE>>' ... '<<TABLE' blocks for the
                                   tables after P<n>, cells joined by ' | ', nested tables as
                                   '{row ;; row}', images as '[IMG imageNN.ext]'.  Input of
                                   construir.py (ranges RANGOS: body P215-P609, Apéndice 1
                                   P1449-P1835, project text P100-P214).
    volcado_formulas(doc, lo, hi)  same walk, but every WMF equation image is replaced by its text
                                   recovered with wmftext.py: '«VRd,max=αcw·bw·z·ν1·fcd·...»'.
                                   Used to read CYPE's formulas (codigo_estructural_spec.md).
    volcado_a10_full(doc)          whole document with python-docx .text only (no glyphs, no nested
                                   tables); the 'a10_full.txt' whose line numbers are quoted in
                                   diseno/python/investigacion/ and diseno/sap/sap_design_spec.md.
    volcado_independiente(doc)     whole document, written by a second, independent walker (other
                                   glyph table, nested tables as '<<NESTED ... NESTED>>'); used only
                                   by validar.py to re-check the JSON with a different parser.

Input
    report/A10_CALC_ESTRUCTURAS.docx (13.7 MB, committed in the repo).

Output
    strings; extraer.py writes them to the (git-ignored) work directory
    diseno/python/extraccion_anejo10/trabajo/.

Run (stand-alone)
    python3 diseno/python/extraccion_anejo10/volcado.py 215 609        # prints the P215-P609 dump
"""

from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]                                   # repo root (diseno/python/extraccion_anejo10 -> ../../..)
DOCX = ROOT / "report" / "A10_CALC_ESTRUCTURAS.docx"

sys.path.insert(0, str(HERE))
import wmftext  # noqa: E402

# dump file name -> paragraph range (inclusive)
RANGOS = {"trasmallo_body.txt": (215, 609),              # Trasmallo results in the body
          "trasmallo_app.txt": (1449, 1835),             # Apéndice 1 'Listados de cálculo muelle de trasmallo'
          "trasmallo_pre.txt": (100, 214)}               # project text: codes, materials, covers, model description
RANGOS_FORMULAS = {"trasmallo_body_formulas.txt": (217, 609),
                   "trasmallo_app_formulas.txt": (1449, 1835)}

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def abrir(path=DOCX):
    return docx.Document(str(path))


# ------------------------------------------------------------------------------------------------
# 1. main dump (input of construir.py)
# ------------------------------------------------------------------------------------------------
SYM = {'\xa3': '≤', '\xb3': '≥', '\xa5': '∞', '\xb0': '°', '\xb4': '×', '\xd7': '·', '\xa2': '′'}
SYM.update(zip('abgdezhqiklmnxoprstufcywjvAGDQLXPSFYWJ', 'αβγδεζηθικλμνξοπρστυφχψωϕϖΑΓΔΘΛΞΠΣΦΨΩϑ'))


def _sym_char(s):
    c = s.get(W + 'char')
    f = s.get(W + 'font')
    try:
        code = int(c, 16)
    except (TypeError, ValueError):
        return '?'
    if code >= 0xF000:
        code -= 0xF000
    ch = chr(code)
    if f and 'Symbol' in f:
        return SYM.get(ch, ch)
    return '{%s:%s}' % (f, c)


def _ptext(p, rels):
    out = []
    for x in p.iter():
        if x.tag == W + 't':
            out.append(''.join((SYM.get(chr(ord(c) - 0xF000), '{U+%04X}' % ord(c)) if 0xF000 <= ord(c) <= 0xF0FF else c)
                               for c in (x.text or '')))
        elif x.tag == W + 'sym':
            out.append(_sym_char(x))
        elif x.tag == W + 'tab':
            out.append(' ')
        elif x.tag == qn('a:blip'):
            rid = x.get(qn('r:embed'))
            out.append('[IMG %s]' % rels[rid].target_ref.split('/')[-1] if rid in rels else '[IMG?]')
        elif x.tag == '{urn:schemas-microsoft-com:vml}imagedata':
            rid = x.get(qn('r:id'))
            out.append('[IMG %s]' % rels[rid].target_ref.split('/')[-1] if rid in rels else '[IMG?]')
    return ''.join(out)


def _cell_text(tc, rels):
    """A cell may hold paragraphs and nested tables: flatten them in order."""
    parts = []
    for ch in tc.iterchildren():
        if ch.tag == W + 'p':
            t = _ptext(ch, rels).strip()
            if t:
                parts.append(t)
        elif ch.tag == W + 'tbl':
            parts.append('{' + _tbl_text(ch, rels, inline=True) + '}')
    return ' / '.join(parts)


def _tbl_text(tbl, rels, inline=False):
    rows = []
    for tr in tbl.iterchildren(W + 'tr'):
        rows.append(' | '.join(_cell_text(tc, rels) for tc in tr.iterchildren(W + 'tc')))
    return (' ;; ' if inline else '\n').join(rows)


def volcado_rango(doc, lo, hi):
    rels = doc.part.rels
    lines = []
    pidx = -1
    for el in doc.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            pidx += 1
            if lo <= pidx <= hi:
                lines.append('P%d\t%s' % (pidx, _ptext(el, rels)))
        elif tag == 'tbl' and lo <= pidx <= hi:
            lines += ['TABLE>>', _tbl_text(el, rels), '<<TABLE']
    return '\n'.join(lines) + '\n'


# ------------------------------------------------------------------------------------------------
# 2. formula dump (WMF equation images -> text)
# ------------------------------------------------------------------------------------------------
def _formula(rel, cache):
    target = rel.target_ref
    if target not in cache:
        cache[target] = wmftext.render(rel.target_part.blob) if target.endswith('.wmf') else '[img:%s]' % target.split('/')[-1]
    return cache[target]


def _runs_text(el, rels, cache):
    out = []
    for node in el.iter():
        t = node.tag
        if t == qn('w:t'):
            out.append(node.text or '')
        elif t == qn('w:tab'):
            out.append(' ')
        elif t == qn('w:sym'):
            c = chr(int(node.get(qn('w:char')), 16) & 0xFF)
            out.append(wmftext.SYM.get(c, c))
        elif t == qn('a:blip'):
            rid = node.get(qn('r:embed'))
            if rid in rels:
                out.append(' «' + _formula(rels[rid], cache) + '» ')
        elif t == '{urn:schemas-microsoft-com:vml}imagedata':
            rid = node.get(qn('r:id'))
            if rid in rels:
                out.append(' «' + _formula(rels[rid], cache) + '» ')
    return ''.join(out)


def volcado_formulas(doc, lo, hi):
    rels = doc.part.rels
    cache = {}
    lines = []
    pidx = -1
    for el in doc.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            pidx += 1
            if lo <= pidx <= hi:
                s = _runs_text(el, rels, cache).strip()
                if s:
                    lines.append(f'P{pidx}\t{s}')
        elif tag == 'tbl' and lo <= pidx <= hi:
            lines.append('TABLE>>')
            for tr in el.iter(qn('w:tr')):
                cells = []
                for tc in tr.iter(qn('w:tc')):
                    s = ' '.join(_runs_text(p, rels, cache) for p in tc.iter(qn('w:p'))).strip()
                    if s and (not cells or cells[-1] != s):
                        cells.append(s)
                if cells:
                    lines.append('   ' + ' | '.join(cells))
            lines.append('<<TABLE')
    return '\n'.join(lines) + '\n'


def formulas_unicas(texto):
    """Distinct lines holding a recovered formula (image placeholders '[img:...]' excluded)."""
    seen, out = set(), []
    for ln in texto.split('\n'):
        if '«' in ln and 'img:' not in ln and ln not in seen:
            seen.add(ln)
            out.append(ln)
    return '\n'.join(out) + '\n'


# ------------------------------------------------------------------------------------------------
# 3. 'a10_full.txt' (python-docx .text only)
# ------------------------------------------------------------------------------------------------
def volcado_a10_full(doc):
    out = []
    pi = 0
    for el in doc.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            out.append(f'P{pi}\t{Paragraph(el, doc).text}')
            pi += 1
        elif tag == 'tbl':
            rows = []
            for r in Table(el, doc).rows:
                cells = []
                prev = None
                for c in r.cells:
                    if c._tc is prev:                      # merged cells repeat the same w:tc
                        continue
                    prev = c._tc
                    cells.append(c.text.replace('\n', ' / '))
                rows.append(' | '.join(cells))
            out.append('TABLE>>\n' + '\n'.join('   ' + x for x in rows) + '\n<<TABLE')
    return '\n'.join(out)


# ------------------------------------------------------------------------------------------------
# 4. independent dump (second parser, for validar.py only)
# ------------------------------------------------------------------------------------------------
_LOWER = {'a': 'α', 'b': 'β', 'c': 'χ', 'd': 'δ', 'e': 'ε', 'f': 'ϕ', 'g': 'γ', 'h': 'η', 'i': 'ι', 'j': 'φ', 'k': 'κ', 'l': 'λ',
          'm': 'μ', 'n': 'ν', 'o': 'ο', 'p': 'π', 'q': 'θ', 'r': 'ρ', 's': 'σ', 't': 'τ', 'u': 'υ', 'v': 'ϖ', 'w': 'ω', 'x': 'ξ',
          'y': 'ψ', 'z': 'ζ'}
_UPPER = {'A': 'Α', 'B': 'Β', 'C': 'Χ', 'D': 'Δ', 'E': 'Ε', 'F': 'Φ', 'G': 'Γ', 'H': 'Η', 'I': 'Ι', 'J': 'ϑ', 'K': 'Κ', 'L': 'Λ',
          'M': 'Μ', 'N': 'Ν', 'O': 'Ο', 'P': 'Π', 'Q': 'Θ', 'R': 'Ρ', 'S': 'Σ', 'T': 'Τ', 'U': 'Υ', 'V': 'ς', 'W': 'Ω', 'X': 'Ξ',
          'Y': 'Ψ', 'Z': 'Ζ'}
_OPS = {0xA3: '≤', 0xB3: '≥', 0xB4: '×', 0xB1: '±', 0xD7: '·', 0xB7: '•', 0xAE: '→', 0xD6: '√', 0xA5: '∞', 0xB0: '°',
        0xB9: '≠', 0xBB: '≈'}


def _i_symchar(el):
    c = el.get(qn('w:char'))
    f = el.get(qn('w:font'))
    v = int(c, 16)
    if v >= 0xF000:
        v -= 0xF000
    ch = chr(v)
    if f and 'Symbol' in f:
        if ch in _LOWER:
            return _LOWER[ch]
        if ch in _UPPER:
            return _UPPER[ch]
        if v in _OPS:
            return _OPS[v]
    return '{sym:%s:%s}' % (f, c)


def _i_runfont(t):
    r = t.getparent()
    rpr = r.find(qn('w:rPr')) if r is not None else None
    if rpr is None:
        return None
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        return None
    return rf.get(qn('w:ascii')) or rf.get(qn('w:hAnsi'))


def _i_symmap(txt):
    o = []
    for ch in txt:
        v = ord(ch)
        if v >= 0xF000:
            v -= 0xF000
            ch = chr(v)
        if ch in _LOWER:
            o.append(_LOWER[ch])
        elif ch in _UPPER:
            o.append(_UPPER[ch])
        elif v in _OPS:
            o.append(_OPS[v])
        else:
            o.append(ch)
    return ''.join(o)


def _i_ptext(p):
    out = []
    for el in p.iter():
        t = el.tag
        if t == qn('w:t'):
            f = _i_runfont(el)
            out.append(_i_symmap(el.text or '') if f and 'Symbol' in f else (el.text or ''))
        elif t == qn('w:sym'):
            out.append(_i_symchar(el))
        elif t == qn('w:tab'):
            out.append('\t')
        elif t.endswith('}blip'):
            out.append('[IMG %s]' % el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'))
    return ''.join(out)


def _i_celltext(tc, depth):
    parts = []
    for ch in tc:
        if ch.tag == qn('w:p'):
            parts.append(_i_ptext(ch))
        elif ch.tag == qn('w:tbl'):
            parts.append('<<NESTED ' + ' // '.join(_i_rowtext(tr, depth + 1) for tr in ch.findall(qn('w:tr'))) + ' NESTED>>')
    return ' '.join(x for x in parts if x)


def _i_rowtext(tr, depth):
    return ' | '.join(_i_celltext(tc, depth) for tc in tr.findall(qn('w:tc')))


def volcado_independiente(doc):
    lines = []
    pi = -1
    for ch in doc.element.body:
        if ch.tag == qn('w:p'):
            pi += 1
            lines.append(f'P{pi}\t{_i_ptext(ch)}')
        elif ch.tag == qn('w:tbl'):
            lines.append('TABLE>>')
            for tr in ch.findall(qn('w:tr')):
                lines.append('   ' + _i_rowtext(tr, 0))
            lines.append('<<TABLE')
    return '\n'.join(lines) + '\n'


def todos(doc, formulas=True):
    """All dumps: {file name: text}."""
    out = {name: volcado_rango(doc, lo, hi) for name, (lo, hi) in RANGOS.items()}
    out["a10_full.txt"] = volcado_a10_full(doc)
    out["a10_independiente.txt"] = volcado_independiente(doc)
    if formulas:
        for name, (lo, hi) in RANGOS_FORMULAS.items():
            out[name] = volcado_formulas(doc, lo, hi)
        out["formulas_unicas.txt"] = formulas_unicas(out["trasmallo_body_formulas.txt"] + out["trasmallo_app_formulas.txt"])
    return out


if __name__ == '__main__':
    lo, hi = int(sys.argv[1]), int(sys.argv[2])
    sys.stdout.write(volcado_rango(abrir(), lo, hi))
